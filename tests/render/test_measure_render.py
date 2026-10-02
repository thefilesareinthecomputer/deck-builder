"""Render tier: real pdftotext word boxes against qa.measure's word-to-shape matching.

These build a minimal .pptx directly with python-pptx (no brand kit needed) so the geometry is
exact, render it with LibreOffice, and measure the real PDF word boxes.
"""
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt

from deck_builder.qa import backends, measure, tools

pytestmark = [
    pytest.mark.render,
    pytest.mark.skipif(tools.soffice() is None or bool(tools.poppler_missing()),
                       reason="needs LibreOffice and poppler"),
]


def _render(tmp_path: Path, build) -> list:
    prs = Presentation()
    prs.slide_width = Pt(960)
    prs.slide_height = Pt(540)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    build(slide)
    pptx = tmp_path / "probe.pptx"
    prs.save(str(pptx))
    pdf = tmp_path / "probe.pdf"
    backends.libreoffice(pptx, pdf)
    return measure.overflow(pptx, pdf)


def test_a_mid_word_line_break_is_still_measured_as_overflow(tmp_path):
    """A single unbroken token too long for its line gets split by the renderer; neither rendered
    fragment equals the source token, so the old exact-match-only logic reported nothing."""
    def build(slide):
        box = slide.shapes.add_textbox(Pt(50), Pt(50), Pt(150), Pt(40))
        box.name = "Probe Box"
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.text = "abcdefghijklmnopqrstuvwxyz0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"

    issues = _render(tmp_path, build)
    assert any(i.code == "OVERFLOW_MEASURED" and "Probe Box" in i.message for i in issues)


def test_a_body_word_spilling_over_a_chart_is_still_measured(tmp_path):
    """A short text box overflows down into a chart placed right below it: the overflowing words'
    centers land inside the chart's box, which used to discard them outright."""
    def build(slide):
        body = slide.shapes.add_textbox(Pt(50), Pt(50), Pt(300), Pt(30))
        body.name = "Body Box"
        tf = body.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.text = "Overflowing body text that runs down several lines past its short box into the chart area below"
        data = CategoryChartData()
        data.categories = ["Q1", "Q2"]
        data.add_series("Series", (1, 2))
        slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Pt(50), Pt(90), Pt(300), Pt(200), data)

    issues = _render(tmp_path, build)
    assert any(i.code == "OVERFLOW_MEASURED" and "Body Box" in i.message for i in issues)
