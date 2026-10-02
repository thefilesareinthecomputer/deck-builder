"""Render tier: a buAutoNum list's own marker must never be mistaken for a shape's text.

Builds a minimal .pptx directly with python-pptx (no brand kit needed), renders it with
LibreOffice, and measures the real PDF word boxes.
"""
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn
from pptx.util import Pt

from deck_builder.qa import backends, measure, tools

pytestmark = [
    pytest.mark.render,
    pytest.mark.skipif(tools.soffice() is None or bool(tools.poppler_missing()),
                       reason="needs LibreOffice and poppler"),
]


def test_an_auto_numbered_marker_is_never_mistaken_for_the_slide_number(tmp_path: Path):
    """A buAutoNum list renders "1." "2." "3." as their own PDF words, far from any shape's stored
    text; "2." on slide 2 used to be matched, by token, to the slide number placeholder's literal "2"."""
    prs = Presentation()
    prs.slide_width = Pt(960)
    prs.slide_height = Pt(540)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    box = slide.shapes.add_textbox(Pt(50), Pt(50), Pt(400), Pt(150))
    box.name = "Numbered Body"
    txBody = box.text_frame._txBody
    for p in txBody.findall(qn("a:p")):
        txBody.remove(p)
    for item in ("First consideration for the plan", "Second consideration for the plan",
                "Third consideration for the plan"):
        txBody.append(parse_xml(
            f'<a:p {nsdecls("a")}><a:pPr marL="457200" indent="-457200"><a:buFont typeface="+mj-lt"/>'
            f'<a:buAutoNum type="arabicPeriod"/></a:pPr><a:r><a:t>{item}</a:t></a:r></a:p>'))

    # a slide-number placeholder far below, with literal text "2" (as deck-builder's own furniture
    # writes it): a bare list marker must never be attributed to it.
    numbox = slide.shapes.add_textbox(Pt(20), Pt(500), Pt(30), Pt(20))
    numbox.name = "Slide Number Placeholder 4"
    numbox._element.nvSpPr.nvPr.append(parse_xml(f'<p:ph {nsdecls("p")} type="sldNum"/>'))
    numbox.text_frame.text = "2"

    pptx = tmp_path / "probe.pptx"
    prs.save(str(pptx))
    pdf = tmp_path / "probe.pdf"
    backends.libreoffice(pptx, pdf)
    assert measure.overflow(pptx, pdf) == []
