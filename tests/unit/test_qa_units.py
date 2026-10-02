"""QA pieces that don't need a renderer: empty placeholders, font matching, PowerPoint failure paths."""
import subprocess
from pathlib import Path

import pytest

from conftest import cli_json, write_deck
from deck_builder.errors import EnvError
from deck_builder.qa import backends, fonts
from deck_builder.qa.render import empty_placeholders, stale_build_check


def test_an_empty_text_field_is_left_out_rather_than_drawn_empty(ws, capsys):
    deck = write_deck(ws, "## Hello\nlayout: title\nsubtitle: ''\n")
    code, out = cli_json(ws, "build", str(deck), capsys=capsys)
    assert code == 0
    assert empty_placeholders(Path(out["output"])) == []


def test_an_empty_placeholder_is_flagged(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title and content; the content left empty
    slide.shapes.title.text = "Hello"
    prs.save(str(tmp_path / "x.pptx"))
    assert [(i.code, i.slide) for i in empty_placeholders(tmp_path / "x.pptx")] == [("EMPTY_PLACEHOLDER", 1)]


def test_font_check_reports_fallback_and_substitution(monkeypatch):
    meta = {"fonts": {"heading": {"family": "Inter", "fallback": "Liberation Sans"},
                      "body": {"family": "Source Serif", "fallback": "Georgia"}}}
    monkeypatch.setattr(fonts, "embedded", lambda pdf: {"liberationsansbold", "dejavusans"})
    issues = fonts.check(Path("x.pdf"), meta)
    assert [i.code for i in issues] == ["MISSING_FONT", "MISSING_FONT"]
    assert "fallback" in issues[0].message and "isn't installed" in issues[1].message
    assert all(i.severity == "warning" for i in issues)


def test_font_check_passes_when_the_family_is_used(monkeypatch):
    monkeypatch.setattr(fonts, "embedded", lambda pdf: {"interbold", "intermedium"})
    assert fonts.check(Path("x.pdf"), {"fonts": {"heading": {"family": "Inter"}}}) == []


def test_stale_build_warns_when_the_source_is_newer_than_the_pptx(tmp_path):
    import os
    import time

    source = tmp_path / "deck.md"
    pptx = tmp_path / "deck.pptx"
    pptx.write_bytes(b"pptx")
    source.write_text("stuff")
    now = time.time()
    os.utime(pptx, (now - 10, now - 10))
    os.utime(source, (now, now))
    manifest = {"input": {"path": str(source)}}
    issues = stale_build_check(pptx, manifest)
    assert [i.code for i in issues] == ["STALE_BUILD"]
    assert issues[0].severity == "warning"


def test_stale_build_is_quiet_when_the_pptx_is_newer(tmp_path):
    import os
    import time

    source = tmp_path / "deck.md"
    pptx = tmp_path / "deck.pptx"
    source.write_text("stuff")
    now = time.time()
    os.utime(source, (now - 10, now - 10))
    pptx.write_bytes(b"pptx")
    os.utime(pptx, (now, now))
    manifest = {"input": {"path": str(source)}}
    assert stale_build_check(pptx, manifest) == []


def test_stale_build_is_quiet_without_a_manifest_or_path():
    assert stale_build_check(Path("x.pptx"), None) == []
    assert stale_build_check(Path("x.pptx"), {"input": {}}) == []


def test_auto_backend_prefers_libreoffice_over_unverified_powerpoint(monkeypatch):
    monkeypatch.setattr(backends.tools, "soffice", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(backends.tools, "powerpoint", lambda: True)
    assert backends.choose("auto") == "libreoffice"


def test_auto_backend_falls_back_to_powerpoint_without_libreoffice(monkeypatch):
    monkeypatch.setattr(backends.tools, "soffice", lambda: None)
    monkeypatch.setattr(backends.tools, "powerpoint", lambda: True)
    assert backends.choose("auto") == "powerpoint"


class FakeProc:
    def __init__(self, returncode=0, stdout="", stderr=""):
        self.returncode, self.stdout, self.stderr = returncode, stdout, stderr


@pytest.fixture
def container(tmp_path, monkeypatch):
    monkeypatch.setattr(backends, "PP_CONTAINER", tmp_path / "container")
    monkeypatch.setattr(backends, "_running", lambda: True)
    pptx = tmp_path / "deck.pptx"
    pptx.write_bytes(b"pptx")
    return pptx


def test_powerpoint_failure_reports_office_repair(container, monkeypatch, tmp_path):
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: FakeProc(returncode=1, stderr="error -1708"))
    with pytest.raises(EnvError) as e:
        backends.powerpoint(container, tmp_path / "out.pdf")
    assert e.value.code == "OFFICE_REPAIR"
    assert not (backends.PP_CONTAINER / ".lock").exists()


def test_powerpoint_permission_denied_says_where_to_grant_it(container, monkeypatch, tmp_path):
    monkeypatch.setattr(subprocess, "run", lambda *a, **k: FakeProc(returncode=1, stderr="Not authorized (-1743)"))
    with pytest.raises(EnvError) as e:
        backends.powerpoint(container, tmp_path / "out.pdf")
    assert "Privacy & Security > Automation" in str(e.value)


def test_powerpoint_timeout_leaves_it_running_and_reports(container, monkeypatch, tmp_path):
    def slow(*a, **k):
        raise subprocess.TimeoutExpired(cmd="osascript", timeout=1)

    monkeypatch.setattr(subprocess, "run", slow)
    with pytest.raises(EnvError) as e:
        backends.powerpoint(container, tmp_path / "out.pdf")
    assert e.value.code == "OFFICE_REPAIR" and "left running" in str(e.value)

