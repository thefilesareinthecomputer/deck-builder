"""`current: n`: the returning map slide of a deck in parts marks where the audience is."""
from __future__ import annotations

import pytest
import yaml
from pptx import Presentation

from conftest import cli_json, codes, write_deck
from deck_builder.brand.kit import contrast
from deck_builder.cli import main
from deck_builder.parse import markdown

BRAND = {"spec_version": 1, "name": "Map", "slug": "map", "version": "1.0.0",
         "palette": {"primary": "1F3A5F", "accent": "E07A2F", "ink": "1B1B1B", "muted": "6B7280",
                     "surface": "F2F2F2", "background": "FFFFFF"},
         "fonts": {"heading": {"family": "Arial"}, "body": {"family": "Arial"}},
         "generate": {"layout_set": "designed"}}

AGENDA = """
---
brand: map
---

## Where we are
layout: agenda
current: 2

- The pilot
- The first bay
- Tester bars
"""

CARDS = """
---
brand: map
---

## The second part: the first bay
layout: cards-3
current: 2
label1: The pilot
label2: The first bay
label3: Tester bars

### body1
- Six stores

### body2
- Sales held

### body3
- $2.40 a week
"""


@pytest.fixture
def root(tmp_path):
    assert main(["init", "--dir", str(tmp_path)]) == 0
    (tmp_path / "map.yaml").write_text(yaml.safe_dump(BRAND), encoding="utf-8")
    assert main(["--config", str(tmp_path / "deck-builder.toml"), "brand", "init", "map", "--from",
                 str(tmp_path / "map.yaml")]) == 0
    (tmp_path / "decks").mkdir()
    return tmp_path


def build(root, body, capsys, name="deck.md"):
    deck = write_deck(root, body, name=name)
    capsys.readouterr()
    code, out = cli_json(root, "build", str(deck), "-o", str(root / "out" / name.replace(".md", ".pptx")),
                         capsys=capsys)
    assert code == 0, out["issues"]
    return deck, Presentation(out["output"])


def test_current_on_an_agenda_mutes_every_other_item_at_readable_contrast(root, capsys):
    _, prs = build(root, AGENDA, capsys)
    slide = prs.slides[0]
    body = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 1)
    colors = [p.runs[0].font.color.rgb if p.runs[0].font.color.type else None for p in body.text_frame.paragraphs]
    assert colors[1] is None and colors[0] == colors[2]  # the current item keeps the text color
    assert contrast(str(colors[0]), "FFFFFF") >= 4.5
    assert slide._element.cSld.get("name") == "current 2"


def test_current_on_cards_leaves_one_label_in_the_primary_ramp(root, capsys):
    _, prs = build(root, CARDS, capsys)
    labels = {sh.text_frame.text: sh for sh in prs.slides[0].placeholders if sh.text_frame.text in
              ("The pilot", "The first bay", "Tester bars")}
    assert labels["The first bay"]._element.spPr.find(
        "{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill") is None  # the layout's ramp fill
    for other in ("The pilot", "Tester bars"):
        assert labels[other].fill.fore_color.theme_color.name == "BACKGROUND_2"


@pytest.mark.parametrize("value,why", [("4", "from 1 to 3"), ("two", "from 1 to 3"), ("1", None)])
def test_current_must_name_an_item_that_exists(root, capsys, value, why):
    deck = write_deck(root, AGENDA.replace("current: 2", f"current: {value}"))
    capsys.readouterr()
    _, out = cli_json(root, "check", str(deck), capsys=capsys)
    if why:
        assert codes(out) == ["PARSE"] and why in out["issues"][0]["message"]
    else:
        assert codes(out) == []
    deck = write_deck(root, "---\nbrand: map\n---\n\n## One number\nlayout: big-number\ncurrent: 1\ncaption: x\n",
                      name="bad.md")
    _, out = cli_json(root, "check", str(deck), capsys=capsys)
    assert codes(out) == ["PARSE"] and "has neither" in out["issues"][0]["message"]


def test_current_round_trips_through_markdown_the_workbook_and_import(root, capsys):
    deck, prs = build(root, CARDS, capsys)
    source = markdown.parse(deck)[0]
    assert source.slides[0].current == 2
    _, out = cli_json(root, "convert", str(deck), str(root / "decks" / "deck.xlsx"), capsys=capsys)
    assert out["ok"], out["issues"]  # convert refuses any change on the way back
    _, out = cli_json(root, "import", str(root / "out" / "deck.pptx"), str(root / "imported"), "--brand", "map",
                      capsys=capsys)
    assert out["ok"], out
    assert markdown.parse(root / "imported" / "deck.md")[0].slides[0].current == 2
