import hashlib
from importlib import resources
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_check import GOOD

from conftest import cli_json, codes, write_deck
from deck_builder import pipeline


def example_decks():
    root = resources.files("deck_builder") / "data" / "decks"
    return [p for p in root.iterdir() if (p / "deck.md").is_file()]


def model(path):
    loaded = pipeline.load_deck(Path(path))
    assert loaded.issues == []
    return loaded.deck


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def roundtrip(ws, capsys, src, chain):
    """Convert src through each extension in chain; return the paths produced."""
    out, cur = [], src
    for i, ext in enumerate(chain):
        nxt = ws / "decks" / f"rt{i}{ext}"
        code, res = cli_json(ws, "convert", str(cur), str(nxt), capsys=capsys)
        assert code == 0, res
        out.append(nxt)
        cur = nxt
    return out


def test_md_to_xlsx_to_md_keeps_the_model_and_the_build(ws, capsys):
    src = write_deck(ws, GOOD)
    xlsx, md = roundtrip(ws, capsys, src, [".xlsx", ".md"])
    assert model(src) == model(xlsx) == model(md)
    builds = []
    for p in (src, xlsx, md):
        code, res = cli_json(ws, "build", str(p), "-o", str(ws / "out" / f"{p.stem}.pptx"), capsys=capsys)
        assert code == 0, res["issues"]
        builds.append(res)
    # one deck, three files in two formats: the same PPTX bytes every time
    assert len({sha(b["output"]) for b in builds}) == 1


def test_canonical_forms_are_byte_stable(ws, capsys):
    src = write_deck(ws, GOOD)
    md1, xlsx1, md2, xlsx2 = roundtrip(ws, capsys, src, [".md", ".xlsx", ".md", ".xlsx"])
    assert md1.read_bytes() == md2.read_bytes()
    assert sha(xlsx1) == sha(xlsx2)


@pytest.mark.parametrize("deck_dir", example_decks(), ids=lambda p: p.name)
def test_examples_round_trip(tmp_path, capsys, deck_dir):
    from deck_builder.cli import main

    main(["init", "--dir", str(tmp_path), "--json"])
    capsys.readouterr()
    cfg = tmp_path / "deck-builder.toml"
    src = tmp_path / "workspace" / "decks" / deck_dir.name / "deck.md"
    xlsx = src.with_name("deck.xlsx")
    assert main(["--config", str(cfg), "convert", str(src), str(xlsx)]) == 0
    capsys.readouterr()
    assert model(src) == model(xlsx)


def test_workbook_is_set_up_for_a_team(ws, capsys):
    src = write_deck(ws, GOOD)
    (xlsx,) = roundtrip(ws, capsys, src, [".xlsx"])
    wb = load_workbook(xlsx)
    assert wb.sheetnames[:2] == ["slides", "deck"]
    assert "README" in wb.sheetnames and wb["_lists"].sheet_state == "hidden"
    ws_ = wb["slides"]
    header = [c.value for c in ws_[1]]
    assert header[:3] == ["slide", "layout", "title"] and "#chars:body" in header
    assert ws_.freeze_panes == "D2"
    assert ws_.data_validations.dataValidation[0].formula1.startswith("='_lists'!")
    assert any(v.startswith("sheet:chart-") for row in ws_.iter_rows(values_only=True) for v in row
               if isinstance(v, str))


def test_text_starting_with_equals_stays_text(ws, capsys):
    src = write_deck(ws, '## A\nlayout: title\nsubtitle: "=SUM(A1:A3) is not a formula"\n')
    (xlsx,) = roundtrip(ws, capsys, src, [".xlsx"])
    assert model(xlsx).slides[0].fields["subtitle"] == "=SUM(A1:A3) is not a formula"


def test_csv_refuses_charts_and_front_matter(ws, capsys):
    src = write_deck(ws, GOOD)
    code, out = cli_json(ws, "convert", str(src), str(ws / "decks" / "x.csv"), capsys=capsys)
    assert code == 1 and set(codes(out)) == {"CONVERT_LOSSY"}
    assert not (ws / "decks" / "x.csv").exists()


def test_text_only_deck_converts_to_csv_and_back(ws, capsys):
    src = write_deck(ws, "## Hello\nlayout: content\n\n- one\n  - two\n\nNotes:\nsource\n")
    csv_, md = roundtrip(ws, capsys, src, [".csv", ".md"])
    assert model(src) == model(csv_) == model(md)


def test_convert_refuses_to_overwrite(ws, capsys):
    src = write_deck(ws, GOOD)
    target = ws / "decks" / "out.xlsx"
    target.write_text("keep me")
    code, out = cli_json(ws, "convert", str(src), str(target), capsys=capsys)
    assert code == 2 and target.read_text() == "keep me"


def test_bulk_from_a_workbook_template(ws, capsys):
    tmpl = write_deck(ws, "---\nbrand: stock\ntitle: '{{client}}'\n---\n\n## {{client}} review\nlayout: title\n",
                      name="pitch.md")
    (xlsx,) = roundtrip(ws, capsys, tmpl, [".xlsx"])
    data = ws / "decks" / "clients.csv"
    data.write_text("client\nHalvorsen Freight\nCopperline Foods\n")
    code, out = cli_json(ws, "build", str(xlsx), "--data", str(data), "--name", "{{client}}.pptx", capsys=capsys)
    assert code == 0, out["issues"]
    assert len(out["outputs"]) == 2
