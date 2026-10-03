"""Code blocks: the fenced format, highlighting, the code layouts and their tokens, check, build, convert and
import. The fixture deck in demo-brands/code/ has the cases in all three brands (test_demo_brands, test_import
and the render tier)."""
from __future__ import annotations

import hashlib
import json
import textwrap
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

from conftest import cli_json, codes, write_deck
from deck_builder import highlight
from deck_builder import template as tpl
from deck_builder.brand.generate import generate
from deck_builder.brand.kit import contrast
from deck_builder.brand.layouts import SETS
from deck_builder.cli import main
from deck_builder.model import Code
from deck_builder.parse import markdown
from deck_builder.write import markdown as md_writer

DEMO = Path(__file__).resolve().parents[1] / "fixtures" / "demo-brands"
SLUGS = ("dumbder-nifftlin", "cubicle-nine", "soap-club")
BRAND = {"spec_version": 1, "name": "Code", "slug": "codes", "version": "1.0.0",
         "palette": {"primary": "1F3A5F", "accent": "E07A2F", "ink": "1B1B1B", "muted": "6B7280",
                     "surface": "F2F2F2", "background": "FFFFFF"},
         "fonts": {"heading": {"family": "Arial"}, "body": {"family": "Arial"}}}


def kit(root: Path, fonts: dict | None = None, **gen: object) -> Path:
    """A workspace at root with brand `codes` generated with the full set and these generate options."""
    assert main(["init", "--dir", str(root)]) == 0
    src = root / "codes-src"
    src.mkdir()
    meta = {**BRAND, "fonts": {**BRAND["fonts"], **(fonts or {})}, "generate": {"layout_set": "full", **gen}}
    (src / "brand.yaml").write_text(yaml.safe_dump(meta), encoding="utf-8")
    assert main(["--config", str(root / "deck-builder.toml"), "brand", "init", "codes", "--from",
                 str(src / "brand.yaml")]) == 0
    (root / "decks").mkdir()
    return root


def tokens(root: Path) -> dict:
    return yaml.safe_load((root / "workspace" / "brands" / "codes" / "tokens.yaml").read_text())


def parse(tmp_path: Path, body: str) -> tuple:
    p = tmp_path / "deck.md"
    p.write_text(textwrap.dedent(body).lstrip("\n"), encoding="utf-8")
    return markdown.parse(p)


def deck_md(slides: str) -> str:
    return "---\nbrand: codes\n---\n\n" + textwrap.dedent(slides).lstrip("\n")


# ---------------------------------------------------------------- the format


def test_a_fenced_block_with_a_language_and_options_is_code(tmp_path):
    deck, issues = parse(tmp_path, """
        ## The query
        layout: code

        ```sql {2,4-5} lines title="orders.sql"
        SELECT region,
          week,
          units
        FROM orders
        WHERE open
        ```
        """)
    assert issues == []
    assert deck.slides[0].fields["body"] == Code(language="sql", text="SELECT region,\n  week,\n  units\nFROM orders\n"
                                                 "WHERE open", highlight=[2, 4, 5], numbers=True, title="orders.sql")


def test_fences_nest_by_length_and_character_and_hide_headings_and_notes_inside(tmp_path):
    """CommonMark: a fence closes only on the same character at least as many times, so a four-backtick block
    holds a markdown example with its own fence; a heading or `Notes:` line inside a block is code."""
    deck, issues = parse(tmp_path, """
        ## Runbook
        layout: code

        ````markdown
        ## Rerun the forecast
        ```bash
        uv run forecast
        ```
        Notes:
        ````

        Notes:
        The real notes.

        ## Second slide
        layout: code

        ~~~
        ```
        not closed by backticks
        ~~~
        """)
    assert issues == []
    assert len(deck.slides) == 2
    first = deck.slides[0]
    assert first.fields["body"].text == "## Rerun the forecast\n```bash\nuv run forecast\n```\nNotes:"
    assert first.notes == "The real notes."
    assert deck.slides[1].fields["body"] == Code(language="", text="```\nnot closed by backticks")


def test_tabs_become_spaces_and_edges_are_trimmed(tmp_path):
    deck, _ = parse(tmp_path, "## T\nlayout: code\n\n```python\n\nif x:   \n\treturn 1\n\n```\n")
    assert deck.slides[0].fields["body"].text == "if x:\n    return 1"


@pytest.mark.parametrize("fence,why", [
    ("```python {9}", "highlights line 9"),
    ("```python {1-999999}", "highlights line 999999"),  # refused before it's expanded, so no huge list
    ("```python {1-9999999}", "unknown option"),  # more digits than any block has lines
    ("```python linez", "unknown option 'linez'"),
    ('```python title="unclosed', "can't read the options"),
])
def test_bad_fence_options_are_parse_errors(tmp_path, fence, why):
    _, issues = parse(tmp_path, f"## T\nlayout: code\n\n{fence}\nx = 1\n```\n")
    assert [i.code for i in issues] == ["PARSE"]
    assert why in issues[0].message


def test_an_unclosed_code_block_is_a_parse_error(tmp_path):
    _, issues = parse(tmp_path, "## T\nlayout: code\n\n````python\nx = 1\n```\n")
    assert [i.code for i in issues] == ["PARSE"] and "unclosed ````python" in issues[0].message


@pytest.mark.parametrize("code", [
    Code("markdown", "# Title\n```bash\nls\n```\n````\nfour", [2], True, 'a "quoted" \\ name'),
    Code("python", "x = 1", title="tick`name.py"),  # a backtick in the info string needs a tilde fence
    Code("", "~~~\nplain", [1, 2]),
])
def test_writing_then_parsing_gives_back_the_same_block(tmp_path, code):
    from deck_builder.model import Deck, Slide

    text = md_writer.write(Deck(meta={"brand": "codes"}, slides=[Slide("T", "code", {"code": code})]))
    p = tmp_path / "deck.md"
    p.write_text(text, encoding="utf-8")
    back, issues = markdown.parse(p)
    assert issues == []
    assert back.slides[0].fields["code"] == code
    assert md_writer.write(back) == text


# ---------------------------------------------------------------- highlighting


def test_tokens_get_roles_and_an_unknown_language_is_plain():
    rows = highlight.lines(Code("python", 'def f():\n    return "x"  # done'))
    assert rows[0][0] == ("keyword", "def ")
    assert ("function", "f") in rows[0]
    assert ("string", '"x"  ') in rows[1] and rows[1][-1] == ("comment", "# done")
    assert highlight.lines(Code("nope", "a b\nc")) == [[("plain", "a b")], [("plain", "c")]]
    assert highlight.known("Python") and highlight.known("") and highlight.known("text")
    assert not highlight.known("nope")


def test_every_role_has_a_style_or_a_color_of_its_own():
    """Keywords bold and comments italic, so color is never the only signal (WCAG 2.2 SC 1.4.1)."""
    assert {"keyword"} == highlight.BOLD and {"comment"} == highlight.ITALIC
    assert set(highlight.ROLES) == {r for _, r in highlight.ROLE_OF} | {"plain"}


# ---------------------------------------------------------------- layouts and tokens


def test_the_code_layouts_are_in_the_full_and_designed_sets_only():
    for name in ("full", "designed"):
        assert {"code", "code-right"} <= set(SETS[name])
    for name in ("minimal", "standard"):
        assert not {"code", "code-right"} & set(SETS[name])


def brand_meta(slug: str, theme: str) -> dict:
    meta = yaml.safe_load((DEMO / "brands" / slug / "brand.yaml").read_text())
    meta["generate"]["code"] = {"theme": theme}
    return meta


PALE = {**BRAND, "palette": {"primary": "9DB4C0", "accent": "FFD000", "ink": "333333", "muted": "CCCCCC",
                             "surface": "F7F7F7", "background": "FFFFFF"}, "generate": {"layout_set": "full"}}


@pytest.mark.parametrize("slug", [*SLUGS, "pale"])
@pytest.mark.parametrize("theme", ["light", "dark"])
def test_every_code_color_reads_on_the_panel_and_on_a_highlighted_line(slug, theme):
    meta = dict(PALE, generate={"layout_set": "full", "code": {"theme": theme}}) if slug == "pale" \
        else brand_meta(slug, theme)
    _, tok = generate(meta, DEMO / "brands" / (slug if slug != "pale" else SLUGS[0]))
    palette = {k: str(v).upper() for k, v in meta["palette"].items()}
    code = tok["code"]
    panel, band = palette.get(code["panel"], code["panel"]), code["highlight"]
    for role, ref in code["colors"].items():
        hexv = palette.get(ref, ref)
        assert contrast(hexv, panel) >= 4.5, (role, hexv, panel)
        assert contrast(hexv, band) >= 4.5, (role, hexv, band)


def test_brand_init_writes_the_code_font_budgets_and_a_no_wrap_monospace_layout(tmp_path):
    root = kit(tmp_path, fonts={"code": {"family": "Consolas", "fallback": "Courier New"}})
    tok = tokens(root)
    assert tok["text"] == {"code_font": "Consolas"}
    assert tok["code"]["theme"] == "light" and tok["code"]["panel"] == "surface"
    for name in ("code", "code-right"):
        field = tok["layouts"][name]["fields"]["code"]
        assert field["kind"] == "code" and field["required"] and field["max_cols"] > 30 and field["max_lines"] >= 10
    prs = tpl.open_template(root / "workspace" / "brands" / "codes" / "template.potx")
    layout = next(lo for lo in prs.slide_layouts if lo.name == "Code")
    ph = next(p for p in layout.placeholders if p.placeholder_format.idx == 1)
    body = ph._element.find(f".//{qn('a:bodyPr')}")
    latin = ph._element.find(f".//{qn('a:latin')}")
    assert body.get("wrap") == "none"
    assert (latin.get("typeface"), latin.get("pitchFamily")) == ("Consolas", "49")
    assert ph._element.find(f".//{qn('a:lnSpc')}/{qn('a:spcPts')}") is not None


def test_a_read_kit_has_smaller_code_and_more_room(tmp_path):
    projected = tokens(kit(tmp_path / "p"))["layouts"]["code"]["fields"]["code"]
    read = tokens(kit(tmp_path / "r", mode="read"))["layouts"]["code"]["fields"]["code"]
    assert read["max_cols"] > projected["max_cols"] and read["max_lines"] > projected["max_lines"]


def test_brand_check_reports_a_code_color_under_4_5_to_1(tmp_path, capsys):
    root = kit(tmp_path)
    path = root / "workspace" / "brands" / "codes" / "tokens.yaml"
    tok = yaml.safe_load(path.read_text())
    tok["code"]["colors"]["string"] = "FFD000"
    path.write_text(yaml.safe_dump(tok, sort_keys=False))
    capsys.readouterr()
    _, out = cli_json(root, "brand", "check", "codes", capsys=capsys)
    assert any(i["code"] == "LOW_CONTRAST" and "code string" in i["message"] for i in out["issues"])


# ---------------------------------------------------------------- check


def check(root: Path, slides: str, capsys) -> dict:
    deck = write_deck(root, deck_md(slides))
    capsys.readouterr()
    return cli_json(root, "check", str(deck), capsys=capsys)[1]


def test_a_line_past_the_panel_and_more_lines_than_it_holds_are_code_long(tmp_path, capsys):
    root = kit(tmp_path)
    field = tokens(root)["layouts"]["code"]["fields"]["code"]
    cols, lines = field["max_cols"], field["max_lines"]
    fits = "x" * cols
    out = check(root, f"## Fits\nlayout: code\n\n```python\n{fits}\n```\n", capsys)
    assert "CODE_LONG" not in codes(out)
    out = check(root, f"## Numbers take columns\nlayout: code\n\n```python lines\n{fits}\n```\n", capsys)
    assert codes(out).count("CODE_LONG") == 1 and "beside its line numbers" in out["issues"][0]["message"]
    many = "\n".join(f"x = {n}" for n in range(lines))
    out = check(root, f"## Fits\nlayout: code\n\n```python\n{many}\n```\n", capsys)
    assert "CODE_LONG" not in codes(out)
    out = check(root, f'## A filename takes two\nlayout: code\n\n```python title="a.py"\n{many}\n```\n', capsys)
    assert codes(out).count("CODE_LONG") == 1 and out["issues"][0]["limit"] == lines - 2
    assert out["ok"] is False


def test_an_unknown_language_warns_and_still_builds(tmp_path, capsys):
    root = kit(tmp_path)
    out = check(root, "## Odd\nlayout: code\n\n```klingon\nqapla\n```\n", capsys)
    assert codes(out) == ["CODE_LANGUAGE"] and out["ok"] is True


def test_more_than_twelve_lines_warns_on_a_projected_deck_only(tmp_path, capsys):
    many = "\n".join(f"x{n} = {n}" for n in range(13))
    body = f"## Long\nlayout: code-right\n\n- One point\n\n### code\n```python\n{many}\n```\n"
    assert codes(check(kit(tmp_path / "p"), body, capsys)) == ["CODE_LINES_MANY"]
    assert codes(check(kit(tmp_path / "r", mode="read"), body, capsys)) == []


def test_code_doesnt_count_as_words_and_a_long_inline_span_warns(tmp_path, capsys):
    root = kit(tmp_path)
    words = "\n".join("alpha beta gamma delta epsilon zeta eta theta" for _ in range(10))
    assert codes(check(root, f"## Many tokens\nlayout: code\n\n```text\n{words}\n```\n", capsys)) == []
    span = "warehouse.orders." + "x" * 120
    out = check(root, f"## Inline\nlayout: content\n\n- Read `{span}` first\n", capsys)
    long = [i for i in out["issues"] if i["code"] == "CODE_LONG"]
    assert len(long) == 1 and long[0]["severity"] == "warning" and "mid-token" in long[0]["message"]


def test_code_in_a_text_field_is_a_kind_mismatch(tmp_path, capsys):
    root = kit(tmp_path)
    out = check(root, "## Wrong\nlayout: content\n\n```python\nx = 1\n```\n", capsys)
    assert "KIND_MISMATCH" in codes(out)


# ---------------------------------------------------------------- build


SLIDE = """
## The forecast
layout: code

```python {4-5} lines title="forecast.py"
def weekly_forecast(orders, weeks=12):
    history = orders.last(weeks)
    median = history.median()
    # one bulk order isn't a trend
    capped = history.clip(upper=3 * median)
    return capped.mean()
```
"""


def build(root: Path, body: str, capsys, name: str = "deck.pptx") -> tuple[dict, Path]:
    deck = write_deck(root, deck_md(body))
    capsys.readouterr()
    code, out = cli_json(root, "build", str(deck), "-o", str(root / "out" / name), capsys=capsys)
    assert code == 0, out["issues"]
    return out, root / "out" / name


def test_code_builds_as_editable_highlighted_text_on_a_fitted_panel(tmp_path, capsys):
    root = kit(tmp_path)
    out, pptx = build(root, SLIDE, capsys)
    slide = Presentation(str(pptx)).slides[0]
    shapes = {sh.name: sh for sh in slide.shapes}
    ph = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 1)
    lines = SLIDE.split("```python")[1].split("\n", 1)[1].rsplit("```", 1)[0].rstrip("\n").split("\n")
    assert [p.text for p in ph.text_frame.paragraphs] == lines  # plain, editable text, line for line
    assert ph._element.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}").get("descr") == "python code"
    assert ph.text_frame._txBody.find(qn("a:bodyPr")).get("wrap") == "none"
    first = ph.text_frame.paragraphs[0].runs
    tok = tokens(root)["code"]
    assert first[0].text == "def " and first[0].font.bold
    palette = BRAND["palette"]
    assert str(first[0].font.color.rgb) == palette[tok["colors"]["keyword"]]
    comment = ph.text_frame.paragraphs[3].runs[-1]
    assert comment.text.startswith("# one") and comment.font.italic
    assert first[0]._r.find(f"{qn('a:rPr')}/{qn('a:latin')}").get("typeface") == "Menlo"
    # the panel fits the code, and the band sits between it and the text
    panel, band = shapes["Code panel 1"], shapes["Code highlight 1 lines 4-5"]
    order = [sh.name for sh in slide.shapes]
    assert order.index(panel.name) < order.index(band.name) < order.index(ph.name)
    layout_ph = ph._base_placeholder
    assert panel.height == ph.height < layout_ph.height
    assert band.height == 2 * Pt(18 * 1.2)  # lines 4 and 5, at the 18 pt code size's exact line pitch
    assert [p.text for p in shapes["Code line numbers 1"].text_frame.paragraphs] == [str(n) for n in range(1, 7)]
    assert shapes["Code title 1"].text_frame.text == "forecast.py"
    fields = json.loads(pptx.with_name("deck.manifest.json").read_text())["slides"][0]["fields"]
    assert fields["code"]["language"] == "python" and fields["code"]["lines"] == 6
    assert out["issues"] == []


def test_a_code_build_is_byte_identical(tmp_path, capsys):
    root = kit(tmp_path)
    _, a = build(root, SLIDE, capsys, "a.pptx")
    _, b = build(root, SLIDE, capsys, "b.pptx")
    assert hashlib.sha256(a.read_bytes()).digest() == hashlib.sha256(b.read_bytes()).digest()


def test_a_kit_without_code_tokens_still_builds_monospace_code(ws, capsys):
    """An adopted kit maps a code field onto any text placeholder: no panel, no colors, but the code font,
    bold keywords and no wrapping."""
    path = ws / "brands" / "stock" / "tokens.yaml"
    tok = yaml.safe_load(path.read_text())
    tok["layouts"]["code"] = {"template_layout": "Title and Content", "fields": {
        "title": {"idx": 0, "kind": "text"}, "code": {"idx": 1, "kind": "code", "max_cols": 60, "max_lines": 10}}}
    path.write_text(yaml.safe_dump(tok, sort_keys=False))
    deck = write_deck(ws, "---\nbrand: stock\n---\n\n## Code\nlayout: code\n\n```sql\nSELECT 1\n```\n")
    code, out = cli_json(ws, "build", str(deck), "-o", str(ws / "out.pptx"), capsys=capsys)
    assert code == 0, out["issues"]
    slide = Presentation(str(ws / "out.pptx")).slides[0]
    assert not any(sh.name.startswith("Code panel") for sh in slide.shapes)
    ph = next(sh for sh in slide.placeholders if sh.placeholder_format.idx == 1)
    run = ph.text_frame.paragraphs[0].runs[0]
    assert run.text.startswith("SELECT") and run.font.bold and run.font.name == "Menlo"


def test_import_clamps_a_highlight_span_from_an_untrusted_file(tmp_path, capsys):
    """A shape named for a span far past the block (a hand-edited or hostile .pptx) imports as the block's
    own lines, without expanding the span first."""
    root = kit(tmp_path)
    _, pptx = build(root, SLIDE, capsys)
    prs = Presentation(str(pptx))
    band = next(sh for sh in prs.slides[0].shapes if sh.name.startswith("Code highlight"))
    band.name = "Code highlight 1 lines 5-999999"
    edited = root / "out" / "edited.pptx"
    prs.save(str(edited))
    capsys.readouterr()
    code, out = cli_json(root, "import", str(edited), str(root / "imported"), "--brand", "codes", capsys=capsys)
    assert code == 0, out
    imported = markdown.parse(root / "imported" / "deck.md")[0]
    assert imported.slides[0].fields["code"].highlight == [5, 6]


# ---------------------------------------------------------------- convert


def test_a_workbook_keeps_code_and_csv_refuses_it(tmp_path, capsys):
    root = kit(tmp_path)
    deck = write_deck(root, deck_md(SLIDE))
    capsys.readouterr()
    code, out = cli_json(root, "convert", str(deck), str(root / "decks" / "deck.xlsx"), capsys=capsys)
    assert code == 0, out["issues"]  # convert reparses its output and refuses anything that changed
    code, out = cli_json(root, "convert", str(root / "decks" / "deck.xlsx"), str(root / "decks" / "back.md"),
                         capsys=capsys)
    assert code == 0, out["issues"]
    assert markdown.parse(root / "decks" / "back.md")[0].slides == markdown.parse(deck)[0].slides
    plain = write_deck(root, "## Code\nlayout: code\n\n```python\nx = 1\n```\n", name="plain.md")
    code, out = cli_json(root, "convert", str(plain), str(root / "decks" / "deck.csv"), capsys=capsys)
    assert code == 1 and codes(out) == ["CONVERT_LOSSY"] and "is a code" in out["issues"][0]["message"]
