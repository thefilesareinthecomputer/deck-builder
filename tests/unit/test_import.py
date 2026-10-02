"""`import`: a .pptx back into deck.md, its images and a report of what needs a decision."""
import copy
import hashlib
import json
import re
from pathlib import Path

import pytest
from PIL import Image as PILImage
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from conftest import cli_json, codes, write_deck
from deck_builder import pipeline, validate
from deck_builder.brand import registry
from deck_builder.cli import main
from deck_builder.config import load as load_config
from deck_builder.errors import EnvError
from deck_builder.model import Image
from deck_builder.parse import markdown
from deck_builder.write import markdown as md_writer

DATA = Path(__file__).resolve().parents[2] / "src" / "deck_builder" / "data" / "decks"
DEMO = Path(__file__).resolve().parents[1] / "fixtures" / "demo-brands"
SLUGS = ("briarfield-paper", "cubicle-nine", "afterhours-soap")

# Front-matter keys a .pptx has no place for.
PPTX_CANT_HOLD = ("spec_version", "default_layout", "slide_level", "template", "tokens", "output")


def comparable_sha(deck, brand, deck_dir: Path) -> str:
    """content_sha256 of a deck after the one normalization an import round trip gets. Exactly three steps:

    1. Drop the front-matter keys in PPTX_CANT_HOLD; a .pptx has nowhere to keep them.
    2. Replace each deck image's path with the SHA-256 of the file it names; import names extracted
       images by slide and content hash, so only the bytes can match.
    3. Resolve the deck as `check` does, so free content under `body` gets the field it fills; deck.md
       can name the same field either way (`![..](..)` after the fields, or `### image`).

    Nothing else. Widening this to make a failing case pass defeats the test.
    """
    deck = copy.deepcopy(deck)
    for k in PPTX_CANT_HOLD:
        deck.meta.pop(k, None)
    resolved, _ = validate.resolve(deck, brand, deck_dir)
    for s in resolved.slides:
        for name, v in s.fields.items():
            if isinstance(v, Image) and not v.ref.startswith("brand:"):
                s.fields[name] = Image(ref=hashlib.sha256((deck_dir / v.ref).read_bytes()).hexdigest(), alt=v.alt)
    return pipeline.content_sha(resolved)


def run(cfg: Path, *argv: str, capsys) -> tuple[int, dict]:
    capsys.readouterr()
    code = main(["--config", str(cfg), *argv, "--json"])
    return code, json.loads(capsys.readouterr().out)


@pytest.fixture(scope="module")
def project(tmp_path_factory):
    """A workspace with the neutral brand and the three demo brands."""
    root = tmp_path_factory.mktemp("import")
    assert main(["init", "--dir", str(root)]) == 0
    for slug in SLUGS:
        assert main(["--config", str(root / "deck-builder.toml"), "brand", "init", slug, "--from",
                     str(DEMO / "brands" / slug / "brand.yaml")]) == 0
    return root


def round_trip(project, capsys, deck_path: Path, row: dict | None, slug: str, name: str) -> tuple[str, str]:
    cfg_path = project / "deck-builder.toml"
    cfg = load_config(str(cfg_path))
    brand = registry.get(cfg, slug)
    source = pipeline.load_deck(deck_path, row).deck
    built = project / "built" / f"{name}.pptx"
    if row is None:
        code, out = run(cfg_path, "build", str(deck_path), "-o", str(built), capsys=capsys)
    else:  # one row of a bulk deck
        data = project / "built" / f"{name}.csv"
        data.parent.mkdir(exist_ok=True)
        data.write_text(",".join(row) + "\n" + ",".join(f'"{v}"' for v in row.values()) + "\n")
        code, out = run(cfg_path, "build", str(deck_path), "--data", str(data), "--name", f"{name}.pptx",
                        "-o", str(built.parent), capsys=capsys)
    assert code == 0, out["issues"]
    dest = project / "imported" / name
    code, out = run(cfg_path, "import", str(built), str(dest), "--brand", slug, capsys=capsys)
    assert code == 0, out
    imported = markdown.parse(dest / "deck.md")[0]
    return comparable_sha(source, brand, deck_path.parent), comparable_sha(imported, brand, dest)


def rows(csv: Path) -> list[dict]:
    lines = csv.read_text().splitlines()
    head = lines[0].split(",")
    return [dict(zip(head, ln.split(","), strict=True)) for ln in lines[1:] if ln]


def test_round_trip_quarterly_review(project, capsys):
    a, b = round_trip(project, capsys, project / "workspace" / "decks" / "quarterly-review" / "deck.md", None,
                      "neutral", "quarterly-review")
    assert a == b


def test_round_trip_bulk_outreach(project, capsys):
    deck = project / "workspace" / "decks" / "bulk-outreach" / "deck.md"
    for n, row in enumerate(rows(deck.parent / "clients.csv")):
        a, b = round_trip(project, capsys, deck, row, "neutral", f"outreach-{n}")
        assert a == b, row


@pytest.mark.parametrize("slug", SLUGS)
def test_round_trip_showcase_in_each_demo_brand(project, capsys, slug):
    row = next(r for r in rows(DEMO / "showcase" / "brands.csv") if r["brand"] == slug)
    a, b = round_trip(project, capsys, DEMO / "showcase" / "deck.md", row, slug, f"showcase-{slug}")
    assert a == b


STRICT = """
---
brand: stock
title: Pemberton Paper quarterly review
author: Sales Operations
date: 2026-01-15
---

## Pemberton Paper quarterly review
layout: title
subtitle: Shipments, margins and next quarter

## Paper volume grew 12% in Q3
layout: content

- Copy paper led the growth, see [the ledger](https://example.com/ledger)
  - Two new *regional* accounts
- **Card stock** held flat at `SKU-12`

Notes:
Source: Q3 shipment ledger.

## 12%
layout: big-number
caption: Volume growth, Q2 to Q3

## Volume by product
layout: chart

### chart
```chart
type: column
number_format: '#,##0'
labels: true
categories: [Jul, Aug, Sep]
series:
  - name: Copy paper
    values: [410, 378, 331.5]
  - name: Card stock
    values: [120, 118, 121]
```

## Open items
layout: table

### table
| Item | Owner | Due |
|---|---|---|
| Warehouse lease | Ops | Oct |
| **Price** review | Sales | Nov |

## Two lists
layout: two-col

### left
- Before

### right
- After

## Logo
layout: image
caption: The current mark.

### image
![Logo](brand:logo/primary)

## Checked
layout: icon
icon: brand:icon/check
"""


def test_a_deck_without_deck_images_round_trips_with_no_normalization(ws, capsys):
    deck = write_deck(ws, STRICT)
    code, out = cli_json(ws, "build", str(deck), capsys=capsys)
    assert code == 0, out["issues"]
    source_sha = out["content_sha256"] if "content_sha256" in out else pipeline.content_sha(markdown.parse(deck)[0])
    code, imp = cli_json(ws, "import", out["output"], str(ws / "imported"), "--brand", "stock", capsys=capsys)
    assert code == 0, imp
    assert (imp["unplaced"], imp["dropped"], imp["images"]) == (0, 0, 0)
    imported = markdown.parse(ws / "imported" / "deck.md")[0]
    assert md_writer.write(imported) == md_writer.write(markdown.parse(deck)[0])  # shows the difference, if any
    assert pipeline.content_sha(imported) == source_sha


# ---------------------------------------------------------------- a messy deck


def messy_pptx(path: Path, image: Path) -> Path:
    """A deck made by hand in the stock template: overrides, a stray text box, a moved placeholder,
    a picture, a table, a chart, notes, and a linked image that must never be fetched."""
    prs = Presentation()
    content = prs.slide_layouts[1]  # Title and Content

    s1 = prs.slides.add_slide(content)
    s1.shapes.title.text = "Volume grew"
    run = s1.shapes.title.text_frame.paragraphs[0].runs[0]
    run.font.size = Pt(54)
    run.font.name = "Comic Sans MS"
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    body = s1.placeholders[1]
    body.left, body.top = Inches(2), Inches(3)  # moved off the layout's position
    body.text_frame.text = "Copy paper led"
    p = body.text_frame.add_paragraph()
    p.text = "Two new accounts"
    p.level = 1
    p.runs[0].font.size = Pt(9)
    box = s1.shapes.add_textbox(Inches(7), Inches(6.5), Inches(2.5), Inches(0.5))
    box.text_frame.text = "Draft: check with finance"
    s1.notes_slide.notes_text_frame.text = "Source: the ledger."

    s2 = prs.slides.add_slide(content)
    s2.shapes.title.text = "Open items"
    ph = s2.placeholders[1]
    tbl = s2.shapes.add_table(3, 2, ph.left, ph.top, ph.width, Inches(1.5)).table
    for r, (a, b) in enumerate([("Item", "Owner"), ("Lease", "Ops"), ("Prices", "Sales")]):
        tbl.cell(r, 0).text, tbl.cell(r, 1).text = a, b
    ph._element.getparent().remove(ph._element)

    s3 = prs.slides.add_slide(content)
    s3.shapes.title.text = "Cases by month"
    ph = s3.placeholders[1]
    data = CategoryChartData()
    data.categories = ["Jul", "Aug", "Sep"]
    data.add_series("Copy paper", (410, 378, 331))
    s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, ph.left, ph.top, ph.width, ph.height, data)
    ph._element.getparent().remove(ph._element)

    s4 = prs.slides.add_slide(prs.slide_layouts[8])  # Picture with Caption
    s4.shapes.title.text = "The warehouse"
    s4.placeholders[1].insert_picture(str(image))
    s4.placeholders[2].text = "Opened in August."
    linked = s4.shapes.add_picture(str(image), Inches(0.2), Inches(0.2), Inches(1), Inches(1))
    rid = s4.part.relate_to("https://example.com/tracking.png", RT.IMAGE, is_external=True)
    blip = linked._element.find(f"{qn('p:blipFill')}/{qn('a:blip')}")
    del blip.attrib[qn("r:embed")]
    blip.set(qn("r:link"), rid)
    prs.save(str(path))
    return path


@pytest.fixture
def messy(ws):
    img = ws / "photo.png"
    PILImage.new("RGB", (1600, 1000), "#2F6F4F").save(img)
    return messy_pptx(ws / "messy.pptx", img)


def test_a_messy_deck_imports_with_every_rough_edge_reported(ws, capsys, messy):
    code, out = cli_json(ws, "import", str(messy), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 0, out
    assert [x["layout"] for x in out["layouts"]] == ["content", "table", "chart", "image"]
    report = (ws / "imp" / "import-report.md").read_text()
    s1 = report.split("## Slide 1")[1].split("## Slide 2")[0]
    assert "Unplaced: Text box" in s1 and "Draft: check with finance" in s1
    for kind in ("color", "font", "position", "size"):
        assert re.search(rf"\d+ {kind}\b", s1), kind
    assert "linked image" in report and "https://example.com/tracking.png" in report and "not fetched" in report
    deck = markdown.parse(ws / "imp" / "deck.md")[0]
    first = deck.slides[0]
    assert first.title == "Volume grew"
    assert first.fields["body"] == [(0, "Copy paper led"), (1, "Two new accounts")]
    assert "Source: the ledger." in first.notes and "Draft: check with finance" in first.notes  # nothing lost
    assert deck.slides[1].fields["table"].rows == [["Lease", "Ops"], ["Prices", "Sales"]]
    assert deck.slides[2].fields["chart"].series[0].values == [410, 378, 331]
    picture = deck.slides[3].fields["image"]
    assert re.fullmatch(r"assets/slide04-[0-9a-f]{12}\.png", picture.ref)
    assert deck.slides[3].fields["caption"] == "Opened in August."


def combo_chart_pptx(path: Path) -> Path:
    """A chart with two plots (a bar plot and a line plot): import reads chart.plots[0] only, so the
    line plot's series must not vanish silently."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # Title and Content
    slide.shapes.title.text = "Cases and margin"
    ph = slide.placeholders[1]
    left, top, width, height = ph.left, ph.top, ph.width, ph.height
    ph._element.getparent().remove(ph._element)

    bar_data = CategoryChartData()
    bar_data.categories = ["Jul", "Aug", "Sep"]
    bar_data.add_series("Cases", (410, 378, 331))
    gframe = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, bar_data)
    bar_chart = gframe.chart

    # A throwaway presentation just to get a valid <c:lineChart> element to graft on as a second plot.
    helper = Presentation()
    hslide = helper.slides.add_slide(helper.slide_layouts[1])
    hph = hslide.placeholders[1]
    line_data = CategoryChartData()
    line_data.categories = ["Jul", "Aug", "Sep"]
    line_data.add_series("Margin", (12, 14, 11))
    hframe = hslide.shapes.add_chart(XL_CHART_TYPE.LINE, hph.left, hph.top, hph.width, hph.height, line_data)
    line_elem = hframe.chart.plots[0]._element

    plot_area = bar_chart.plots[0]._element.getparent()
    plot_area.append(copy.deepcopy(line_elem))
    assert len(bar_chart.plots) == 2

    prs.save(str(path))
    return path


def test_import_lists_a_combo_charts_extra_plot_as_unplaced(ws, capsys, tmp_path):
    src = combo_chart_pptx(tmp_path / "combo.pptx")
    code, out = cli_json(ws, "import", str(src), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 0, out
    assert out["unplaced"] >= 1
    deck = markdown.parse(ws / "imp" / "deck.md")[0]
    chart = deck.slides[0].fields["chart"]
    assert chart.series[0].name == "Cases" and chart.series[0].values == [410, 378, 331]  # the first plot, kept
    assert "Margin" in deck.slides[0].notes and "12" in deck.slides[0].notes  # the second plot, not dropped
    report = (ws / "imp" / "import-report.md").read_text()
    assert "Margin" in report


def test_a_deck_from_another_template_maps_onto_a_brand_by_placeholder_types(ws, capsys, messy):
    src = DEMO / "brands" / "briarfield-paper" / "brand.yaml"
    assert cli_json(ws, "brand", "init", "briarfield-paper", "--from", str(src), capsys=capsys)[0] == 0
    code, out = cli_json(ws, "import", str(messy), str(ws / "rebranded"), "--brand", "briarfield-paper",
                         capsys=capsys)
    assert code == 0, out
    assert [x["layout"] for x in out["layouts"]] == ["content", "table", "chart", "image"]
    assert all(x["match"].startswith("best match on placeholder types") for x in out["layouts"])
    assert out["unplaced"] == 1  # the stray text box goes to the notes, not into a spare field
    deck = markdown.parse(ws / "rebranded" / "deck.md")[0]
    assert "Draft: check with finance" in deck.slides[0].notes
    code, out = cli_json(ws, "check", str(ws / "rebranded"), capsys=capsys)
    assert code == 0, out["issues"]


def test_the_rebuilt_messy_deck_uses_only_template_styling(ws, capsys, messy):
    cli_json(ws, "import", str(messy), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    code, out = cli_json(ws, "build", str(ws / "imp" / "deck.md"), "-o", str(ws / "rebuilt.pptx"), capsys=capsys)
    assert code == 0, out["issues"]
    for slide in Presentation(out["output"]).slides:
        for sh in slide.placeholders:
            if not sh.has_text_frame:
                continue
            assert sh._element.find(f"{qn('p:spPr')}/{qn('a:xfrm')}") is None  # positions come from the layout
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    assert (r.font.size, r.font.name, r.font.color.type) == (None, None, None), r.text


def test_import_is_deterministic(ws, capsys, messy):
    for d in ("a", "b"):
        assert cli_json(ws, "import", str(messy), str(ws / d), "--brand", "stock", capsys=capsys)[0] == 0
    files = sorted(p.relative_to(ws / "a") for p in (ws / "a").rglob("*") if p.is_file())
    assert files == sorted(p.relative_to(ws / "b") for p in (ws / "b").rglob("*") if p.is_file())
    assert all((ws / "a" / f).read_bytes() == (ws / "b" / f).read_bytes() for f in files)


def test_import_never_replaces_a_deck_without_force(ws, capsys, messy):
    assert cli_json(ws, "import", str(messy), str(ws / "imp"), "--brand", "stock", capsys=capsys)[0] == 0
    code, out = cli_json(ws, "import", str(messy), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 2 and "pass --force" in out["error"]
    assert cli_json(ws, "import", str(messy), str(ws / "imp"), "--brand", "stock", "--force", capsys=capsys)[0] == 0


def test_import_takes_only_a_pptx(ws, capsys):
    deck = write_deck(ws, STRICT)
    code, out = cli_json(ws, "import", str(deck), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 2 and "import takes an existing .pptx file" in out["error"]


DANGEROUS_NOTES = (
    "Before the gap.\n"
    "## Appendix\n"
    "layout: title\n"
    "Not a real slide; this is notes text, not slide 2.\n"
    "Notes:\n"
    "Nested marker, still notes.\n"
    "```fence\n"
    "inside notes\n"
    "```\n"
    "![not an image](nowhere.png)\n"
    "After the gap."
)


def notes_pptx(path: Path, notes: str) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Opening"
    slide.placeholders[1].text_frame.text = "Point one"
    slide.notes_slide.notes_text_frame.text = notes
    prs.save(str(path))
    return path


def link_pptx(path: Path, address: str) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Links"
    box = slide.placeholders[1].text_frame
    box.text = "see docs"
    run = box.paragraphs[0].runs[0]
    run.hyperlink.address = address
    prs.save(str(path))
    return path


@pytest.mark.parametrize("address", ["https://example.com/docs", "mailto:a@example.com", "HTTP://example.com"])
def test_a_safe_scheme_hyperlink_is_kept_on_import(ws, capsys, tmp_path, address):
    src = link_pptx(tmp_path / "safe.pptx", address)
    code, out = cli_json(ws, "import", str(src), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 0, out
    text = Path(out["output"]).read_text()
    assert f"({address})" in text


@pytest.mark.parametrize("address", ["javascript:alert(1)", "file:///etc/passwd", "\\\\evil\\share\\a"])
def test_an_unsafe_scheme_hyperlink_becomes_plain_text_on_import(ws, capsys, tmp_path, address):
    src = link_pptx(tmp_path / "unsafe.pptx", address)
    code, out = cli_json(ws, "import", str(src), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 0, out
    text = Path(out["output"]).read_text()
    assert address not in text and "see docs" in text


def test_notes_with_heading_like_lines_round_trip_as_one_slide(ws, capsys, tmp_path):
    src = notes_pptx(tmp_path / "danger.pptx", DANGEROUS_NOTES)
    code, out = cli_json(ws, "import", str(src), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 0, out
    assert out["slides"] == 1
    imported = markdown.parse(ws / "imp" / "deck.md")[0]
    assert len(imported.slides) == 1
    assert imported.slides[0].notes == DANGEROUS_NOTES


def test_import_fails_loudly_when_the_round_trip_changes_structure(ws, capsys, messy, monkeypatch):
    import deck_builder.commands as commands_mod

    monkeypatch.setattr(commands_mod.md_writer, "write",
                        lambda deck: "---\nbrand: stock\n---\n\n## Only\nlayout: title\n")
    code, out = cli_json(ws, "import", str(messy), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 1
    assert "IMPORT_LOSSY" in codes(out)


def test_a_template_theme_cant_read_local_files_through_xml_entities(tmp_path):
    import zipfile

    from deck_builder.brand.inspect import theme

    secret = tmp_path / "secret.txt"
    secret.write_text("TOPSECRET")
    xml = (f'<?xml version="1.0"?><!DOCTYPE a:theme [<!ENTITY x SYSTEM "file://{secret}">]>'
           '<a:theme xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><a:themeElements>'
           '<a:clrScheme name="x"><a:dk1><a:srgbClr val="&x;"/></a:dk1></a:clrScheme></a:themeElements></a:theme>')
    pptx = tmp_path / "evil.pptx"
    with zipfile.ZipFile(pptx, "w") as z:
        z.writestr("ppt/theme/theme1.xml", xml)
    with pytest.raises(EnvError, match="isn't plain XML"):
        theme(pptx)


def test_adopt_makes_a_kit_from_the_decks_own_layouts(ws, capsys, messy):
    code, out = cli_json(ws, "import", str(messy), str(ws / "imp"), "--adopt", "messy-co", capsys=capsys)
    assert code == 0, out
    kit = ws / "brands" / "messy-co"
    assert len(Presentation(str(kit / "template.pptx")).slides) == 0  # masters and layouts only
    assert all(x["match"].startswith("matched by layout name") for x in out["layouts"])
    assert (ws / "imp" / "deck.md").read_text().startswith("---\nbrand: messy-co\n")


def test_import_adopt_refuses_to_regenerate_an_existing_kit(ws, capsys, messy):
    code, out = cli_json(ws, "import", str(messy), str(ws / "imp"), "--adopt", "messy-co", capsys=capsys)
    assert code == 0, out
    tokens_path = ws / "brands" / "messy-co" / "tokens.yaml"
    tuned = tokens_path.read_text() + "\n# tuned by hand after adopt\n"
    tokens_path.write_text(tuned, encoding="utf-8")
    code, out = cli_json(ws, "import", str(messy), str(ws / "imp2"), "--adopt", "messy-co", "--force", capsys=capsys)
    assert code == 2
    assert "--brand" in out["error"]
    assert tokens_path.read_text() == tuned  # the tuned kit was never touched


def test_import_force_with_an_existing_brand_never_touches_the_kit(ws, capsys, messy):
    code, out = cli_json(ws, "import", str(messy), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 0, out
    kit = ws / "brands" / "stock"
    before = {p: p.read_bytes() for p in kit.rglob("*") if p.is_file()}
    code, out = cli_json(ws, "import", str(messy), str(ws / "imp"), "--brand", "stock", "--force", capsys=capsys)
    assert code == 0, out
    after = {p: p.read_bytes() for p in kit.rglob("*") if p.is_file()}
    assert before == after


def test_import_refuses_to_overwrite_an_existing_report_or_assets_without_force(ws, capsys, messy):
    code, out = cli_json(ws, "import", str(messy), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 0, out
    report_before = (ws / "imp" / "import-report.md").read_text()
    asset = next((ws / "imp" / "assets").iterdir())
    asset_before = asset.read_bytes()
    (ws / "imp" / "deck.md").unlink()  # deck.md gone, but the report and assets remain from the first run
    code, out = cli_json(ws, "import", str(messy), str(ws / "imp"), "--brand", "stock", capsys=capsys)
    assert code == 2
    assert "--force" in out["error"]
    assert (ws / "imp" / "import-report.md").read_text() == report_before
    assert asset.read_bytes() == asset_before
