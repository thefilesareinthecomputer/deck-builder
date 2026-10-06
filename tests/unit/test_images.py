"""Images a slide names and where they land: image cues in the notes, IMAGE_NO_SLOT, IMAGE_CROPPED, `fit:`."""
from pathlib import Path

import pytest
import yaml
from PIL import Image
from pptx import Presentation

from conftest import cli_json, codes, write_deck
from deck_builder import images
from deck_builder.model import Image as ImageValue
from deck_builder.model import Slide


def _png(ws: Path, rel: str, size: tuple[int, int], mode: str = "RGB") -> None:
    p = ws / "decks" / rel
    p.parent.mkdir(parents=True, exist_ok=True)
    Image.new(mode, size, "#556677" if mode == "RGB" else (85, 102, 119, 0)).save(p)


def _issue(out: dict, code: str) -> dict:
    hits = [i for i in out["issues"] if i["code"] == code]
    assert hits, out["issues"]
    return hits[0]


def test_cues_read_the_path_when_the_first_part_names_an_image():
    found = images.cues("Source: the ledger.\nSCREENSHOT: assets/screenshots/run.png | shows: the run page\n"
                        "> diagram: the flow from intake to invoice\n- Screenshot: `assets/b.jpg`\nNo cue here.")
    assert [(c.kind, c.path) for c in found] == [("screenshot", "assets/screenshots/run.png"), ("diagram", None),
                                                 ("screenshot", "assets/b.jpg")]


def test_a_cue_on_a_layout_with_no_image_slot_warns_with_the_layouts_and_what_moving_costs(ws, capsys):
    p = write_deck(ws, """
        ---
        brand: stock
        ---
        ## Runs finish in four minutes
        layout: content

        - The scheduler starts each run at six
        - Retries cover the two flaky sources
        - Alerts go to the on-call channel

        Notes:
        SCREENSHOT: assets/screenshots/run.png | shows: the run page | capture: after a full run
        """)
    code, out = cli_json(ws, "check", str(p), capsys=capsys)
    assert code == 0  # a warning: the deck still builds
    hit = _issue(out, "IMAGE_NO_SLOT")
    assert hit["severity"] == "warning" and hit["slide"] == 1 and hit["actual"] == 1 and hit["limit"] == 0
    # the stock kit's only image layout is `image`: one caption of 120 characters takes the three bullets' place
    assert "layout 'content' has no image slot" in hit["message"]
    assert "image (image): cut 2 of 3 items" in hit["message"]


def test_no_warning_when_the_layout_shows_the_image_the_notes_name(ws, capsys):
    _png(ws, "assets/screenshots/run.png", (1200, 900))
    p = write_deck(ws, """
        ---
        brand: stock
        ---
        ## Runs finish in four minutes
        layout: image

        ![The run page](assets/screenshots/run.png)

        Notes:
        SCREENSHOT: assets/screenshots/run.png | shows: the run page
        """)
    _, out = cli_json(ws, "check", str(p), capsys=capsys)
    assert "IMAGE_NO_SLOT" not in codes(out), out["issues"]
    p.write_text(p.read_text(encoding="utf-8").replace("SCREENSHOT: assets/screenshots/run.png |",
                                                       "SCREENSHOT:"), encoding="utf-8")
    _, out = cli_json(ws, "check", str(p), capsys=capsys)  # a cue with no path describes the image shown
    assert "IMAGE_NO_SLOT" not in codes(out), out["issues"]


def test_more_images_than_any_layout_holds_says_to_split_the_slide(ws, capsys):
    _png(ws, "assets/a.png", (1200, 900))
    p = write_deck(ws, """
        ---
        brand: stock
        ---
        ## Two screens side by side
        layout: image

        ![The run page](assets/a.png)

        Notes:
        SCREENSHOT: assets/a.png
        SCREENSHOT: assets/b.png | shows: the log
        """)
    _, out = cli_json(ws, "check", str(p), capsys=capsys)
    hit = _issue(out, "IMAGE_NO_SLOT")
    assert "name 2 images" in hit["message"] and "holds 1" in hit["message"]
    assert "split the slide across image (1)" in hit["message"]


def test_an_image_in_a_field_the_layout_lacks_names_the_image_layouts(ws, capsys):
    _png(ws, "assets/a.png", (1200, 900))
    p = write_deck(ws, "---\nbrand: stock\n---\n## A\nlayout: content\n\n### photo\n![A](assets/a.png)\n")
    _, out = cli_json(ws, "check", str(p), capsys=capsys)
    assert "layouts with an image slot: image" in _issue(out, "UNKNOWN_FIELD")["message"]
    p = write_deck(ws, "---\nbrand: stock\n---\n## A\nlayout: two-col\n\n![A](assets/a.png)\n")
    _, out = cli_json(ws, "check", str(p), capsys=capsys)
    assert "layouts with an image slot: image" in _issue(out, "KIND_MISMATCH")["message"]


@pytest.mark.parametrize("rel,fit,warned", [
    ("assets/tall.png", None, True),  # a portrait image in the stock 4:3 box loses 44% of its height
    ("assets/tall.png", "cover", False),  # the crop is chosen
    ("assets/tall.png", "contain", False),  # shown whole
    ("assets/screenshots/tall.png", None, False),  # a screenshot is contained by default
    ("assets/wide.png", None, False),  # 16:9 in a 4:3 box loses 25% of its width: only its sides
    ("assets/four-three.png", None, False),
])
def test_image_cropped_reports_what_a_crop_to_fill_takes(ws, capsys, rel, fit, warned):
    size = {"tall": (900, 1200), "wide": (1600, 900), "four-three": (1200, 900)}[Path(rel).stem]
    _png(ws, rel, size)
    setting = f"fit: {fit}\n" if fit else ""
    p = write_deck(ws, f"---\nbrand: stock\n---\n## A\nlayout: image\n{setting}\n![A]({rel})\n")
    _, out = cli_json(ws, "check", str(p), capsys=capsys)
    assert ("IMAGE_CROPPED" in codes(out)) is warned, out["issues"]
    if warned and rel.endswith("tall.png"):
        hit = _issue(out, "IMAGE_CROPPED")
        assert hit["actual"] == 44 and "of its height, from the top and bottom" in hit["message"]


def test_lint_max_crop_sets_the_threshold(ws, capsys):
    brand = ws / "brands" / "stock" / "brand.yaml"
    meta = yaml.safe_load(brand.read_text(encoding="utf-8"))
    meta["lint"]["max_crop"] = 50
    brand.write_text(yaml.safe_dump(meta, sort_keys=False), encoding="utf-8")
    _png(ws, "assets/tall.png", (900, 1200))  # loses 44% of its height
    p = write_deck(ws, "---\nbrand: stock\n---\n## A\nlayout: image\n\n![A](assets/tall.png)\n")
    _, out = cli_json(ws, "check", str(p), capsys=capsys)
    assert "IMAGE_CROPPED" not in codes(out), out["issues"]


def test_fit_takes_only_contain_or_cover(ws, capsys):
    p = write_deck(ws, "---\nbrand: stock\n---\n## A\nlayout: title\nfit: stretch\n")
    code, out = cli_json(ws, "check", str(p), capsys=capsys)
    assert code == 1 and "fit: 'stretch' must be contain" in _issue(out, "PARSE")["message"]


def test_a_contained_screenshot_is_shown_whole_and_a_photo_fills_its_box(ws, capsys):
    _png(ws, "assets/screenshots/tall.png", (900, 1200))
    _png(ws, "assets/photo.png", (1600, 900))
    p = write_deck(ws, """
        ---
        brand: stock
        ---
        ## The run page
        layout: image

        ![The run page](assets/screenshots/tall.png)

        ## The warehouse
        layout: image
        fit: cover

        ![The warehouse](assets/photo.png)
        """)
    out_pptx = ws / "out" / "deck.pptx"
    code, out = cli_json(ws, "build", str(p), "-o", str(out_pptx), capsys=capsys)
    assert code == 0, out
    manifest = yaml.safe_load((ws / "out" / "deck.manifest.json").read_text(encoding="utf-8"))
    assert manifest["slides"][0]["fields"]["image"]["fit"] == "contain"
    assert "fit" not in manifest["slides"][1]["fields"]["image"]
    shot, photo = (next(sh for sh in s.shapes if sh.shape_type == 13 or "Picture" in sh.name)
                   for s in Presentation(str(out_pptx)).slides)
    assert (shot.crop_top, shot.crop_bottom) == (0, 0)
    assert shot.height <= 4.5 * 914400 + 1 and shot.width < shot.height  # inside the 6 x 4.5 in box, whole
    assert photo.crop_left > 0.1  # 16:9 cropped to fill a 4:3 box


def test_import_keeps_a_screenshot_shown_whole(ws, capsys):
    """The imported image sits in assets/, outside its screenshots/ folder, so the slide says `fit: contain`."""
    _png(ws, "assets/screenshots/tall.png", (900, 1200))
    p = write_deck(ws, "---\nbrand: stock\n---\n## The run page\nlayout: image\n\n![The run page]"
                       "(assets/screenshots/tall.png)\n")
    code, out = cli_json(ws, "build", str(p), "-o", str(ws / "out" / "shot.pptx"), capsys=capsys)
    assert code == 0, out
    code, out = cli_json(ws, "import", str(ws / "out" / "shot.pptx"), str(ws / "decks" / "back"), "--brand", "stock",
                         capsys=capsys)
    assert code == 0, out
    assert "fit: contain" in (ws / "decks" / "back" / "deck.md").read_text(encoding="utf-8")


def test_fit_round_trips_through_the_workbook_and_csv(ws, capsys):
    _png(ws, "assets/photo.png", (1200, 900))
    p = write_deck(ws, "## A\nlayout: image\nfit: contain\nimage: assets/photo.png\n")  # CSV holds no front matter
    for ext in ("xlsx", "csv"):
        code, out = cli_json(ws, "convert", str(p), str(ws / "decks" / f"deck.{ext}"), capsys=capsys)
        assert code == 0, out["issues"]
        back = ws / "decks" / f"back-{ext}.md"
        code, out = cli_json(ws, "convert", str(ws / "decks" / f"deck.{ext}"), str(back), capsys=capsys)
        assert code == 0, out["issues"]
        assert "fit: contain" in back.read_text(encoding="utf-8")


def test_move_cost_keeps_the_shared_fields_and_counts_the_rest():
    s = Slide(title="Runs finish in four minutes", layout="process-3", fields={
        "title": "Runs finish in four minutes", "kicker": "Operations", "takeaway": "Nothing waits on a person",
        "step1": "Start", "text1": "The scheduler starts each run at six every morning",
        "step2": "Retry", "text2": "Retries cover the two flaky sources", "step3": "Alert", "text3": "On call",
        "shot": ImageValue(ref="assets/a.png")})
    target = {"fields": {"title": {"kind": "text"}, "kicker": {"kind": "text"}, "image": {"kind": "image"},
                         "body": {"kind": "bullets", "max_bullets": 4, "max_chars": 120, "max_bullet_chars": 40}}}
    said = images.move_cost(s, target)
    assert said == "cut 2 of 6 items and shorten 1 item to 40 characters; drops the takeaway"
    tight = {"fields": {**target["fields"], "body": {**target["fields"]["body"], "max_chars": 100}}}
    assert images.move_cost(s, tight).startswith("cut 2 of 6 items and cut 7 of 107 characters and")
    roomy = {"fields": {**target["fields"], "takeaway": {"kind": "text"},
                        "body": {"kind": "bullets", "max_bullets": 6, "max_chars": 300, "max_bullet_chars": 60}}}
    assert images.move_cost(s, roomy) == "fits as is"


def test_the_image_report_says_where_each_image_lands(ws, capsys):
    _png(ws, "assets/screenshots/run.png", (900, 1200))
    _png(ws, "assets/photo.png", (1600, 900))
    _png(ws, "assets/log.png", (1200, 900))
    _png(ws, "assets/spare.png", (1200, 900))
    p = write_deck(ws, """
        ---
        brand: stock
        ---
        ## The run page
        layout: image

        ![The run page](assets/screenshots/run.png)

        ## The warehouse
        layout: image

        ![The warehouse](assets/photo.png)

        ## Runs retry twice
        layout: content

        - Retries cover the flaky sources

        Notes:
        SCREENSHOT: assets/log.png | shows: the retry log
        DIAGRAM: assets/flow.png | shows: the flow, not drawn yet
        """)
    code, out = cli_json(ws, "assets", str(p), "--images", capsys=capsys)
    assert code == 0, out
    rows = [(r["slide"], r["where"], r["path"], r["status"], r.get("cropped"), r.get("missing")) for r in out["images"]]
    assert rows == [(1, "image", "assets/screenshots/run.png", "shown", None, None),
                    (2, "image", "assets/photo.png", "cropped", 25, None),
                    (3, "notes", "assets/log.png", "notes only", None, None),
                    (3, "notes", "assets/flow.png", "notes only", None, True),
                    (None, None, "assets/spare.png", "unused", None, None)]


def test_the_image_report_needs_a_deck(ws, capsys):
    code, out = cli_json(ws, "assets", "stock", "--images", capsys=capsys)
    assert code == 2 and "--images takes a deck" in out["error"]


def test_crop_loss_names_the_sides():
    assert images.crop_loss((1600, 900), (1200, 900)) == pytest.approx((0.25, "left and right"))
    assert images.crop_loss((900, 1200), (1200, 900)) == pytest.approx((0.4375, "top and bottom"))
    assert images.crop_loss((1200, 900), (1200, 900))[0] == 0
