"""Status dots: a body cell matching a tokens.yaml table.status key gets a colored "● " run."""
import time
from pathlib import Path

from pptx.dml.color import RGBColor
from test_visual_output import build, frames, set_table_tokens

from conftest import cli_json, write_deck
from deck_builder.parse import markdown

STATUS = {"Green": "15803D", "Amber": "B45309", "Red": "B91C1C"}

STATUS_DECK = """
## Open items
layout: table

| Item | Status | Due |
|---|---|---|
| Warehouse lease | Green | Oct |
| Card stock order | Purple | Nov |
"""

HEADER_MATCH_DECK = """
## Open items
layout: table

| Green | Owner | Due |
|---|---|---|
| Warehouse lease | Ops | Oct |
"""

WHITESPACE_DECK = """
## Open items
layout: table

### table
```table
header: [Item, Status, Due]
rows:
  - [Warehouse lease, "  green ", Oct]
```
"""


def cell_runs(prs):
    """The table's cells as lists of (text, color-or-None) tuples, per run."""
    table = next(frames(prs)).table
    out = []
    for row in table.rows:
        row_out = []
        for cell in row.cells:
            runs = cell.text_frame.paragraphs[0].runs
            row_out.append([(r.text, r.font.color.rgb if r.font.color.type else None) for r in runs])
        out.append(row_out)
    return table, out


def test_a_matching_cell_gets_a_dot_run_in_its_status_color(ws, capsys):
    set_table_tokens(ws, status=STATUS, text="ink")
    prs = build(ws, capsys, STATUS_DECK)
    table, rows = cell_runs(prs)
    dot, word = rows[1][1]  # row 1 (Warehouse lease), column 1 (Status) = "Green"
    assert dot == ("● ", RGBColor.from_string(STATUS["Green"]))
    assert word == ("Green", RGBColor.from_string("1B1B1B"))  # stock palette ink


def test_an_unmatched_value_is_unchanged(ws, capsys):
    set_table_tokens(ws, status=STATUS, text="ink")
    prs = build(ws, capsys, STATUS_DECK)
    _, rows = cell_runs(prs)
    assert rows[2][1] == [("Purple", RGBColor.from_string("1B1B1B"))]  # no dot, no error


def test_matching_is_case_and_whitespace_insensitive(ws, capsys):
    set_table_tokens(ws, status=STATUS, text="ink")
    prs = build(ws, capsys, WHITESPACE_DECK)
    _, rows = cell_runs(prs)
    dot, word = rows[1][1]
    assert dot[0] == "● " and dot[1] == RGBColor.from_string(STATUS["Green"])
    assert word[0] == "  green "  # the word itself is kept verbatim


def test_the_header_row_never_gets_a_dot_even_if_its_text_matches_a_status_key(ws, capsys):
    set_table_tokens(ws, status=STATUS, text="ink")
    prs = build(ws, capsys, HEADER_MATCH_DECK)
    _, rows = cell_runs(prs)
    # one run, header_text color (stock palette "background"), never the status color or a dot
    assert rows[0][0] == [("Green", RGBColor.from_string("FFFFFF"))]


def test_the_dot_run_matches_the_body_runs_size_and_the_cell_fill_is_unchanged(ws, capsys):
    set_table_tokens(ws, status=STATUS, text="ink", band_fill="F2F2F2", row_fill="FFFFFF")
    prs = build(ws, capsys, STATUS_DECK)
    table = next(frames(prs)).table
    status_cell = table.cell(1, 1)
    runs = status_cell.text_frame.paragraphs[0].runs
    assert len(runs) == 2
    assert runs[0].font.size == runs[1].font.size
    # table row index 1 (the first data row) takes row_fill, same as every other cell in that row
    title_cell = table.cell(1, 0)
    assert status_cell.fill.fore_color.rgb == title_cell.fill.fore_color.rgb == RGBColor.from_string("FFFFFF")


def test_two_builds_with_status_tokens_are_byte_identical(ws, capsys):
    set_table_tokens(ws, status=STATUS)
    deck = write_deck(ws, STATUS_DECK)
    code, out1 = cli_json(ws, "build", str(deck), "-o", str(ws / "one.pptx"), capsys=capsys)
    assert code == 0, out1["issues"]
    time.sleep(2.1)  # past the zip timestamp resolution, so a clock leak would show
    code, out2 = cli_json(ws, "build", str(deck), "-o", str(ws / "two.pptx"), capsys=capsys)
    assert code == 0, out2["issues"]
    assert Path(out1["output"]).read_bytes() == Path(out2["output"]).read_bytes()


def test_import_of_a_built_status_table_gives_back_the_decks_own_cell_text(ws, capsys):
    set_table_tokens(ws, status=STATUS)
    deck = write_deck(ws, STATUS_DECK)
    code, out = cli_json(ws, "build", str(deck), capsys=capsys)
    assert code == 0, out["issues"]
    code, imp = cli_json(ws, "import", out["output"], str(ws / "imported"), "--brand", "stock", capsys=capsys)
    assert code == 0, imp
    imported = markdown.parse(ws / "imported" / "deck.md")[0]
    table = imported.slides[0].fields["table"]
    assert table.rows == [["Warehouse lease", "Green", "Oct"], ["Card stock order", "Purple", "Nov"]]


# ---------------------------------------------------------------- contrast and color_refs


def low_contrast(out):
    return [i for i in out["issues"] if i["code"] == "LOW_CONTRAST"]


def test_the_default_status_set_passes_contrast_against_every_table_fill(ws, capsys):
    set_table_tokens(ws, status=STATUS, band_fill="F2F2F2", row_fill="FFFFFF")
    code, out = cli_json(ws, "brand", "check", "stock", capsys=capsys)
    assert code == 0, out["issues"]
    assert [i for i in low_contrast(out) if i["message"].startswith("table status")] == []


def test_a_light_status_color_warns_low_contrast_on_the_background(ws, capsys):
    set_table_tokens(ws, status={"Amber": "FFD966"})
    code, out = cli_json(ws, "brand", "check", "stock", capsys=capsys)
    assert code == 0  # a warning, not an error
    issue = next(i for i in low_contrast(out) if i["message"].startswith("table status Amber"))
    assert issue["message"] == "table status Amber on background (#FFD966 on #FFFFFF) is 1.37:1; needs 3.0:1"
    assert issue["limit"] == 3.0


def test_a_status_color_that_only_fails_against_band_fill_is_labeled_band_fill(ws, capsys):
    set_table_tokens(ws, status={"Red": "B91C1C"}, band_fill="B91C1C")
    code, out = cli_json(ws, "brand", "check", "stock", capsys=capsys)
    assert code == 0
    issues = [i for i in low_contrast(out) if i["message"].startswith("table status")]
    assert len(issues) == 1
    assert issues[0]["message"].startswith("table status Red on band_fill")


def test_an_unknown_status_palette_name_is_unknown_asset(ws, capsys):
    set_table_tokens(ws, status={"Green": "not-a-color"})
    code, out = cli_json(ws, "brand", "check", "stock", capsys=capsys)
    assert code == 1
    assert any(i["code"] == "UNKNOWN_ASSET" and "table.status.Green" in i["message"] for i in out["issues"])
