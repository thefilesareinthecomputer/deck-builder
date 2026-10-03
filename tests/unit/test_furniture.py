"""Slide numbers and footers: generated kits, built slides, adopted templates. Restraint is the bar."""
import shutil
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Emu
from test_check import GOOD

from conftest import cli_json, write_deck
from deck_builder import template as tpl

DEMO = Path(__file__).resolve().parents[1] / "fixtures" / "demo-brands" / "brands"
EMU = 914400
HIDES_MASTER = {"Title", "Section", "Closing", "Image Full"}  # image-full: nothing over the photo


def init_brand(ws, capsys, slug="dumbder-nifftlin", **generate):
    src = ws / "src" / slug
    shutil.copytree(DEMO / slug, src, dirs_exist_ok=True)  # assets resolve beside the brand.yaml
    if generate:
        meta = yaml.safe_load((src / "brand.yaml").read_text())
        meta["generate"] = {**(meta.get("generate") or {}), **generate}
        (src / "brand.yaml").write_text(yaml.safe_dump(meta, sort_keys=False))
    code, out = cli_json(ws, "brand", "init", slug, "--from", str(src / "brand.yaml"), capsys=capsys)
    assert code == 0, out["issues"]
    return ws / "brands" / slug


def furniture(shapes) -> dict[str, object]:
    return {ph.placeholder_format.type.name: ph for ph in shapes
            if ph.placeholder_format.type in (PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE)}


# ---------------------------------------------------------------- generated kits


def test_the_master_has_a_quiet_number_and_footer_and_no_date(ws, capsys):
    kit = init_brand(ws, capsys)
    prs = tpl.open_template(kit / "template.potx")
    master = furniture(prs.slide_masters[0].placeholders)
    assert set(master) == {"SLIDE_NUMBER", "FOOTER"}
    num, ftr = master["SLIDE_NUMBER"], master["FOOTER"]
    h = prs.slide_height
    assert num.left == round(0.6 * EMU)  # left-aligned at the margin
    assert abs((num.top + num.height / 2) - (h - Emu(round(0.46 * EMU)))) < EMU * 0.01  # on the logo's line
    assert ftr.top == num.top and ftr.left == num.left + num.width  # same baseline, after the number
    for ph in (num, ftr):
        lvl = ph.text_frame._txBody.find(".//{*}lstStyle/{*}lvl1pPr")
        assert lvl.get("algn") == "l" and lvl.find("{*}buNone") is not None
        rpr = lvl.find("{*}defRPr")
        assert (rpr.get("sz"), rpr.get("b")) == ("1000", "0")
        assert rpr.find("{*}solidFill/{*}srgbClr").get("val") == "6A7971"  # dumbder-nifftlin's muted, 4.58:1
        assert ph._element.find("{*}spPr/{*}ln") is None and ph._element.find("{*}spPr/{*}solidFill") is None


def test_only_layouts_that_show_the_master_get_furniture(ws, capsys):
    prs = tpl.open_template(init_brand(ws, capsys) / "template.potx")
    for layout in prs.slide_masters[0].slide_layouts:
        got = set(furniture(layout.placeholders))
        assert got == (set() if layout.name in HIDES_MASTER else {"SLIDE_NUMBER", "FOOTER"}), layout.name


def test_tokens_record_the_furniture_color_choice(ws, capsys):
    tokens = yaml.safe_load((init_brand(ws, capsys) / "tokens.yaml").read_text())
    assert tokens["furniture"] == {"slide_numbers": True, "color": "muted"}


def test_a_muted_color_under_4_5_to_1_falls_back_to_ink(ws, capsys):
    kit = init_brand(ws, capsys, slug="soap-club")  # muted 8D8180 is 3.76:1 on white
    assert yaml.safe_load((kit / "tokens.yaml").read_text())["furniture"]["color"] == "ink"
    num = furniture(tpl.open_template(kit / "template.potx").slide_masters[0].placeholders)["SLIDE_NUMBER"]
    assert num.text_frame._txBody.find(".//{*}defRPr/{*}solidFill/{*}srgbClr").get("val") == "211F24"


def test_the_brand_can_turn_slide_numbers_off(ws, capsys):
    kit = init_brand(ws, capsys, slide_numbers=False)
    prs = tpl.open_template(kit / "template.potx")
    master = furniture(prs.slide_masters[0].placeholders)
    assert set(master) == {"FOOTER"} and master["FOOTER"].left == round(0.6 * EMU)
    assert all("SLIDE_NUMBER" not in furniture(lay.placeholders) for lay in prs.slide_masters[0].slide_layouts)
    assert yaml.safe_load((kit / "tokens.yaml").read_text())["furniture"]["slide_numbers"] is False


# ---------------------------------------------------------------- built decks

DECK = """
---
brand: dumbder-nifftlin
title: Furniture test
{extra}---

## Q3 review
layout: title
subtitle: Volume and delivery

## Volume grew
layout: content

- Copy paper led

## Part 2
layout: section
kicker: Part 2

## 18%
layout: big-number
caption: Growth in volume

## Thank you
layout: closing
subtitle: Questions to ops@example.com
"""


def build(ws, capsys, extra="", out="deck.pptx"):
    deck = write_deck(ws, DECK.format(extra=extra))
    code, res = cli_json(ws, "build", str(deck), "-o", str(ws / "out" / out), capsys=capsys)
    assert code == 0, res["issues"]
    return Path(res["output"])


def slide_furniture(pptx: Path) -> list[dict[str, str]]:
    """Per slide: furniture type name -> its text."""
    return [{k: ph.text_frame.text for k, ph in furniture(s.placeholders).items()}
            for s in Presentation(str(pptx)).slides]


def test_content_slides_get_their_number_and_title_section_closing_get_none(ws, capsys):
    init_brand(ws, capsys)
    got = slide_furniture(build(ws, capsys))
    assert got == [{}, {"SLIDE_NUMBER": "2"}, {}, {"SLIDE_NUMBER": "4"}, {}]


def test_the_number_keeps_the_layouts_position_and_style(ws, capsys):
    init_brand(ws, capsys)
    slide = Presentation(str(build(ws, capsys))).slides[1]
    num = furniture(slide.placeholders)["SLIDE_NUMBER"]._element
    src = furniture(slide.slide_layout.placeholders)["SLIDE_NUMBER"]._element
    for part in ("{*}spPr", "{*}txBody/{*}bodyPr", "{*}txBody/{*}lstStyle"):
        shape = [(e.tag, dict(e.attrib)) for e in num.find(part).iter()]
        assert shape == [(e.tag, dict(e.attrib)) for e in src.find(part).iter()], part
    assert num.find(".//{*}fld").get("type") == "slidenum"


def test_a_footer_appears_only_when_the_deck_sets_one(ws, capsys):
    init_brand(ws, capsys)
    got = slide_furniture(build(ws, capsys, "footer: Confidential\n"))
    assert got == [{}, {"SLIDE_NUMBER": "2", "FOOTER": "Confidential"}, {},
                   {"SLIDE_NUMBER": "4", "FOOTER": "Confidential"}, {}]


def test_a_deck_can_turn_numbers_off(ws, capsys):
    init_brand(ws, capsys)
    assert slide_furniture(build(ws, capsys, "slide_numbers: false\n")) == [{}] * 5


def test_an_adopted_templates_own_slide_number_is_copied_with_its_field(ws, capsys):
    deck = write_deck(ws, GOOD)  # the stock kit wraps python-pptx's template, whose layouts have sldNum
    code, res = cli_json(ws, "build", str(deck), capsys=capsys)
    assert code == 0, res["issues"]
    prs = Presentation(res["output"])
    slide = prs.slides[1]
    num = furniture(slide.placeholders)["SLIDE_NUMBER"]
    layout_fld = furniture(slide.slide_layout.placeholders)["SLIDE_NUMBER"]._element.find(".//{*}fld")
    assert num.text_frame.text == "2"
    assert num._element.find(".//{*}fld").get("id") == layout_fld.get("id")
    assert "DATE" not in furniture(slide.placeholders)


def test_builds_with_furniture_stay_byte_identical(ws, capsys):
    init_brand(ws, capsys)
    a = build(ws, capsys, "footer: Confidential\n", "a.pptx").read_bytes()
    assert build(ws, capsys, "footer: Confidential\n", "b.pptx").read_bytes() == a


def test_furniture_isnt_an_empty_placeholder(ws, capsys):
    from deck_builder.qa.render import empty_placeholders

    init_brand(ws, capsys)
    assert empty_placeholders(build(ws, capsys, "footer: Confidential\n")) == []


def test_import_gives_back_the_footer_and_the_off_switch(ws, capsys):
    from deck_builder.parse import markdown

    init_brand(ws, capsys)
    for n, (extra, meta) in enumerate([("footer: Confidential\n", {"footer": "Confidential"}),
                                       ("slide_numbers: false\n", {"slide_numbers": False})]):
        pptx = build(ws, capsys, extra)
        dest = ws / "imp" / str(n)
        code, res = cli_json(ws, "import", str(pptx), str(dest), "--brand", "dumbder-nifftlin", capsys=capsys)
        assert code == 0, res
        got = markdown.parse(dest / "deck.md")[0].meta
        assert {k: got.get(k) for k in meta} == meta


# ---------------------------------------------------------------- reporting


def test_brand_show_says_which_layouts_show_numbers_and_footers(ws, capsys):
    init_brand(ws, capsys)
    code, out = cli_json(ws, "brand", "show", "dumbder-nifftlin", capsys=capsys)
    assert code == 0
    for kind in ("slide_numbers", "footer"):
        assert "content" in out["furniture"][kind]
        assert not {"title", "section", "closing"} & set(out["furniture"][kind])


@pytest.mark.parametrize("extra, numbers, footer", [
    ("", True, None), ("footer: Confidential\n", True, "Confidential"), ("slide_numbers: false\n", False, None)])
def test_check_reports_whether_numbers_and_footer_are_on(ws, capsys, extra, numbers, footer):
    init_brand(ws, capsys)
    deck = write_deck(ws, DECK.format(extra=extra))
    for command in ("check", "build"):
        code, out = cli_json(ws, command, str(deck), capsys=capsys)
        assert code == 0, out["issues"]
        assert (out["slide_numbers"], out["footer"]) == (numbers, footer)


@pytest.mark.parametrize("slug", ["dumbder-nifftlin"])
def test_generation_stays_deterministic(ws, capsys, slug):
    first = (init_brand(ws, capsys, slug) / "template.potx").read_bytes()
    code, out = cli_json(ws, "brand", "init", slug, "--from", str(ws / "src" / slug / "brand.yaml"), "--force",
                         capsys=capsys)
    assert code == 0
    assert (ws / "brands" / slug / "template.potx").read_bytes() == first
