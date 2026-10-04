import hashlib
from pathlib import Path

import pytest
import yaml
from PIL import Image
from pptx import Presentation

from conftest import cli_json, write_deck
from deck_builder.template import open_template

BRAND = {
    "spec_version": 1,
    "name": "Pemberton Paper Co.",
    "slug": "pemberton",
    "version": "1.0.0",
    "palette": {"primary": "1F3A5F", "accent": "E07A2F", "ink": "1B1B1B", "muted": "6B7280",
                "surface": "F4F5F7", "background": "FFFFFF"},
    "fonts": {"heading": {"family": "Georgia", "fallback": "Times New Roman"},
              "body": {"family": "Arial", "fallback": "Liberation Sans"}},
    "logos": {"primary": "assets/logo.png"},
    "icons": {"dir": "assets/icons", "default_color": "accent", "source": "made for tests"},
    "lint": {"max_slides": 40, "banned_patterns": []},
    "generate": {"slide_size": "16:9", "layout_set": "standard", "logo_on_master": "primary"},
}


def source(tmp_path, **gen):
    src = tmp_path / "src"
    (src / "assets" / "icons").mkdir(parents=True)
    Image.new("RGB", (600, 200), "#1F3A5F").save(src / "assets" / "logo.png")
    Image.new("RGBA", (256, 256), (0, 0, 0, 255)).save(src / "assets" / "icons" / "truck.png")
    meta = {**BRAND, "generate": {**BRAND["generate"], **gen}}
    (src / "brand.yaml").write_text(yaml.safe_dump(meta, sort_keys=False))
    return src / "brand.yaml"


@pytest.mark.parametrize("layout_set", ["minimal", "standard", "full"])
def test_generated_kit_passes_brand_check(ws, capsys, tmp_path, layout_set):
    src = source(tmp_path, layout_set=layout_set)
    code, out = cli_json(ws, "brand", "init", "pemberton", "--from", str(src), capsys=capsys)
    assert code == 0, out["issues"]
    code, out = cli_json(ws, "brand", "check", "pemberton", capsys=capsys)
    assert code == 0, out["issues"]


def test_template_is_16_9_with_named_layouts_theme_and_logo(ws, capsys, tmp_path):
    cli_json(ws, "brand", "init", "pemberton", "--from", str(source(tmp_path)), capsys=capsys)
    kit = ws / "brands" / "pemberton"
    prs = open_template(kit / "template.potx")
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)
    names = [lay.name for lay in prs.slide_layouts]
    assert names == ["Title", "Section", "Content", "Two Column", "Big Number", "Chart", "Table", "Image",
                     "Quote", "Closing"]
    from deck_builder.brand.inspect import theme

    th = theme(kit / "template.potx")
    assert th["colors"]["dk2"] == "1F3A5F" and th["colors"]["accent2"] == "E07A2F"
    assert th["fonts"] == {"heading": "Georgia", "body": "Arial"}
    master_pics = [s for s in prs.slide_masters[0].shapes if s.shape_type == 13]
    assert len(master_pics) == 1
    assert (kit / "assets" / "icons" / "truck.png").is_file()


def test_generation_is_deterministic(ws, capsys, tmp_path):
    src = source(tmp_path)
    cli_json(ws, "brand", "init", "pemberton", "--from", str(src), capsys=capsys)
    kit = ws / "brands" / "pemberton"
    first = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (kit / "template.potx", kit / "tokens.yaml")}
    cli_json(ws, "brand", "init", "pemberton", "--from", str(src), "--force", capsys=capsys)
    again = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (kit / "template.potx", kit / "tokens.yaml")}
    assert first == again


def _layout_xml(kit, name):
    from lxml import etree

    prs = open_template(kit / "template.potx")
    layout = next(lay for lay in prs.slide_layouts if lay.name == name)
    return etree.tostring(layout._element).decode()


def _kit(ws, capsys, tmp_path, **gen):
    code, out = cli_json(ws, "brand", "init", "pemberton", "--from", str(source(tmp_path, **gen)), "--force",
                         capsys=capsys)
    assert code == 0, out
    kit = ws / "brands" / "pemberton"
    return kit, yaml.safe_load((kit / "tokens.yaml").read_text())


def test_default_type_scale_is_sized_for_projection(ws, capsys, tmp_path):
    kit, tokens = _kit(ws, capsys, tmp_path, layout_set="full")
    content = _layout_xml(kit, "Content")
    assert 'sz="3200" b="1"' in content  # bold 32 pt titles
    assert 'sz="2400"' in content and 'anchor="ctr"' in content  # 24 pt body, centered in its area
    assert 'tIns="45720" bIns="594360"' in content  # lifted 0.3 in to the optical center
    assert '<a:buClr><a:schemeClr val="tx2"/></a:buClr>' in content  # bullets in the primary color
    assert "buAutoNum" in _layout_xml(kit, "Agenda")
    assert tokens["table"]["font_size"] == 18 and tokens["table"]["header_font_size"] == 18  # projected floor


def test_read_mode_sets_document_sizes_a_measure_and_top_placement(ws, capsys, tmp_path):
    kit, tokens = _kit(ws, capsys, tmp_path, layout_set="full", mode="read")
    content = _layout_xml(kit, "Content")
    assert 'sz="2800"' in content and 'sz="1400"' in content and 'anchor="t"' in content
    assert tokens["table"]["font_size"] == 12
    from lxml import etree

    body = next(p for p in open_template(kit / "template.potx").slide_layouts.get_by_name("Content").placeholders
                if p.placeholder_format.idx == 1)
    assert body.width == 9 * 914400  # about 90 characters a line at 14 pt
    assert etree.tostring(body._element)  # parses


def test_table_columns_never_break_a_word():
    from deck_builder.build.visuals import column_widths
    from deck_builder.model import Table

    t = Table(header=["Account", "Renewal", "Status"],
              rows=[["Riverside Community College District Office", "November", "Green"]])
    widths = column_widths(t, 12 * 914400, 18)
    assert sum(widths) == 12 * 914400
    november = 8 * 18 * 0.55 * 12700  # the longest word's estimated width at 18 pt
    assert widths[1] > november and widths[2] > 5 * 18 * 0.55 * 12700


def test_type_scale_and_placement_come_from_brand_yaml(ws, capsys, tmp_path):
    _, default_tokens = _kit(ws, capsys, tmp_path / "a", layout_set="full")
    kit, tokens = _kit(ws, capsys, tmp_path / "b", layout_set="full", body_anchor="top", big_number="dark",
                       type={"title": 28, "title_bold": False, "body": 20, "table": 12})
    content = _layout_xml(kit, "Content")
    assert 'sz="2800"' in content and 'sz="3200"' not in content and 'b="1"' not in content.split("idx")[0]
    assert 'sz="2000"' in content and 'anchor="t"' in content
    big = _layout_xml(kit, "Big Number")
    assert 'showMasterSp="0"' in big and '<a:schemeClr val="tx2"/></a:solidFill>' in big  # dark background
    assert tokens["table"]["font_size"] == 12
    # budgets follow the sizes: smaller type fits more characters
    body, default_body = (t["layouts"]["content"]["fields"]["body"] for t in (tokens, default_tokens))
    assert body["max_chars"] > default_body["max_chars"]


def test_a_kit_records_its_recipe_regenerates_from_it_and_reports_drift(ws, capsys, tmp_path):
    kit, tokens = _kit(ws, capsys, tmp_path)
    assert tokens["generated"]["by"].startswith("deck-builder ")
    assert len(tokens["generated"]["inputs_sha256"]) == 64

    def stale():
        code, out = cli_json(ws, "brand", "check", "pemberton", capsys=capsys)
        return code, [i for i in out["issues"] if i["code"] == "KIT_STALE"]

    assert stale() == (0, [])
    # a changed ingredient (the logo) makes the kit stale, as a warning
    Image.new("RGB", (600, 200), "#E07A2F").save(kit / "assets" / "logo.png")
    code, found = stale()
    assert code == 0 and [i["severity"] for i in found] == ["warning"]
    # the kit regenerates from its own brand.yaml and assets, with no --from
    code, out = cli_json(ws, "brand", "init", "pemberton", "--force", capsys=capsys)
    assert code == 0, out
    assert stale() == (0, [])
    # an icon is read at build time, not baked into the template, so adding one never makes the kit stale
    Image.new("RGBA", (256, 256), (0, 0, 0, 255)).save(kit / "assets" / "icons" / "box.png")
    assert stale() == (0, [])
    code, out = cli_json(ws, "brand", "init", "pemberton", capsys=capsys)
    assert code == 2  # an existing kit is never regenerated without --force


def test_regenerating_a_kit_keeps_a_copy_of_the_kit_it_replaces(ws, capsys, tmp_path):
    kit, _ = _kit(ws, capsys, tmp_path)
    tuned = (kit / "tokens.yaml").read_text().replace("max_chars: ", "max_chars: 1", 1)  # a budget tuned by hand
    (kit / "tokens.yaml").write_text(tuned)
    code, out = cli_json(ws, "brand", "init", "pemberton", "--force", capsys=capsys)
    assert code == 0, out
    first = Path(out["backup"])
    assert first.parent == kit / "backups" and (first / "tokens.yaml").read_text() == tuned
    assert len(out["budget_changes"]) == 1 and " max_chars 1" in out["budget_changes"][0]  # the tuned one, named
    assert {"brand.yaml", "template.potx", "assets"} <= {p.name for p in first.iterdir()}
    code, out = cli_json(ws, "brand", "init", "pemberton", "--force", capsys=capsys)
    second = Path(out["backup"])
    assert second != first and "backups" not in {p.name for p in second.iterdir()}  # no backups of backups
    code, out = cli_json(ws, "brand", "list", capsys=capsys)
    assert sorted(b["slug"] for b in out["brands"]) == ["pemberton", "stock"]  # a backup is never a second kit


def test_brand_copy_makes_a_renamed_kit_and_leaves_the_original(ws, capsys, tmp_path):
    kit, _ = _kit(ws, capsys, tmp_path)
    tuned = (kit / "tokens.yaml").read_text().replace("max_chars: ", "max_chars: 1", 1)
    (kit / "tokens.yaml").write_text(tuned)
    before = {p.relative_to(kit): p.read_bytes() for p in kit.rglob("*") if p.is_file()}
    code, out = cli_json(ws, "brand", "copy", "pemberton", "pemberton-2026", capsys=capsys)
    assert code == 0, out
    new = ws / "brands" / "pemberton-2026"
    assert (new / "tokens.yaml").read_text() == tuned  # tuned budgets and template edits come along
    assert yaml.safe_load((new / "brand.yaml").read_text())["slug"] == "pemberton-2026"
    assert {p.relative_to(kit): p.read_bytes() for p in kit.rglob("*") if p.is_file()} == before  # untouched
    code, out = cli_json(ws, "brand", "check", "pemberton-2026", capsys=capsys)
    assert code == 0 and "KIT_STALE" not in [i["code"] for i in out["issues"]]
    code, out = cli_json(ws, "brand", "copy", "pemberton", "pemberton-2026", capsys=capsys)
    assert code == 2 and "already" in out["error"]  # never over an existing kit


def test_replacing_an_asset_keeps_a_copy_of_the_kit(ws, capsys, tmp_path):
    kit, _ = _kit(ws, capsys, tmp_path)
    old = (kit / "assets" / "logo.png").read_bytes()
    new = tmp_path / "new.png"
    Image.new("RGB", (600, 200), "#E07A2F").save(new)
    code, out = cli_json(ws, "brand", "add-asset", "pemberton", str(new), "--as", "logo/primary", "--force",
                         capsys=capsys)
    assert code == 0, out
    assert (Path(out["backup"]) / "assets" / "logo.png").read_bytes() == old


def test_a_kit_generated_before_kits_recorded_their_maker_reports_stale(ws, capsys, tmp_path):
    kit, tokens = _kit(ws, capsys, tmp_path)
    del tokens["generated"]  # what brand init wrote before 0.2.0: its header, and no generated record
    (kit / "tokens.yaml").write_text("# Generated by `deck-builder brand init`. max_chars and bullet budgets are "
                                     "estimates;\n" + yaml.safe_dump(tokens, sort_keys=False))
    code, out = cli_json(ws, "brand", "check", "pemberton", capsys=capsys)
    found = [i for i in out["issues"] if i["code"] == "KIT_STALE"]
    assert code == 0 and len(found) == 1 and "older than 0.2.0" in found[0]["message"]


def test_unknown_type_key_is_refused(ws, capsys, tmp_path):
    code, out = cli_json(ws, "brand", "init", "pemberton", "--from", str(source(tmp_path, type={"huge": 99})),
                         capsys=capsys)
    assert code != 0
    assert "huge" in str(out)


def test_columns_get_structure_and_title_slides_an_accent_rule(ws, capsys, tmp_path):
    kit, _ = _kit(ws, capsys, tmp_path, layout_set="full")
    comparison = _layout_xml(kit, "Comparison")
    panels = comparison.count('<a:schemeClr val="bg2"/>')
    assert panels == 2 and comparison.index("Decoration") < comparison.index("<p:ph ")  # drawn behind
    assert "lumMod" in _layout_xml(kit, "Two Column")  # the divider is a tint, not a hard line
    for name in ("Title", "Section", "Closing", "Big Number"):
        assert '<a:schemeClr val="accent2"/>' in _layout_xml(kit, name), name


def test_takeaway_band_is_drawn_only_when_used(ws, capsys, tmp_path):
    kit, tokens = _kit(ws, capsys, tmp_path)
    for key in ("content", "two-col", "chart", "table"):
        assert tokens["layouts"][key]["fields"]["takeaway"]["kind"] == "text", key
    assert '<a:solidFill><a:schemeClr val="tx2"/></a:solidFill>' in _layout_xml(kit, "Content")
    deck = write_deck(ws, """
    ---
    brand: pemberton
    ---

    ## Copy paper led growth
    layout: content
    takeaway: Two contracts explain the growth

    - Two new accounts

    ## Margins held
    layout: content

    - Flat quarter
    """)
    code, out = cli_json(ws, "build", str(deck), capsys=capsys)
    assert code == 0, out["issues"]
    with_band, without = Presentation(out["output"]).slides
    assert [p.text_frame.text for p in with_band.placeholders if p.placeholder_format.idx == 7] == [
        "Two contracts explain the growth"]
    assert [p for p in without.placeholders if p.placeholder_format.idx == 7] == []
    code, imp = cli_json(ws, "import", out["output"], str(ws / "back"), "--brand", "pemberton", capsys=capsys)
    assert code == 0, imp
    assert "takeaway: Two contracts explain the growth" in (ws / "back" / "deck.md").read_text()


def test_tables_have_rules_and_content_led_column_widths(ws, capsys, tmp_path):
    from lxml import etree

    _, tokens = _kit(ws, capsys, tmp_path)
    assert {"rule", "row_fill", "text", "row_height_factor"} <= set(tokens["table"])
    assert "band_fill" not in tokens["table"]
    assert 3 <= tokens["layouts"]["table"]["fields"]["table"]["max_rows"] <= 10
    deck = write_deck(ws, """
    ---
    brand: pemberton
    ---

    ## Accounts
    layout: table

    | Account | Tier |
    |---|---|
    | Northfield Regional School District | A |
    | Harbor County | B |
    """)
    code, out = cli_json(ws, "build", str(deck), capsys=capsys)
    assert code == 0, out["issues"]
    frame = next(s for s in Presentation(out["output"]).slides[0].shapes if s.has_table)
    wide, narrow = (col.width for col in frame.table.columns)
    assert wide > 2 * narrow and wide + narrow == frame.width
    header, body = (etree.tostring(frame.table.cell(r, 0)._tc).decode() for r in (0, 1))
    assert "<a:lnB" in body and "srgbClr" in body.split("<a:lnB")[1].split("</a:lnB>")[0]
    assert "<a:lnL" in body and "noFill" in body.split("<a:lnL")[1].split("</a:lnL>")[0]
    assert "noFill" in header.split("<a:lnB")[1].split("</a:lnB>")[0]


def test_numeric_columns_right_align_and_bars_start_at_zero(ws, capsys, tmp_path):
    from pptx.enum.text import PP_ALIGN

    _kit(ws, capsys, tmp_path)
    deck = write_deck(ws, """
    ---
    brand: pemberton
    ---

    ## Margins
    layout: table

    | Line | Margin | Change |
    |---|---|---|
    | Core | 23% | +2 pts |
    | Specialty | 31% | flat |

    ## Volume
    layout: chart

    ```chart
    type: column
    categories: [Jul, Aug]
    series: [{name: Core, values: [9800, 11250]}]
    ```
    """)
    code, out = cli_json(ws, "build", str(deck), capsys=capsys)
    assert code == 0, out["issues"]
    table_slide, chart_slide = Presentation(out["output"]).slides
    table = next(s for s in table_slide.shapes if s.has_table).table
    aligns = [table.cell(1, c).text_frame.paragraphs[0].alignment for c in range(3)]
    assert aligns == [None, PP_ALIGN.RIGHT, None]  # "flat" keeps the change column left-aligned
    assert table.cell(0, 1).text_frame.paragraphs[0].alignment == PP_ALIGN.RIGHT  # its header too
    chart = next(s for s in chart_slide.shapes if s.has_chart).chart
    assert chart.value_axis.minimum_scale == 0
    assert chart.plots[0].vary_by_categories is False
    rounded = chart._chartSpace.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}roundedCorners")
    assert rounded is not None and rounded.get("val") == "0"
    order = [el.tag.split("}")[1] for el in chart._chartSpace]
    assert order.index("roundedCorners") < order.index("chart")


def test_generated_kit_builds_every_layout(ws, capsys, tmp_path):
    cli_json(ws, "brand", "init", "pemberton", "--from", str(source(tmp_path, layout_set="full")), capsys=capsys)
    deck = write_deck(ws, """
    ---
    brand: pemberton
    ---

    ## Pemberton Paper review
    layout: title
    subtitle: Q3 shipments

    ## Agenda
    layout: agenda

    - Volume
    - Margins

    ## Volume
    layout: section
    kicker: Part 1

    ## Copy paper led growth
    layout: content

    - Two new accounts

    ## Before and after
    layout: two-col

    ### left
    - Manual counts

    ### right
    - Scanned counts

    ## Options
    layout: comparison
    left-heading: Lease
    right-heading: Buy

    ### left
    - Lower upfront

    ### right
    - Lower total

    ## 12%
    layout: big-number
    caption: Volume growth

    ## By month
    layout: chart

    ```chart
    categories: [Jul, Aug]
    series: [{name: Copy, values: [1, 2]}]
    ```

    ## Open items
    layout: table

    | Item | Owner |
    |---|---|
    | Lease | Ops |

    ## Warehouse
    layout: image

    ### image
    ![Logo](brand:logo/primary)

    ### caption
    The main site.

    ## Routes
    layout: image-right

    - North route

    ### image
    ![Logo](brand:logo/primary)

    ## How we ship
    layout: icon-row
    icon1: brand:icon/truck
    text1: Daily trucks

    ## Team
    layout: team
    name1: Ops lead

    ## "Paper is still the fastest way to sign a deal."
    layout: quote
    attribution: A regional buyer

    ## Thank you
    layout: closing
    subtitle: ops@example.com
    """)
    code, out = cli_json(ws, "build", str(deck), capsys=capsys)
    assert code == 0, out["issues"]
    prs = Presentation(out["output"])
    assert len(prs.slides) == 15
    assert prs.slides[13].placeholders[1].text_frame.text.startswith('"Paper is still')
