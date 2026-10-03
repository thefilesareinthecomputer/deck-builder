"""`build:`: the parts of one slide fade in on click. Never a transition between slides."""
from __future__ import annotations

import hashlib

import pytest
import yaml
from pptx import Presentation
from pptx.oxml.ns import qn

from conftest import cli_json, codes, write_deck
from deck_builder.cli import main
from deck_builder.parse import markdown

BRAND = {"spec_version": 1, "name": "Fade", "slug": "fade", "version": "1.0.0",
         "palette": {"primary": "1F3A5F", "accent": "E07A2F", "ink": "1B1B1B", "muted": "6B7280",
                     "surface": "F2F2F2", "background": "FFFFFF"},
         "fonts": {"heading": {"family": "Arial"}, "body": {"family": "Arial"}},
         "generate": {"layout_set": "designed"}}

DECK = """
---
brand: fade
---

## Three things changed
layout: content
build: body

- A second bay
- A tester bar
- Weekly restocks

## Pilot sales pulled away
layout: chart
build: chart

```chart
type: line
categories: [W1, W2, W3]
series:
  - name: Pilot
    values: [22, 27, 31]
  - name: Control
    values: [17, 18, 17]
```

## The rollout runs in three steps
layout: process-3
build: slots
step1: Survey
text1: Measure every shelf
step2: Order
text2: Bays and testers
step3: Fit
text3: One store a day
"""


@pytest.fixture
def root(tmp_path):
    assert main(["init", "--dir", str(tmp_path)]) == 0
    (tmp_path / "fade.yaml").write_text(yaml.safe_dump(BRAND), encoding="utf-8")
    assert main(["--config", str(tmp_path / "deck-builder.toml"), "brand", "init", "fade", "--from",
                 str(tmp_path / "fade.yaml")]) == 0
    (tmp_path / "decks").mkdir()
    return tmp_path


def build(root, capsys, name="deck.pptx"):
    deck = write_deck(root, DECK)
    capsys.readouterr()
    code, out = cli_json(root, "build", str(deck), "-o", str(root / "out" / name), capsys=capsys)
    assert code == 0, out["issues"]
    return deck, root / "out" / name


def clicks(slide) -> list[list[str]]:
    """Per click, the ids of the shapes that fade in, with a paragraph index where the build is per item."""
    seq = slide._element.find(f".//{qn('p:cTn')}[@nodeType='mainSeq']")
    out = []
    for step in seq.find(qn("p:childTnLst")):
        targets = []
        for t in step.iter(qn("p:spTgt")):
            pr = t.find(f"{qn('p:txEl')}/{qn('p:pRg')}")
            targets.append(t.get("spid") + (f"#{pr.get('st')}" if pr is not None else ""))
        out.append(list(dict.fromkeys(targets)))
    return out


def test_a_list_fades_in_item_by_item_a_chart_whole_and_slots_one_group_per_click(root, capsys):
    _, pptx = build(root, capsys)
    slides = Presentation(str(pptx)).slides
    items = clicks(slides[0])
    assert len(items) == 3 and all(len(c) == 1 and c[0].endswith(f"#{i}") for i, c in enumerate(items))
    assert slides[0]._element.find(f".//{qn('p:bldP')}").get("build") == "p"
    assert len(clicks(slides[1])) == 1 and slides[1]._element.find(f".//{qn('p:bldGraphic')}") is not None
    steps = clicks(slides[2])
    assert len(steps) == 3 and all(len(s) == 2 for s in steps)  # each step's label and its text together
    for slide in slides:
        fades = slide._element.findall(f".//{qn('p:animEffect')}")
        assert fades and all(f.get("filter") == "fade" and f.get("transition") == "in" for f in fades)
        assert {d.get("dur") for f in fades for d in f.iter(qn("p:cTn"))} == {"500"}
        assert slide._element.find(qn("p:transition")) is None  # never a transition between slides


def test_a_build_is_byte_identical_and_round_trips_through_the_workbook_and_import(root, capsys):
    deck, a = build(root, capsys, "a.pptx")
    _, b = build(root, capsys, "b.pptx")
    assert hashlib.sha256(a.read_bytes()).digest() == hashlib.sha256(b.read_bytes()).digest()
    _, out = cli_json(root, "convert", str(deck), str(root / "decks" / "deck.xlsx"), capsys=capsys)
    assert out["ok"], out["issues"]
    _, out = cli_json(root, "import", str(a), str(root / "imported"), "--brand", "fade", capsys=capsys)
    assert out["ok"], out
    assert [s.build for s in markdown.parse(root / "imported" / "deck.md")[0].slides] == ["body", "chart", "slots"]


@pytest.mark.parametrize("line,why", [
    ("build: notes", "must be `slots` or a field"),
    ("build: slots", "two or more filled slots"),
])
def test_build_must_name_a_field_or_slots_the_slide_has(root, capsys, line, why):
    deck = write_deck(root, f"---\nbrand: fade\n---\n\n## One list\nlayout: content\n{line}\n\n- One\n- Two\n")
    capsys.readouterr()
    _, out = cli_json(root, "check", str(deck), capsys=capsys)
    assert codes(out) == ["PARSE"] and why in out["issues"][0]["message"]


def test_a_code_block_never_builds(root, capsys):
    deck = write_deck(root, "---\nbrand: fade\n---\n\n## Code\nlayout: code\nbuild: code\n\n```python\nx = 1\n```\n")
    capsys.readouterr()
    _, out = cli_json(root, "check", str(deck), capsys=capsys)
    assert codes(out) == ["PARSE"] and "doesn't build" in out["issues"][0]["message"]
