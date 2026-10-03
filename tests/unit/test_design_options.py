"""The designed set's generate options and the build features behind them: counts, takeaway styles, icon
tiles, process icons, band label shapes, bold in the primary color, logo rows and fitting, the deck-wide
kicker, first_slide_number, starter icons, and the convention warnings from `docs design`."""
from __future__ import annotations

import itertools
import json
from pathlib import Path

import pytest
import yaml
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from conftest import cli_json, codes, write_deck
from deck_builder.brand.layouts import FOOTER, MARGIN, SETS, Style, layout_set
from deck_builder.cli import main

W, H = 13.333, 7.5
BRAND = {"spec_version": 1, "name": "Options", "slug": "opts", "version": "1.0.0",
         "palette": {"primary": "1F3A5F", "accent": "E07A2F", "ink": "1B1B1B", "surface": "F2F2F2",
                     "background": "FFFFFF"},
         "fonts": {"heading": {"family": "Arial"}, "body": {"family": "Arial"}}}


def kit(root: Path, **gen: object) -> Path:
    """A workspace at root with brand `opts` generated with the designed set and these generate options."""
    assert main(["init", "--dir", str(root)]) == 0
    src = root / "opts-src"
    src.mkdir()
    meta = {**BRAND, "generate": {"layout_set": "designed", **gen}}
    (src / "brand.yaml").write_text(yaml.safe_dump(meta), encoding="utf-8")
    assert main(["--config", str(root / "deck-builder.toml"), "brand", "init", "opts", "--from",
                 str(src / "brand.yaml")]) == 0
    (root / "decks").mkdir()
    return root


def tokens(root: Path) -> dict:
    return yaml.safe_load((root / "workspace" / "brands" / "opts" / "tokens.yaml").read_text())


def logo(path: Path, size: tuple[int, int] = (1600, 400)) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    im = Image.new("RGBA", size, (0, 0, 0, 0))
    im.paste((31, 58, 95, 255), (0, size[1] // 4, size[0], 3 * size[1] // 4))
    im.save(path)


def by_key(**style: object) -> dict:
    return {ld.key: ld for ld in layout_set("designed", W, H, style=Style(**style))}  # type: ignore[arg-type]


# ---------------------------------------------------------------- geometry


def test_the_designed_set_has_every_count_and_the_logo_row():
    assert {"cards-2", "cards-5", "process-3", "process-6", "bands-2", "bands-4", "logos"} <= set(SETS["designed"])


@pytest.mark.parametrize("takeaway,tile,label,icons", list(itertools.product(
    ("band", "quote"), ("none", "square", "circle"), ("parallelogram", "rectangle"), (False, True))))
def test_every_option_keeps_every_box_inside_the_margins(takeaway, tile, label, icons):
    st = Style(takeaway=takeaway, icon_tile=tile, band_label=label, process_icons=icons)
    for ld in layout_set("designed", W, H, style=st):
        for ph in ld.phs:
            assert ph.x >= MARGIN - 1e-6 and ph.x + ph.w <= W - MARGIN + 1e-6, (ld.key, ph.field)
            assert ph.y + ph.h <= H - FOOTER + 1e-6, (ld.key, ph.field)
        idxs = [ph.idx for ph in ld.phs if ph.kind != "title"]
        assert len(idxs) == len(set(idxs)), ld.key


def test_options_left_out_are_todays_layouts():
    assert Style.from_meta({}) == Style()
    full = {ld.key: ld for ld in layout_set("full", W, H)}
    take = next(ph for ph in full["content"].phs if ph.field == "takeaway")
    assert take.fill == "tx2" and take.bold and not take.italic
    assert full["icon-row"].decor == [] and all(ph.color is None for ph in full["icon-row"].phs if ph.icon)


def test_a_quote_takeaway_is_an_italic_line_with_no_band():
    take = next(ph for ph in by_key(takeaway="quote")["cards-3"].phs if ph.field == "takeaway")
    assert take.fill is None and take.italic and take.align == "ctr" and take.color == "tx2"


def test_icon_tiles_sit_behind_white_icons_and_stay_square_under_the_header():
    ld = by_key(icon_tile="circle")["icon-row"]
    tiles = [d for d in ld.decor if d.geom == "ellipse"]
    icons = [ph for ph in ld.phs if ph.icon]
    assert len(tiles) == 3 and all(abs(t.w - t.h) < 1e-9 for t in tiles)
    for t, ph in zip(tiles, icons, strict=True):
        assert ph.color == "bg1"
        assert abs((t.x + t.w / 2) - (ph.x + ph.w / 2)) < 1e-6 and abs((t.y + t.h / 2) - (ph.y + ph.h / 2)) < 1e-6


def test_process_icons_sit_centered_in_their_chevrons_with_the_title_below():
    ld = by_key(process_icons=True)["process-4"]
    chevrons = [d for d in ld.decor if d.geom in ("homePlate", "chevron")]
    for i, chev in enumerate(chevrons, start=1):
        fields = {ph.field: ph for ph in ld.phs}
        icon, step = fields[f"icon{i}"], fields[f"step{i}"]
        assert icon.icon and icon.color == "bg1" and icon.required
        assert abs((icon.y + icon.h / 2) - (chev.y + chev.h / 2)) < 1e-6
        assert step.y >= chev.y + chev.h and step.color == "tx2"


def test_band_labels_can_be_rectangles():
    geoms = {d.geom for d in by_key(band_label="rectangle")["bands-3"].decor}
    assert "rect" in geoms and "parallelogram" not in geoms


@pytest.mark.parametrize("primary", ["23443D", "293D59", "302F35", "1F3A5F", "E07A2F", "2A9D8F"])
def test_white_label_text_keeps_its_contrast_on_every_ramp_step(primary):
    """The ramp's lightest step is set from the brand's primary, so white bold labels stay readable."""
    from deck_builder.brand.generate import RAMP_CONTRAST, ramp_floor, tint
    from deck_builder.brand.kit import contrast

    floor = ramp_floor(primary, "FFFFFF")
    process = {ld.key: ld for ld in layout_set("designed", W, H, style=Style(ramp_floor=floor))}["process-6"]
    steps = [int(d.fill.split("@")[1]) if "@" in d.fill else 100 for d in process.decor]
    assert min(steps) == floor
    if contrast(primary, "FFFFFF") >= RAMP_CONTRAST:
        assert all(contrast(tint(primary, pct), "FFFFFF") >= RAMP_CONTRAST for pct in steps)
    else:  # a primary too light for white text keeps every step at the primary itself
        assert floor == 100 and set(steps) == {100}


# ---------------------------------------------------------------- tokens


def test_tokens_record_icon_colors_emphasis_fitting_and_the_logo_row(tmp_path):
    t = tokens(kit(tmp_path, icons={"tile": "square"}, process={"icons": True}))["layouts"]
    assert t["icon-row"]["fields"]["icon1"] == {"idx": 10, "kind": "icon", "color": "background", "required": True}
    assert t["process-4"]["fields"]["icon1"]["kind"] == "icon"
    assert t["cards-3"]["fields"]["body1"]["emphasis"] == "primary"
    assert t["bands-2"]["fields"]["text1"]["emphasis"] == "primary"
    assert t["image-right"]["fields"]["image"]["fit_max"] == 0.6
    assert t["comparison"]["fields"]["left-logo"]["fit"] is True
    assert t["logos"]["row"] is True and t["logos"]["fields"]["logo3"]["fit"] is True


def test_no_generated_layout_repeats_a_placeholder_idx_with_its_furniture(tmp_path):
    """The kicker and subtitle once took the footer's and slide number's idx (20, 21) on every layout."""
    from deck_builder.template import open_template

    root = kit(tmp_path, icons={"tile": "circle"}, process={"icons": True})
    prs = open_template(root / "workspace" / "brands" / "opts" / "template.potx")
    assert len(prs.slide_layouts) == len(SETS["designed"])
    for layout in prs.slide_layouts:
        idxs = [ph.placeholder_format.idx for ph in layout.placeholders]
        assert len(idxs) == len(set(idxs)), (layout.name, sorted(idxs))


def test_emphasis_ink_leaves_bold_in_the_text_color(tmp_path):
    t = tokens(kit(tmp_path, emphasis="ink"))["layouts"]
    assert "emphasis" not in t["cards-3"]["fields"]["body1"]


# ---------------------------------------------------------------- build


def built(root: Path, body: str, capsys: pytest.CaptureFixture[str]) -> tuple[dict, Presentation]:
    deck = write_deck(root, body)
    capsys.readouterr()  # init's own output
    code, out = cli_json(root, "build", str(deck), "-o", str(root / "out" / "deck.pptx"), capsys=capsys)
    assert code == 0, out
    return out, Presentation(str(root / "out" / "deck.pptx"))


def texts(slide) -> list[str]:
    return [sh.text_frame.text for sh in slide.shapes if sh.has_text_frame]


def test_a_front_matter_kicker_labels_every_slide_that_sets_none(tmp_path, capsys):
    root = kit(tmp_path)
    _, prs = built(root, """
        ---
        brand: opts
        kicker: Q4 plan
        ---
        ## The plan is three bets
        layout: content

        - One

        ## Bet one is invoicing
        layout: content
        kicker: Bet 1

        - Two
        """, capsys)
    assert "Q4 plan" in texts(prs.slides[0])
    assert "Bet 1" in texts(prs.slides[1]) and "Q4 plan" not in texts(prs.slides[1])


def test_first_slide_number_numbers_an_excerpt_from_there(tmp_path, capsys):
    root = kit(tmp_path)
    body = ("---\nbrand: opts\nfirst_slide_number: 88\n---\n"
            "## One\nlayout: content\n\n- A\n\n## Two\nlayout: content\n\n- B\n")
    _, prs = built(root, body, capsys)
    assert prs.part._element.get("firstSlideNum") == "88"
    assert "89" in texts(prs.slides[1])
    p = write_deck(root, body.replace("88", "0"), "bad.md")
    _, out = cli_json(root, "check", str(p), capsys=capsys)
    assert "PARSE" in codes(out) and "first_slide_number" in str(out["issues"])


def test_bold_on_cards_takes_the_primary_color(tmp_path, capsys):
    root = kit(tmp_path)
    _, prs = built(root, """
        ---
        brand: opts
        ---
        ## Two ways to ship
        layout: cards-2
        label1: Now
        label2: Next

        ### body1
        - **Weekly** routes

        ### body2
        - **Daily** routes
        """, capsys)
    runs = [r for sh in prs.slides[0].shapes if sh.has_text_frame for p in sh.text_frame.paragraphs for r in p.runs]
    bold = [r for r in runs if r.font.bold and r.text in ("Weekly", "Daily")]
    assert len(bold) == 2 and all(str(r.font.color.rgb) == "1F3A5F" for r in bold)


def test_a_logo_row_spaces_its_filled_slots_evenly(tmp_path, capsys):
    root = kit(tmp_path)
    for name in ("a", "b"):
        logo(root / "decks" / "assets" / f"{name}.png")
    _, prs = built(root, """
        ---
        brand: opts
        ---
        ## The tools this runs on
        layout: logos
        caption1: Ingestion
        caption2: Modeling

        ### logo1
        ![Tool A](assets/a.png)

        ### logo2
        ![Tool B](assets/b.png)
        """, capsys)
    pics = sorted((sh for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE),
                  key=lambda sh: sh.left)
    centers = [(p.left + p.width / 2) / 914400 for p in pics]
    left, right = MARGIN, W - MARGIN
    assert len(centers) == 2
    quarter = (right - left) / 4
    assert abs(centers[0] - (left + quarter)) < 0.05 and abs(centers[1] - (left + 3 * quarter)) < 0.05
    captions = {sh.text_frame.text: (sh.left + sh.width / 2) / 914400 for sh in prs.slides[0].shapes
                if sh.has_text_frame and sh.text_frame.text in ("Ingestion", "Modeling")}
    assert abs(captions["Ingestion"] - centers[0]) < 0.05 and abs(captions["Modeling"] - centers[1]) < 0.05


def test_a_transparent_logo_in_a_large_box_takes_at_most_the_logo_share(tmp_path, capsys):
    root = kit(tmp_path)
    logo(root / "decks" / "assets" / "a.png")
    _, prs = built(root, """
        ---
        brand: opts
        ---
        ## The new account
        layout: image-right

        - Signed in September

        ### image
        ![Account logo](assets/a.png)
        """, capsys)
    pic = next(sh for sh in prs.slides[0].shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    box_w = next(ph for ph in prs.slides[0].slide_layout.placeholders if ph.placeholder_format.idx == 2).width
    assert abs(pic.width / box_w - 0.6) < 0.01 and pic.crop_left == 0


def test_a_starter_icon_fills_in_where_the_kit_has_none(ws, capsys):
    deck = write_deck(ws, "---\nbrand: stock\n---\n## Data lands here\nlayout: icon\nicon: brand:icon/database\n")
    code, out = cli_json(ws, "build", str(deck), "-o", str(ws / "out" / "deck.pptx"), capsys=capsys)
    if not (Path(__file__).resolve().parents[2] / "src" / "deck_builder" / "data" / "icons" / "database.png").is_file():
        pytest.skip("run scripts/make_starter_icons.py to draw the full starter set")
    assert code == 0, out
    manifest = json.loads((ws / "out" / "deck.manifest.json").read_text())
    assert manifest["slides"][0]["fields"]["icon"]["asset"]["source"] == "starter/database.png"
    code, out = cli_json(ws, "check", str(deck), capsys=capsys)
    assert code == 0, out
    code, shown = cli_json(ws, "brand", "show", "stock", capsys=capsys)
    assert "database" in shown["starter_icons"] and "check" not in shown["starter_icons"]  # the kit has its own


def test_an_icon_neither_the_kit_nor_the_starter_set_has_is_unknown(ws, capsys):
    deck = write_deck(ws, "---\nbrand: stock\n---\n## A\nlayout: icon\nicon: brand:icon/no-such-icon\n")
    _, out = cli_json(ws, "check", str(deck), capsys=capsys)
    assert "UNKNOWN_ASSET" in codes(out)


# ---------------------------------------------------------------- convention warnings


CHART_9 = "\n".join(f"  - name: S{i}\n    values: [1, 2]" for i in range(9))


def test_convention_warnings_flag_crowded_slides_and_runs_without_failing(ws, capsys):
    long = " ".join(["word"] * 32)
    deck = write_deck(ws, f"""---
brand: stock
---
## A
layout: content

- One
- Two
- Three
- Four
- Five

## B
layout: content

- One

## C
layout: content

- One

## D
layout: content

- One

## E
layout: two-col

### left
- {long}

### right
- {long}

## F
layout: chart

```chart
type: column
categories: [a, b]
series:
{CHART_9}
```
""")
    code, out = cli_json(ws, "check", str(deck), capsys=capsys)
    found = {(i["code"], i["slide"]) for i in out["issues"]}
    assert {("BULLETS_MANY", 1), ("LAYOUT_RUN", 4), ("WORDS_MANY", 5), ("SERIES_MANY", 6)} <= found
    assert code == 0 and {i["severity"] for i in out["issues"]} == {"warning"}


# ---------------------------------------------------------------- accessibility (WCAG 2.2 AA)


def test_check_flags_missing_alt_duplicate_titles_and_color_only_charts(ws, capsys):
    (ws / "decks" / "assets").mkdir()
    Image.new("RGB", (1600, 900), "#888888").save(ws / "decks" / "assets" / "p.png")
    deck = write_deck(ws, """---
brand: stock
---
## A photo with nothing said about it
layout: image

![](assets/p.png)

## Volume by month
layout: chart

```chart
type: column
legend: false
categories: [Jul, Aug]
series:
  - name: Core
    values: [1, 2]
  - name: Specialty
    values: [2, 1]
```

## Volume by month
layout: chart

```chart
type: column
legend: false
labels: true
categories: [Jul, Aug]
series:
  - name: Core
    values: [1, 2]
  - name: Specialty
    values: [2, 1]
```
""")
    code, out = cli_json(ws, "check", str(deck), capsys=capsys)
    found = {(i["code"], i["slide"]) for i in out["issues"]}
    assert {("MISSING_ALT", 1), ("COLOR_ONLY", 2), ("TITLE_DUPLICATE", 3)} <= found
    assert ("COLOR_ONLY", 3) not in found  # labels name the series
    assert code == 0


def test_brand_check_simulates_color_blindness_on_chart_colors(ws, capsys):
    from deck_builder.brand.kit import cvd_delta

    red, green, navy, orange = "D62728", "2CA02C", "1F3A5F", "E07A2F"
    assert cvd_delta(red, green, "deuteranopia") < cvd_delta(red, green, "tritanopia")
    assert cvd_delta(navy, orange, "deuteranopia") > 10  # the default navy and orange stay apart
    tok = yaml.safe_load((ws / "brands" / "stock" / "tokens.yaml").read_text())
    tok["chart"]["colors"] = ["5B8C2A", "C0602A"]  # a green and an orange-red that converge for red-green CVD
    (ws / "brands" / "stock" / "tokens.yaml").write_text(yaml.safe_dump(tok))
    _, out = cli_json(ws, "brand", "check", "stock", capsys=capsys)
    assert "CVD_CONFUSABLE" in codes(out)


def test_brand_check_flags_type_under_the_legibility_floor(ws, capsys):
    meta = yaml.safe_load((ws / "brands" / "stock" / "brand.yaml").read_text())
    meta["generate"] = {"type": {"table": 14, "kicker": 18}}
    (ws / "brands" / "stock" / "brand.yaml").write_text(yaml.safe_dump(meta))
    _, out = cli_json(ws, "brand", "check", "stock", capsys=capsys)
    small = [i for i in out["issues"] if i["code"] == "TYPE_SMALL"]
    assert [i["actual"] for i in small] == [14]
    meta["generate"]["mode"] = "read"  # 14 pt is fine in a read deck
    (ws / "brands" / "stock" / "brand.yaml").write_text(yaml.safe_dump(meta))
    _, out = cli_json(ws, "brand", "check", "stock", capsys=capsys)
    assert "TYPE_SMALL" not in codes(out)


def test_projected_kits_have_no_size_under_18_pt():
    for ld in layout_set("designed", W, H):
        for ph in ld.phs:
            assert ph.size is None or ph.size >= 18, (ld.key, ph.field, ph.size)


# ---------------------------------------------------------------- fixes found by the brand fixture decks


def test_budgets_come_from_the_width_text_gets():
    from deck_builder.brand.layouts import BULLET_HANG, DEFAULT_INSET, PH

    plain_ph = PH("x", "body", 1, 0, 0, 2.0, 1.0, size=13)
    listed = PH("x", "body", 1, 0, 0, 2.0, 1.0, size=13, bullets=True)
    assert plain_ph.text_w == pytest.approx(2.0 - 2 * DEFAULT_INSET)
    assert listed.text_w == pytest.approx(2.0 - 2 * DEFAULT_INSET - BULLET_HANG)


def test_check_wraps_bullets_line_by_line_against_the_box(tmp_path, capsys):
    """Six short bullets each one word past a narrow card's line wrap to twelve lines, plus their spacing,
    though their character count is under budget: BUDGET_LINES, which the character budget can't see."""
    root = kit(tmp_path, mode="read")
    spec = tokens(root)["layouts"]["cards-5"]["fields"]["body1"]
    per = spec["line_chars"]
    word = "x" * (per - 2)  # with " ab" each bullet is one character past the line, so it wraps
    bullets = "\n".join(f"- {word} ab" for _ in range(min(6, spec["max_bullets"])))
    labels = "".join(f"label{i}: L{i}\n" for i in range(1, 6))
    bodies = "".join(f"\n### body{i}\n- One\n" for i in range(2, 6))
    deck = write_deck(root, f"---\nbrand: opts\n---\n## Five\nlayout: cards-5\n{labels}\n"
                            f"### body1\n{bullets}\n{bodies}")
    capsys.readouterr()
    _, out = cli_json(root, "check", str(deck), capsys=capsys)
    lines = [i for i in out["issues"] if i["code"] == "BUDGET_LINES"]
    assert lines and lines[0]["field"] == "body1"
    assert "BUDGET_CHARS" not in codes(out)


def test_bold_italic_markup_renders_as_one_bold_italic_run():
    from pptx import Presentation as P

    from deck_builder.build.text import add_runs

    prs = P()
    para = prs.slides.add_slide(prs.slide_layouts[1]).placeholders[1].text_frame.paragraphs[0]
    add_runs(para, "a ***bold italic*** b")
    runs = [(r.text, r.font.bold, r.font.italic) for r in para.runs]
    assert ("bold italic", True, True) in runs and "*" not in "".join(t for t, _, _ in runs)


def test_a_split_word_inside_its_own_box_isnt_another_fields_overflow():
    from deck_builder.qa.measure import Box, TextShape, _candidates

    chevron = TextShape("step3", Box(0, 0, 1, 1), {"load-tested"}, chart=False, protected=False)
    below = TextShape("text3", Box(0, 2, 1, 3), {"tested", "nightly"}, chart=False, protected=False)
    (d, i), *_ = sorted(_candidates("tested", 0.5, 0.5, [chevron, below]))
    assert (d, i) == (0, 0)  # inside the chevron that holds "load-tested", not the exact match below


def test_chart_fixes_negative_bars_label_contrast_and_a_light_accent(ws, capsys):
    from lxml import etree

    from deck_builder.brand.generate import chart_color
    from deck_builder.brand.kit import contrast

    assert chart_color("accent", "E89445", "FFFFFF") != "accent"
    assert contrast(chart_color("accent", "E89445", "FFFFFF"), "FFFFFF") >= 3.0
    assert chart_color("primary", "1F3A5F", "FFFFFF") == "primary"
    deck = write_deck(ws, """---
brand: stock
---
## Cost per bar against last year
layout: chart

```chart
type: stacked-column
labels: true
categories: [Oil, Wrap]
series:
  - name: Up
    values: [3, -1]
  - name: Down
    values: [1, -2]
```
""")
    code, out = cli_json(ws, "build", str(deck), "-o", str(ws / "out" / "c.pptx"), capsys=capsys)
    assert code == 0, out
    chart = next(sh.chart for sh in Presentation(str(ws / "out" / "c.pptx")).slides[0].shapes if sh.has_chart)
    xml = etree.tostring(chart._chartSpace).decode()
    assert 'invertIfNegative val="0"' in xml and 'tickLblPos val="low"' in xml
    assert xml.count("<c:showVal val=\"1\"/>") >= 2  # each series' labels shown, in white or ink by its fill


def _charts(ws, capsys, body: str) -> list:
    """Build the chart slides in body on the stock brand; the built charts, one per slide."""
    deck = write_deck(ws, "---\nbrand: stock\n---\n" + body)
    code, out = cli_json(ws, "build", str(deck), "-o", str(ws / "out" / "c.pptx"), capsys=capsys)
    assert code == 0, out
    return [sh.chart for s in Presentation(str(ws / "out" / "c.pptx")).slides for sh in s.shapes if sh.has_chart]


def _chart_slide(title: str, spec: str) -> str:
    return f"## {title}\nlayout: chart\n\n```chart\n{spec}```\n\n"


def test_charts_are_quiet_hairline_gridlines_no_ticks_no_markers_and_muted_axis_text(ws, capsys):
    from lxml import etree

    tokens = ws / "brands" / "stock" / "tokens.yaml"
    tok = yaml.safe_load(tokens.read_text())
    tok["chart"].update(axis_text_color="6B7280", gridline_color="E5E7EB")
    tokens.write_text(yaml.safe_dump(tok))
    line, column = _charts(ws, capsys, _chart_slide("Lines", "type: line\ncategories: [a, b, c]\nseries:\n"
                                                    "  - name: x\n    values: [1, 2, 3]\n") +
                           _chart_slide("Bars", "type: column\ncategories: [a, b]\nseries:\n"
                                        "  - name: x\n    values: [1, 2]\n  - name: y\n    values: [2, 1]\n"))
    xml = etree.tostring(line._chartSpace).decode()
    assert '<c:symbol val="none"/>' in xml  # lines have no markers
    assert 'majorTickMark val="none"' in xml and "6B7280" in xml  # no ticks; axis text in the muted color
    assert '<c:majorGridlines><c:spPr><a:ln w="6350">' in xml and "E5E7EB" in xml  # hairline gridlines
    assert column.plots[0].overlap == -8 and column.plots[0].gap_width == 70  # a cluster's bars stand apart


def test_a_labeled_bar_chart_drops_its_scale_and_reads_top_to_bottom(ws, capsys):
    from lxml import etree

    (chart,) = _charts(ws, capsys, _chart_slide("Shares", "type: bar\nlabels: true\ncategories: [a, b, c]\n"
                                                "series:\n  - name: x\n    values: [5, 3, 2]\n"))
    assert not chart.value_axis.has_major_gridlines  # each bar shows its own number
    assert chart.category_axis.reverse_order  # a, b, c from the top, in the order written
    val_ax = etree.tostring(chart.value_axis._element).decode()
    assert 'tickLblPos val="none"' in val_ax and 'crosses val="max"' in val_ax  # any scale stays below


def test_a_thin_stacked_segment_shows_no_label(ws, capsys):
    from lxml import etree

    (chart,) = _charts(ws, capsys, _chart_slide("Stack", "type: stacked-column\nlabels: true\ncategories: [a, b]\n"
                                                "series:\n  - name: big\n    values: [100, 90]\n"
                                                "  - name: small\n    values: [3, 40]\n"))
    small = etree.tostring(chart.plots[0].series[1]._element).decode()
    assert '<c:dLbl><c:idx val="0"/><c:delete val="1"/></c:dLbl>' in small  # 3 of 103: too thin for its label
    assert 'c:idx val="1"/><c:delete' not in small  # 40 keeps its label


def test_brand_init_writes_muted_axis_text_and_faint_gridlines():
    from deck_builder.brand.generate import tokens_for
    from deck_builder.brand.kit import contrast

    meta = {"palette": {"ink": "1B1B1B", "background": "FFFFFF", "muted": "6B7280", "primary": "1F3A5F"}}
    chart = tokens_for([], meta)["chart"]
    assert chart["axis_text_color"] == "muted"  # 4.8:1 on white
    assert contrast(chart["gridline_color"], "FFFFFF") < 1.5  # a faint tint that stays behind the data
    too_light = {"palette": {**meta["palette"], "muted": "B0B0B0"}}
    assert tokens_for([], too_light)["chart"]["axis_text_color"] == "ink"


def test_pie_and_doughnut_charts_still_build_with_slice_colors_and_readable_labels(ws, capsys):
    """The design rules use bars for shares, but decks can still hold a pie or doughnut (an imported deck)."""
    spec = "labels: true\ncategories: [a, b]\nseries:\n  - name: share\n    values: [60, 40]\n"
    pie, doughnut = _charts(ws, capsys, _chart_slide("Pie", f"type: pie\n{spec}") +
                            _chart_slide("Doughnut", f"type: doughnut\n{spec}"))
    for chart in (pie, doughnut):
        points = chart.plots[0].series[0].points
        assert [str(points[i].format.fill.fore_color.rgb) for i in range(2)] == ["1F3A5F", "E07A2F"]
        assert str(points[0].data_label.font.color.rgb) == "FFFFFF"  # white on the dark primary slice
        assert chart.has_legend  # on by default, so slices aren't told apart by color alone


def test_a_read_deck_skips_the_text_limits(ws, capsys):
    meta = yaml.safe_load((ws / "brands" / "stock" / "brand.yaml").read_text())
    meta["generate"] = {"mode": "read"}
    (ws / "brands" / "stock" / "brand.yaml").write_text(yaml.safe_dump(meta))
    deck = write_deck(ws, "---\nbrand: stock\n---\n## A\nlayout: content\n\n- 1\n- 2\n- 3\n- 4\n- 5\n")
    _, out = cli_json(ws, "check", str(deck), capsys=capsys)
    assert "BULLETS_MANY" not in codes(out)
