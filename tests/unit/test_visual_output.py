"""What the builder writes into charts and tables: alt text and cell fills."""
from pathlib import Path

import pytest
import yaml
from pptx import Presentation
from pptx.oxml.ns import qn
from test_check import GOOD

from conftest import cli_json, write_deck

TABLE_DECK = """
## Open items
layout: table

| Item | Owner | Due |
|---|---|---|
| Warehouse lease | Ops | Oct |
| Card stock order | Purchasing | Nov |
| Price review | Sales | Dec |
"""


def build(ws, capsys, body):
    code, out = cli_json(ws, "build", str(write_deck(ws, body)), capsys=capsys)
    assert code == 0, out["issues"]
    return Presentation(out["output"])


def frames(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_chart or shape.has_table:
                yield shape


def test_charts_and_tables_carry_alt_text(ws, capsys):
    descr = {("chart" if f.has_chart else "table"): f._element.nvGraphicFramePr.cNvPr.get("descr")
             for f in frames(build(ws, capsys, GOOD))}
    assert descr == {"chart": "Column chart of Copy paper by Jul, Aug, Sep",
                     "table": "Table with columns Item, Owner, Due, 1 row"}


def set_table_tokens(ws, **tokens):
    path = ws / "brands" / "stock" / "tokens.yaml"
    data = yaml.safe_load(path.read_text())
    data["table"].update(tokens)
    path.write_text(yaml.safe_dump(data, sort_keys=False))


def fills(prs):
    """Per row: the first cell's fill, as a hex, 'none' or None (left to the table style)."""
    table = next(frames(prs)).table
    out = []
    for row in table.rows:
        tcPr = row.cells[0]._tc.find(qn("a:tcPr"))
        solid = tcPr.find(qn("a:solidFill")) if tcPr is not None else None
        if solid is not None:
            out.append(solid.find(qn("a:srgbClr")).get("val"))
        elif tcPr is not None and tcPr.find(qn("a:noFill")) is not None:
            out.append("none")
        else:
            out.append(None)
    return table, out


@pytest.mark.parametrize("tokens, expected", [
    ({"band_fill": "F2F2F2"}, ["1F3A5F", "none", "F2F2F2", "none"]),
    ({"band_fill": "F2F2F2", "row_fill": "FFFFFF"}, ["1F3A5F", "FFFFFF", "F2F2F2", "FFFFFF"]),
    ({}, ["1F3A5F", "none", "none", "none"]),
])
def test_table_rows_take_only_the_token_fills(ws, capsys, tokens, expected):
    set_table_tokens(ws, **tokens)
    table, got = fills(build(ws, capsys, TABLE_DECK))
    assert got == expected
    assert table.horz_banding is False  # the default table style's bands stay off
    assert table.first_row is True


def test_built_table_file_has_no_style_banding(ws, capsys):
    set_table_tokens(ws, band_fill="F2F2F2")
    prs = build(ws, capsys, TABLE_DECK)
    tblPr = next(frames(prs)).table._tbl.tblPr
    assert tblPr.get("bandRow") in (None, "0")
    assert Path(ws / "out" / "decks.pptx").is_file()
