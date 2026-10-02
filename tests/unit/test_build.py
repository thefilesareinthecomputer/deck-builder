import hashlib
import json
import time
import zipfile
from pathlib import Path

from pptx import Presentation
from test_check import GOOD

from conftest import cli_json, codes, write_deck


def build(ws, capsys, body=GOOD, *extra):
    deck = write_deck(ws, body)
    code, out = cli_json(ws, "build", str(deck), *extra, capsys=capsys)
    return code, out


def test_build_writes_pptx_and_manifest(ws, capsys):
    code, out = build(ws, capsys)
    assert code == 0, out["issues"]
    pptx, manifest = Path(out["output"]), Path(out["manifest"])
    assert pptx == ws / "out" / "decks.pptx"  # decks/deck.md takes its folder's name
    assert pptx.is_file() and manifest.is_file()
    prs = Presentation(str(pptx))
    assert len(prs.slides) == 7
    assert prs.slides[0].shapes.title.text == "Pemberton Paper quarterly review"
    m = json.loads(manifest.read_text())
    assert [s["layout"] for s in m["slides"]][:3] == ["title", "content", "big-number"]
    assert m["brand"]["slug"] == "stock"
    assert Path(m["input"]["path"]) == (ws / "decks" / "deck.md").resolve()


def test_default_output_uses_folder_name_for_deck_md(ws, capsys):
    deck = ws / "decks" / "q3-review" / "deck.md"
    deck.parent.mkdir()
    deck.write_text(GOOD.lstrip("\n"))
    code, out = cli_json(ws, "build", str(deck), capsys=capsys)
    assert code == 0
    assert Path(out["output"]) == ws / "out" / "q3-review.pptx"


def test_default_output_from_inside_the_deck_folder(ws, capsys, monkeypatch):
    # `build deck.md` run from the deck's own folder, as the bulk-outreach example documents
    deck = ws / "decks" / "q3-review" / "deck.md"
    deck.parent.mkdir()
    deck.write_text(GOOD.lstrip("\n"))
    monkeypatch.chdir(deck.parent)
    code, out = cli_json(ws, "build", "deck.md", capsys=capsys)
    assert code == 0
    assert Path(out["output"]).resolve() == (ws / "out" / "q3-review.pptx").resolve()


def test_inline_markup_becomes_runs(ws, capsys):
    _, out = build(ws, capsys)
    body = Presentation(out["output"]).slides[1].placeholders[1].text_frame
    bold = [r for p in body.paragraphs for r in p.runs if r.font.bold]
    assert [r.text for r in bold] == ["Card stock"]
    assert body.paragraphs[1].level == 1


def test_chart_table_and_pictures_are_native(ws, capsys):
    _, out = build(ws, capsys)
    slides = Presentation(out["output"]).slides
    assert any(sh.has_chart for sh in slides[3].shapes)
    assert any(sh.has_table for sh in slides[4].shapes)
    assert any(sh.shape_type == 13 for sh in slides[5].shapes)  # picture


def test_no_empty_placeholders_left(ws, capsys):
    _, out = build(ws, capsys)
    for slide in Presentation(out["output"]).slides:
        for ph in slide.placeholders:
            assert ph.has_text_frame is False or ph.text_frame.text.strip(), ph.name


def test_rebuild_is_byte_identical(ws, capsys):
    _, first = build(ws, capsys)
    a = Path(first["output"]).read_bytes()
    time.sleep(2.1)  # past the zip timestamp resolution, so a clock leak would show
    _, second = build(ws, capsys)
    assert Path(second["output"]).read_bytes() == a


def test_properties_are_clean_and_stamped(ws, capsys):
    _, out = build(ws, capsys)
    with zipfile.ZipFile(out["output"]) as z:
        core = z.read("docProps/core.xml").decode()
        custom = z.read("docProps/custom.xml").decode()
    assert "python-pptx" not in core and "Steve Canny" not in core
    assert "2000-01-01T00:00:00Z" in core
    assert 'name="deck-builder:brand"' in custom and ">stock<" in custom


def test_manifest_hashes_match_embedded_media(ws, capsys):
    _, out = build(ws, capsys)
    m = json.loads(Path(out["manifest"]).read_text())
    declared = {f["asset"]["sha256"] for s in m["slides"] for f in s["fields"].values() if "asset" in f}
    with zipfile.ZipFile(out["output"]) as z:
        media = {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith("ppt/media/")}
    assert declared and declared == media


def test_small_icon_warns_low_res(ws, capsys):
    code, out = build(ws, capsys)
    assert code == 0
    assert "ASSET_LOW_RES" in codes(out)


def test_build_stops_on_check_errors(ws, capsys):
    code, out = build(ws, capsys, "## A\nlayout: nope\n")
    assert code == 1
    assert "output" not in out


def test_bulk_builds_one_deck_per_row(ws, capsys):
    tmpl = write_deck(ws, "---\nbrand: stock\n---\n\n## {{client}} review\nlayout: title\nsubtitle: For {{contact}}\n",
                      name="pitch.md")
    data = ws / "decks" / "clients.csv"
    data.write_text("client,contact\nHalvorsen Freight,Ops lead\nCopperline Foods,Buyer\n")
    code, out = cli_json(ws, "build", str(tmpl), "--data", str(data), "--name", "{{client}}.pptx", capsys=capsys)
    assert code == 0, out["issues"]
    names = sorted(Path(o["output"]).name for o in out["outputs"])
    assert names == ["Copperline-Foods.pptx", "Halvorsen-Freight.pptx"]
