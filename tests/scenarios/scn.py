"""Shared helpers for the scenario suite.

Every scenario drives the real CLI (`deck_builder.cli.main`) through a tmp workspace, the way a user
or agent would: `init`, a brand step, a deck, then `check`, `build` and a render step. The helpers
below build the binaries each scenario needs (templates, logos, icons) at test time with python-pptx
and Pillow; nothing here is checked into the repo as a new binary file.

Named `scn.py`, not `conftest.py`: a second file named `conftest.py` under tests/ would collide with
tests/conftest.py under plain `from conftest import ...` imports (pytest's no-`__init__.py` import
mode resolves both to the same module name). `ws`, `write_deck`, `cli_json`, `codes` and `make_kit`
still come from tests/conftest.py; `ws` is a fixture and needs no import.
"""
from __future__ import annotations

import json
import textwrap
from pathlib import Path
from typing import Any

import pytest
import yaml
from PIL import Image
from pptx import Presentation

from deck_builder.qa import tools

# ---------------------------------------------------------------- render marker

render = pytest.mark.render
needs_render = pytest.mark.skipif(tools.soffice() is None or bool(tools.poppler_missing()),
                                  reason="needs LibreOffice and poppler")


def json_bytes(payload: dict[str, Any]) -> int:
    """The byte size of a step's --json output, the way an agent receives it: compact, no indent."""
    return len(json.dumps(payload).encode("utf-8"))


def init_root(tmp_path: Path) -> Path:
    """`deck-builder init` into tmp_path; returns tmp_path itself (deck-builder.toml lives there)."""
    from deck_builder.cli import main

    assert main(["init", "--dir", str(tmp_path)]) == 0
    return tmp_path


def write_deck_md(root: Path, rel_name: str, body: str) -> Path:
    """Write a deck.md under <root>/workspace/decks/<rel_name>, creating its folder. rel_name is
    usually '<deck-folder>/deck.md'. body is dedented and leading-newline-stripped."""
    p = root / "workspace" / "decks" / rel_name
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(textwrap.dedent(body).lstrip("\n"), encoding="utf-8")
    return p


# ---------------------------------------------------------------- brand kit sources (not checked in)

DEFAULT_PALETTE = {"primary": "1F3A5F", "accent": "E07A2F", "ink": "1B1B1B", "muted": "6B7280",
                   "surface": "F4F5F7", "background": "FFFFFF"}
DEFAULT_FONTS = {"heading": {"family": "Georgia", "fallback": "Times New Roman"},
                 "body": {"family": "Arial", "fallback": "Liberation Sans"}}

LONG_GUIDE = textwrap.dedent("""\
    # Brand guide

    ## Who we are

    Fictional operations company invented for deck-builder's scenario tests. This guide exists only
    so the fixture folder looks like a real brand hand-off: a designer's notes beside the machine
    files, read by a person, never by the engine.

    ## Voice

    Plain, operational language. Numbers over adjectives. One idea per slide.

    ## Color

    Primary is a deep slate blue, used for titles and section dividers. Accent is a warm orange,
    reserved for callouts, the one chart series that matters, and status marks. Ink is near-black,
    used for body copy only. Surface is a pale warm gray for secondary panels.

    ## Type

    Headings in the brand's serif; body copy in the brand's sans. Never mix a third family in a
    generated deck.

    ## Logo

    Clear space equal to the cap height of the wordmark on every side. Never stretch, recolor outside
    the approved marks, or place on a background under 3:1 contrast.

    ## Icons

    Single-weight line icons, square canvas, transparent background. Default color is the accent.

    ## Photography

    Warm, documentary, never stock-looking. Horizon lines level. No text baked into photography.
    """)

TERSE_GUIDE = "# Brand guide\n\nSlate blue and warm orange. Georgia headings, Arial body. See brand.yaml.\n"


def brand_source(tmp_path: Path, slug: str, *, layout_set: str = "full", with_logo: bool = True,
                 with_icons: bool = True, icon_ids: tuple[str, ...] = ("truck",), extra_icons: int = 0,
                 svg_siblings: bool = False, references_dir: bool = False, guide: str | None = None,
                 slide_numbers: bool | None = None, max_slides: int = 40) -> Path:
    """Write a brand.yaml (and whatever it points at) under tmp_path/<slug>-src; return the brand.yaml path.

    Every knob defaults to the full, realistic shape; scenarios turn pieces off or add to them so the
    kits differ structurally instead of being clones of each other.
    """
    src = tmp_path / f"{slug}-src"
    src.mkdir(parents=True, exist_ok=True)
    meta: dict[str, Any] = {
        "spec_version": 1,
        "name": slug.replace("-", " ").title(),
        "slug": slug,
        "version": "1.0.0",
        "description": "Fictional company invented for deck-builder's scenario tests",
        "palette": DEFAULT_PALETTE,
        "fonts": DEFAULT_FONTS,
        "lint": {"max_slides": max_slides, "banned_patterns": ["—"]},
        "generate": {"slide_size": "16:9", "layout_set": layout_set},
    }
    if with_logo:
        (src / "assets").mkdir(exist_ok=True)
        Image.new("RGB", (600, 200), "#1F3A5F").save(src / "assets" / "logo.png")
        meta["logos"] = {"primary": "assets/logo.png"}
        meta["generate"]["logo_on_master"] = "primary"
    if with_icons:
        icons_dir = src / "assets" / "icons"
        icons_dir.mkdir(parents=True, exist_ok=True)
        for name in icon_ids:
            Image.new("RGBA", (256, 256), (0, 0, 0, 255)).save(icons_dir / f"{name}.png")
        for n in range(extra_icons):
            Image.new("RGBA", (256, 256), (0, 0, 0, 255)).save(icons_dir / f"extra-{n:02d}.png")
        meta["icons"] = {"dir": "assets/icons", "default_color": "accent", "source": "made for tests"}
    if svg_siblings and with_icons:
        svg_dir = src / "assets" / "svg"
        svg_dir.mkdir(parents=True, exist_ok=True)
        for name in icon_ids:
            (svg_dir / f"{name}.svg").write_text(
                '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 10 10">'
                '<rect width="10" height="10"/></svg>', encoding="utf-8")
    if references_dir:
        refs = src / "references"
        refs.mkdir(exist_ok=True)
        for n in range(3):
            Image.new("RGB", (300, 200), "#777777").save(refs / f"moodboard-{n}.png")
    if slide_numbers is not None:
        meta["generate"]["slide_numbers"] = slide_numbers
    if guide is not None:
        (src / "brand-guide.md").write_text(guide, encoding="utf-8")
    brand_path = src / "brand.yaml"
    brand_path.write_text(yaml.safe_dump(meta, sort_keys=False), encoding="utf-8")
    return brand_path


def minimal_brand_source(tmp_path: Path, slug: str) -> Path:
    """palette and fonts only: no logo, no icons, no guide, the smallest layout set."""
    src = tmp_path / f"{slug}-src"
    src.mkdir(parents=True, exist_ok=True)
    meta = {
        "spec_version": 1,
        "name": slug.replace("-", " ").title(),
        "slug": slug,
        "version": "1.0.0",
        "palette": {"primary": "2F3E4D", "accent": "2A7DE1"},
        "fonts": {"heading": {"family": "Arial"}, "body": {"family": "Arial"}},
        "generate": {"slide_size": "16:9", "layout_set": "minimal"},
    }
    brand_path = src / "brand.yaml"
    brand_path.write_text(yaml.safe_dump(meta, sort_keys=False), encoding="utf-8")
    return brand_path


# ---------------------------------------------------------------- client-like templates (not checked in)

RENAMES = {
    "Title Slide": "Cover",
    "Title and Content": "Standard Body",
    "Section Header": "Divider",
    "Picture with Caption": "Photo Panel",
}


def client_template(path: Path) -> Path:
    """A client-style .pptx: python-pptx's stock master, with its layouts renamed the way a design
    team names its own deck's layouts. No slides; a template for `brand adopt` to wrap."""
    prs = Presentation()
    for layout in prs.slide_masters[0].slide_layouts:
        if layout.name in RENAMES:
            layout.name = RENAMES[layout.name]
    prs.save(str(path))
    return path


def write_hero(deck_dir: Path, name: str = "hero.png") -> Path:
    """A deck-local image big enough to clear ASSET_LOW_RES in the generator's full-width image layout."""
    assets = deck_dir / "assets"
    assets.mkdir(parents=True, exist_ok=True)
    path = assets / name
    Image.new("RGB", (2000, 1125), "#334455").save(path)
    return path


def hidden_notes_pptx(path: Path) -> Path:
    """A 2-slide deck made by hand: a normal slide with speaker notes, and a second slide marked
    hidden (as PowerPoint's 'Hide Slide') with its own notes. Both use the stock Title and Content
    layout, so the `stock` kit from tests/conftest.py (STOCK_TOKENS) maps onto it directly."""
    prs = Presentation()
    content = prs.slide_layouts[1]  # Title and Content

    s1 = prs.slides.add_slide(content)
    s1.shapes.title.text = "Volume grew in Q3"
    s1.placeholders[1].text_frame.text = "Copy paper led the growth"
    s1.notes_slide.notes_text_frame.text = "Source: the Q3 shipment ledger."

    s2 = prs.slides.add_slide(content)
    s2.shapes.title.text = "Draft: pricing options"
    s2.placeholders[1].text_frame.text = "Not for this review"
    s2.notes_slide.notes_text_frame.text = "Hidden until finance signs off."
    s2._element.set("show", "0")

    prs.save(str(path))
    return path


def rename_tokens_layouts(kit_dir: Path, renames: dict[str, str]) -> None:
    """Rename top-level layout keys in a kit's tokens.yaml, the way a user tunes a starter file after
    `brand adopt`. renames maps the generated key (slugified template layout name) to the short alias."""
    path = kit_dir / "tokens.yaml"
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    layouts = data["layouts"]
    data["layouts"] = {renames.get(k, k): v for k, v in layouts.items()}
    path.write_text(yaml.safe_dump(data, sort_keys=False, width=110), encoding="utf-8")


# ---------------------------------------------------------------- generated deck.md bodies

FULL_CYCLE = ("title", "agenda", "section", "content", "two-col", "comparison", "big-number", "chart",
             "table", "image", "image-right", "icon-row", "team", "quote", "closing")


def _slide(layout: str, n: int, icon_id: str) -> str:
    if layout == "title":
        return f"## Northfield Ops review, part {n}\nlayout: title\nsubtitle: Volume, margin and the plan\n"
    if layout == "agenda":
        return f"## Agenda {n}\nlayout: agenda\n\n- Volume\n- Margins\n- Delivery\n"
    if layout == "section":
        return f"## Section {n}\nlayout: section\nkicker: Part {n}\n"
    if layout == "content":
        return (f"## Copy paper led the growth, week {n}\nlayout: content\n\n"
                f"- Volume rose {n}% over the prior period\n- Two accounts renewed early\n")
    if layout == "two-col":
        return (f"## Before and after, cycle {n}\nlayout: two-col\n\n### left\n- Manual count\n\n"
                f"### right\n- Scanned count\n")
    if layout == "comparison":
        return (f"## Lease or buy, option {n}\nlayout: comparison\nleft-heading: Lease\nright-heading: Buy\n\n"
                f"### left\n- Lower upfront\n\n### right\n- Lower total\n")
    if layout == "big-number":
        return f"## {n}%\nlayout: big-number\ncaption: Growth this cycle, period {n}\n"
    if layout == "chart":
        return (f"## Cases by month, cycle {n}\nlayout: chart\n\n```chart\ncategories: [Jul, Aug]\n"
                f"series: [{{name: Copy, values: [{n}, {n + 1}]}}]\n```\n")
    if layout == "table":
        return f"## Open items, cycle {n}\nlayout: table\n\n| Item | Owner |\n|---|---|\n| Lease {n} | Ops |\n"
    if layout == "image":
        return (f"## Warehouse, cycle {n}\nlayout: image\n\n### image\n![Warehouse](assets/hero.png)\n\n"
                f"### caption\nSite {n}.\n")
    if layout == "image-right":
        return (f"## Routes, cycle {n}\nlayout: image-right\n\n- North route {n}\n\n"
                f"### image\n![Warehouse](assets/hero.png)\n")
    if layout == "icon-row":
        return f"## How we ship, cycle {n}\nlayout: icon-row\nicon1: brand:icon/{icon_id}\ntext1: Daily trucks\n"
    if layout == "team":
        return f"## Team, cycle {n}\nlayout: team\nname1: Ops lead {n}\n"
    if layout == "quote":
        return (f'## "Paper is still the fastest way to sign a deal, cycle {n}."\nlayout: quote\n'
                f"attribution: A regional buyer\n")
    return f"## Thank you, cycle {n}\nlayout: closing\nsubtitle: ops@example.com\n"  # closing


def gen_deck_md(n_slides: int, brand_slug: str, icon_id: str = "truck", sentinel: str | None = None) -> str:
    """n_slides of deck.md body, cycling through every layout of the generator's `full` set.

    Needs a deck-local assets/hero.png beside deck.md (write_hero) once any cycle reaches image or
    image-right. sentinel, if given, is dropped into one slide's body as a long, distinctive run of
    text: the cost-of-output scenarios check that no --json step echoes it back.
    """
    front = f"---\nbrand: {brand_slug}\n---\n\n"
    slides = [_slide(FULL_CYCLE[i % len(FULL_CYCLE)], i + 1, icon_id) for i in range(n_slides)]
    if sentinel:
        slides[0] = slides[0].rstrip("\n") + f"\n\nNotes:\n{sentinel}\n"
    return front + "\n".join(slides)


# ---------------------------------------------------------------- per-scenario setup (no capsys: plain
# functions, not fixtures, so a module-scoped fixture never has to request the function-scoped capsys)


def setup_adopted_kit(tmp_path: Path) -> tuple[Path, Path]:
    """A client .pptx with renamed layouts, adopted, with its tokens.yaml layout keys renamed again
    the way a user would; a 4-slide deck using the renamed keys. Returns (root, deck)."""
    from deck_builder.cli import main

    root = init_root(tmp_path)
    client = tmp_path / "client.pptx"
    client_template(client)
    assert main(["--config", str(root / "deck-builder.toml"), "brand", "adopt", "acme-client",
                "--template", str(client)]) == 0
    kit = root / "workspace" / "brands" / "acme-client"
    rename_tokens_layouts(kit, {"cover": "title", "standard-body": "content", "divider": "section",
                                "photo-panel": "image"})
    write_hero(root / "workspace" / "decks" / "acme")
    deck = write_deck_md(root, "acme/deck.md", """
        ---
        brand: acme-client
        ---

        ## Acme quarterly review
        layout: title
        subtitle: Q3 volume and the plan for Q4

        ## Volume grew in Q3
        layout: content

        - Up 12% over Q2
        - Two new accounts signed

        ## Part 1: volume
        layout: section

        - What drove the growth

        ## The new warehouse
        layout: image

        ### image
        ![Warehouse](assets/hero.png)

        ### body
        - Opened in August
        """)
    return root, deck


def setup_minimal_kit(tmp_path: Path) -> tuple[Path, Path]:
    """palette-and-fonts-only kit, minimal layout set, a 4-slide deck. Returns (root, deck)."""
    from deck_builder.cli import main

    root = init_root(tmp_path)
    src = minimal_brand_source(tmp_path, "northfield-minimal")
    assert main(["--config", str(root / "deck-builder.toml"), "brand", "init", "northfield-minimal",
                "--from", str(src)]) == 0
    deck = write_deck_md(root, "min/deck.md", """
        ---
        brand: northfield-minimal
        ---

        ## Northfield review
        layout: title
        subtitle: Q3

        ## Part 1
        layout: section
        kicker: Part 1

        ## Volume grew
        layout: content

        - Up 12%

        ## Thank you
        layout: closing
        subtitle: ops@example.com
        """)
    return root, deck


def setup_many_icons_kit(tmp_path: Path) -> tuple[Path, Path]:
    """3 named icons plus 15 more, and a references/ folder the engine never touches. 3-slide deck."""
    from deck_builder.cli import main

    root = init_root(tmp_path)
    src = brand_source(tmp_path, "northfield-signage", layout_set="full",
                       icon_ids=("truck", "box", "dock"), extra_icons=15, references_dir=True)
    assert main(["--config", str(root / "deck-builder.toml"), "brand", "init", "northfield-signage",
                "--from", str(src)]) == 0
    deck = write_deck_md(root, "signage/deck.md", """
        ---
        brand: northfield-signage
        ---

        ## Northfield signage review
        layout: title
        subtitle: Icon inventory

        ## How we ship
        layout: icon-row
        icon1: brand:icon/dock
        text1: Loading dock

        ## Thank you
        layout: closing
        subtitle: ops@example.com
        """)
    return root, deck


def setup_svg_source_kit(tmp_path: Path) -> tuple[Path, Path]:
    """Icons with PNG-plus-SVG source art beside them; a 2-slide deck. Returns (root, deck)."""
    from deck_builder.cli import main

    root = init_root(tmp_path)
    src = brand_source(tmp_path, "riverton-studio", layout_set="standard", svg_siblings=True,
                       icon_ids=("spark",))
    assert main(["--config", str(root / "deck-builder.toml"), "brand", "init", "riverton-studio",
                "--from", str(src)]) == 0
    deck = write_deck_md(root, "studio/deck.md", """
        ---
        brand: riverton-studio
        ---

        ## Riverton Studio review
        layout: title
        subtitle: Icons, PNG and SVG source

        ## Thank you
        layout: closing
        subtitle: ops@example.com
        """)
    return root, deck


def setup_cycle_decks(tmp_path: Path, slug: str = "northfield-cycle",
                      sentinel: str | None = None) -> tuple[Path, Path, Path]:
    """One `full`-layout-set kit, a 3-slide deck and a 40-slide deck, both cycling through every
    generated layout. Returns (root, deck3, deck40). sentinel is embedded in deck40 only."""
    from deck_builder.cli import main

    root = init_root(tmp_path)
    src = brand_source(tmp_path, slug, layout_set="full")
    assert main(["--config", str(root / "deck-builder.toml"), "brand", "init", slug, "--from", str(src)]) == 0
    write_hero(root / "workspace" / "decks" / "small")
    write_hero(root / "workspace" / "decks" / "large")
    deck3 = write_deck_md(root, "small/deck.md", gen_deck_md(3, slug))
    deck40 = write_deck_md(root, "large/deck.md", gen_deck_md(40, slug, sentinel=sentinel))
    return root, deck3, deck40
