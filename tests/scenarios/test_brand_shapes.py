"""Brand kits that differ structurally, each taken through init, a brand step, a deck, check, build
and a render step, the way a user or agent actually works.

Every kit here is generated at test time (python-pptx, Pillow, yaml); nothing is a new binary fixture
file. `tests/fixtures/demo-brands/` still covers the "one full kit, three brand skins" shape; these
scenarios cover kit shapes that fixture set never exercises: an adopted client template, a kit with
nothing but palette and fonts, a kit with far more icons than any demo brand plus an unrelated assets
folder, PNG-only vs. PNG-plus-SVG-source icons, and a long vs. a terse brand guide.

Each scenario is a plain setup function in scn.py (not a fixture): capsys is function-scoped, and a
module-scoped fixture can't request it, so the non-render and render tests below each call setup
themselves. Setup is a few CLI calls on tiny decks, so repeating it costs well under a second.
"""
from __future__ import annotations

import hashlib

import yaml
from pptx import Presentation
from scn import (
    LONG_GUIDE,
    TERSE_GUIDE,
    brand_source,
    init_root,
    json_bytes,
    needs_render,
    render,
    setup_adopted_kit,
    setup_many_icons_kit,
    setup_minimal_kit,
    setup_svg_source_kit,
)

from conftest import cli_json

JSON_BUDGET = 4096  # a tiny deck's check/build --json; generous, but catches a step that echoes content
RENDER_JSON_BUDGET = 8192


# ---------------------------------------------------------------- adopted client kit, layouts renamed


def test_adopted_kit_has_the_clients_own_renamed_layouts(tmp_path):
    root, _ = setup_adopted_kit(tmp_path)
    kit = root / "workspace" / "brands" / "acme-client"
    tokens = yaml.safe_load((kit / "tokens.yaml").read_text())
    assert tokens["layouts"]["title"]["template_layout"] == "Cover"
    assert tokens["layouts"]["content"]["template_layout"] == "Standard Body"
    assert tokens["layouts"]["section"]["template_layout"] == "Divider"
    assert tokens["layouts"]["image"]["template_layout"] == "Photo Panel"


def test_adopted_kit_checks_and_builds_four_slides(tmp_path, capsys):
    root, deck = setup_adopted_kit(tmp_path)
    capsys.readouterr()
    code, out = cli_json(root, "check", str(deck), capsys=capsys)
    assert code == 0, out["issues"]
    assert json_bytes(out) < JSON_BUDGET
    code, build_out = cli_json(root, "build", str(deck), capsys=capsys)
    assert code == 0, build_out["issues"]
    assert json_bytes(build_out) < JSON_BUDGET
    assert Presentation(build_out["output"]).slides[0].shapes.title.text == "Acme quarterly review"


@render
@needs_render
def test_adopted_kit_renders_with_no_errors(tmp_path, capsys):
    root, deck = setup_adopted_kit(tmp_path)
    capsys.readouterr()
    code, out = cli_json(root, "check", str(deck), "--render", capsys=capsys)
    errors = [i for i in out["issues"] if i["severity"] == "error"]
    assert code == 0, errors
    assert out["flagged_slides"] == []
    assert json_bytes(out) < RENDER_JSON_BUDGET


# ---------------------------------------------------------------- minimal kit: palette and fonts only


def test_minimal_kit_has_no_logo_or_icons(tmp_path):
    root, _ = setup_minimal_kit(tmp_path)
    kit = root / "workspace" / "brands" / "northfield-minimal"
    assert not (kit / "assets").exists()
    meta = yaml.safe_load((kit / "brand.yaml").read_text())
    assert "logos" not in meta and "icons" not in meta


def test_minimal_kit_checks_and_builds_four_slides(tmp_path, capsys):
    root, deck = setup_minimal_kit(tmp_path)
    capsys.readouterr()
    code, out = cli_json(root, "check", str(deck), capsys=capsys)
    assert code == 0, out["issues"]
    assert json_bytes(out) < JSON_BUDGET
    code, build_out = cli_json(root, "build", str(deck), capsys=capsys)
    assert code == 0, build_out["issues"]
    assert Presentation(build_out["output"]).slides[0].shapes.title.text == "Northfield review"


@render
@needs_render
def test_minimal_kit_renders_with_no_errors(tmp_path, capsys):
    root, deck = setup_minimal_kit(tmp_path)
    capsys.readouterr()
    code, out = cli_json(root, "check", str(deck), "--render", capsys=capsys)
    errors = [i for i in out["issues"] if i["severity"] == "error"]
    assert code == 0, errors
    assert out["flagged_slides"] == []


# ---------------------------------------------------------------- many icons, plus an unrelated folder


def test_kit_copies_every_icon_but_not_the_references_folder(tmp_path):
    root, _ = setup_many_icons_kit(tmp_path)
    kit = root / "workspace" / "brands" / "northfield-signage"
    icons = sorted(p.name for p in (kit / "assets" / "icons").glob("*.png"))
    assert len(icons) == 18  # 3 named + 15 extra
    assert not (kit / "references").exists()


def test_many_icons_kit_checks_and_builds(tmp_path, capsys):
    root, deck = setup_many_icons_kit(tmp_path)
    capsys.readouterr()
    code, out = cli_json(root, "check", str(deck), capsys=capsys)
    assert code == 0, out["issues"]
    code, build_out = cli_json(root, "build", str(deck), capsys=capsys)
    assert code == 0, build_out["issues"]
    assert build_out["slides"] == 3


@render
@needs_render
def test_many_icons_kit_renders_with_no_errors(tmp_path, capsys):
    root, deck = setup_many_icons_kit(tmp_path)
    capsys.readouterr()
    code, out = cli_json(root, "check", str(deck), "--render", capsys=capsys)
    errors = [i for i in out["issues"] if i["severity"] == "error"]
    assert code == 0, errors
    assert out["flagged_slides"] == []


# ---------------------------------------------------------------- PNG-only vs. PNG-plus-SVG-source icons


def test_svg_source_art_beside_the_icons_never_reaches_the_kit(tmp_path):
    root = init_root(tmp_path)
    png_only = brand_source(tmp_path, "riverton", layout_set="standard", svg_siblings=False)
    with_svg = brand_source(tmp_path, "riverton", layout_set="standard", svg_siblings=True)
    cfg = str(root / "deck-builder.toml")
    from deck_builder.cli import main

    assert main(["--config", cfg, "brand", "init", "riverton", "--from", str(png_only),
                "--out", str(tmp_path / "kit-png")]) == 0
    assert main(["--config", cfg, "brand", "init", "riverton", "--from", str(with_svg),
                "--out", str(tmp_path / "kit-svg")]) == 0
    kit_png, kit_svg = tmp_path / "kit-png" / "riverton", tmp_path / "kit-svg" / "riverton"

    def sha(p):
        return hashlib.sha256(p.read_bytes()).hexdigest()

    assert sha(kit_png / "template.potx") == sha(kit_svg / "template.potx")
    assert sha(kit_png / "tokens.yaml") == sha(kit_svg / "tokens.yaml")
    assert not (kit_svg / "assets" / "svg").exists()  # the engine never copies or reads it


def test_svg_source_kit_checks_and_builds(tmp_path, capsys):
    root, deck = setup_svg_source_kit(tmp_path)
    capsys.readouterr()
    code, out = cli_json(root, "check", str(deck), capsys=capsys)
    assert code == 0, out["issues"]
    code, build_out = cli_json(root, "build", str(deck), capsys=capsys)
    assert code == 0, build_out["issues"]


@render
@needs_render
def test_svg_source_kit_renders_with_no_errors(tmp_path, capsys):
    root, deck = setup_svg_source_kit(tmp_path)
    capsys.readouterr()
    code, out = cli_json(root, "check", str(deck), "--render", capsys=capsys)
    errors = [i for i in out["issues"] if i["severity"] == "error"]
    assert code == 0, errors
    assert out["flagged_slides"] == []


# ---------------------------------------------------------------- long vs. terse brand guide


def test_a_long_or_a_terse_brand_guide_generates_the_identical_kit(tmp_path):
    """brand-guide.md is for people; the engine never reads it. Same brand.yaml, same name and slug,
    a long guide beside one source and a three-line guide beside the other: the generated kit (the
    only thing `brand init` writes) must be byte-identical either way."""
    root = init_root(tmp_path)
    long_src = brand_source(tmp_path, "northfield-guide", layout_set="standard", guide=LONG_GUIDE)
    terse_src = brand_source(tmp_path, "northfield-guide", layout_set="standard", guide=TERSE_GUIDE)
    cfg = str(root / "deck-builder.toml")
    from deck_builder.cli import main

    assert main(["--config", cfg, "brand", "init", "northfield-guide", "--from", str(long_src),
                "--out", str(tmp_path / "kit-long")]) == 0
    assert main(["--config", cfg, "brand", "init", "northfield-guide", "--from", str(terse_src),
                "--out", str(tmp_path / "kit-terse")]) == 0
    a = tmp_path / "kit-long" / "northfield-guide"
    b = tmp_path / "kit-terse" / "northfield-guide"

    def sha(p):
        return hashlib.sha256(p.read_bytes()).hexdigest()

    assert sha(a / "template.potx") == sha(b / "template.potx")
    assert sha(a / "tokens.yaml") == sha(b / "tokens.yaml")
    assert not (a / "brand-guide.md").exists()  # the guide sits beside brand.yaml, not inside the kit
