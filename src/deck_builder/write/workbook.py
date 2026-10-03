"""Deck -> canonical deck.xlsx, set up for a team to edit in Excel, SharePoint or LibreOffice.

Sheets: `slides` (one row per slide), `deck` (front matter), one sheet per chart or table, `README`,
and a hidden `_lists` sheet that feeds the layout dropdown and the live character counts.
"""
from __future__ import annotations

import io
import re
from typing import Any

import yaml
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from deck_builder.brand.registry import Brand
from deck_builder.build.normalize import normalize_workbook
from deck_builder.model import SLIDE_KEYS, Chart, Deck, Table
from deck_builder.write.cells import cell_text

FIXED_DATE = "2000-01-01T00:00:00Z"
SPARE_ROWS = 50  # dropdowns and counters also cover rows a person adds
LISTS = "'_lists'"
README = [
    "How to edit this deck",
    "",
    "One row per slide on the slides sheet. Change text in any cell; add or delete rows to add or remove slides.",
    "layout: pick from the dropdown. The brand's layouts and their fields: deck-builder brand show <slug>.",
    "Bullets: one per line, starting with '- '. Indent two spaces per level.",
    "Images: ![alt text](assets/file.png), or ![alt](brand:logo/<id>). Icons: brand:icon/<id>.",
    "A cell holding sheet:<name> places the chart or table on that sheet.",
    "Charts: option rows, a blank row, then categories down the first column and one series per column.",
    "Columns starting with # are helpers (live character counts against the layout's budget) and are ignored.",
    "The deck sheet holds the deck's settings as YAML values; quotes there are part of the value's syntax.",
    "",
    "Build: deck-builder build deck.xlsx. Back to markdown: deck-builder convert deck.xlsx deck.md.",
    "deck-builder check deck.xlsx is the authority on budgets; the counts here are a guide.",
]


def yaml_scalar(v: Any) -> str:
    return str(yaml.safe_dump(v, default_flow_style=True, width=10_000, allow_unicode=True)).removesuffix(
        "\n...\n").strip()


def _sheet_name(kind: str, n: int, field: str, used: set[str]) -> str:
    base = re.sub(r"[\[\]:*?/\\]", "-", f"{kind}-{n:02d}-{field}")[:31]
    name, k = base, 2
    while name.lower() in used:
        suffix = f"-{k}"
        name, k = base[: 31 - len(suffix)] + suffix, k + 1
    used.add(name.lower())
    return name


def _text(ws: Any, row: int, col: int, value: str) -> Any:
    cell = ws.cell(row=row, column=col, value=value)
    cell.data_type = "s"  # never a formula, even when the text starts with '='
    return cell


def _chart_sheet(ws: Any, c: Chart) -> None:
    opts = [("type", c.type), ("number_format", c.number_format), ("labels", "true" if c.labels else "false"),
            ("legend", "" if c.legend is None else ("true" if c.legend else "false")), ("title", c.title or ""),
            ("colors", ", ".join(c.colors or []))]
    for r, (k, v) in enumerate(opts, start=1):
        _text(ws, r, 1, k).font = Font(bold=True)
        _text(ws, r, 2, v)
    head = len(opts) + 2
    _text(ws, head, 1, "category").font = Font(bold=True)
    for j, s in enumerate(c.series, start=2):
        _text(ws, head, j, s.name).font = Font(bold=True)
    for i, cat in enumerate(c.categories, start=head + 1):
        _text(ws, i, 1, cat)
        for j, s in enumerate(c.series, start=2):
            if i - head - 1 < len(s.values):
                ws.cell(row=i, column=j, value=s.values[i - head - 1])
    ws.column_dimensions["A"].width = 18


def _table_sheet(ws: Any, t: Table) -> None:
    for j, h in enumerate(t.header, start=1):
        _text(ws, 1, j, h).font = Font(bold=True)
    for i, row in enumerate(t.rows, start=2):
        for j, v in enumerate(row, start=1):
            _text(ws, i, j, v)


def _budgets(brand: Brand | None) -> dict[str, dict[str, int]]:
    """field -> {layout: max_chars}, for every budgeted text field in the brand."""
    out: dict[str, dict[str, int]] = {}
    if brand is None:
        return out
    for lname, ls in (brand.tokens.get("layouts") or {}).items():
        for fname, fs in (ls.get("fields") or {}).items():
            if fs.get("max_chars"):
                out.setdefault(fname, {})[lname] = int(fs["max_chars"])
    out.setdefault("title", {})
    for lname, ls in (brand.tokens.get("layouts") or {}).items():
        hf = ls.get("heading_field", "title")
        if hf in ls.get("fields", {}) and ls["fields"][hf].get("max_chars"):
            out["title"][lname] = int(ls["fields"][hf]["max_chars"])
    return {k: v for k, v in out.items() if v}


def write(deck: Deck, brand: Brand | None) -> bytes:
    """Deck -> canonical deck.xlsx; with a brand, layout dropdowns and live character counts for the team."""
    wb = Workbook()
    ws = wb.active
    ws.title = "slides"
    fields: list[str] = []
    for s in deck.slides:
        fields += [k for k in s.fields if k not in fields]
    budgets = {f: b for f, b in _budgets(brand).items() if f in fields or f == "title"}
    headers = ["slide", "layout", "title"]
    if "title" in budgets:
        headers.append("#chars:title")
    headers += [k for k in SLIDE_KEYS if any(getattr(s, k) is not None for s in deck.slides)]
    for f in fields:
        headers.append(f)
        if f in budgets:
            headers.append(f"#chars:{f}")
    headers.append("notes")
    col = {h: i for i, h in enumerate(headers, start=1)}
    for h, i in col.items():
        cell = _text(ws, 1, i, h)
        cell.font = Font(bold=True, color="6B7280" if h.startswith("#") else None)
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = (8 if h in ("slide", *SLIDE_KEYS) else 16 if h == "layout"
                                              else 12 if h.startswith("#") else 40)
    ws.freeze_panes = "D2"

    used = {"slides", "deck", "readme", "_lists"}
    wrap = Alignment(wrap_text=True, vertical="top")
    for n, s in enumerate(deck.slides, start=1):
        r = n + 1
        ws.cell(row=r, column=col["slide"], value=n)
        _text(ws, r, col["layout"], s.layout)
        _text(ws, r, col["title"], s.title).alignment = wrap
        for k in SLIDE_KEYS:
            if getattr(s, k) is not None:
                ws.cell(row=r, column=col[k], value=getattr(s, k))
        for f in fields:
            if f not in s.fields:
                continue
            v = s.fields[f]
            if isinstance(v, Chart | Table):
                kind = "chart" if isinstance(v, Chart) else "table"
                name = _sheet_name(kind, n, f, used)
                vs = wb.create_sheet(name)
                _chart_sheet(vs, v) if isinstance(v, Chart) else _table_sheet(vs, v)
                _text(ws, r, col[f], f"sheet:{name}")
            else:
                _text(ws, r, col[f], cell_text(v)).alignment = wrap
        if s.notes:
            _text(ws, r, col["notes"], s.notes).alignment = wrap

    if brand is not None:
        _team_features(wb, ws, brand, budgets, col, len(deck.slides) + 1 + SPARE_ROWS)

    meta_ws = wb.create_sheet("deck", 1)
    _text(meta_ws, 1, 1, "key").font = Font(bold=True)
    _text(meta_ws, 1, 2, "value").font = Font(bold=True)
    for i, (k, v) in enumerate(deck.meta.items(), start=2):
        _text(meta_ws, i, 1, str(k))
        _text(meta_ws, i, 2, yaml_scalar(v))
    meta_ws.column_dimensions["A"].width = 18
    meta_ws.column_dimensions["B"].width = 60

    readme = wb.create_sheet("README")
    for i, line in enumerate(README, start=1):
        _text(readme, i, 1, line).font = Font(bold=(i == 1))
    readme.column_dimensions["A"].width = 120

    buf = io.BytesIO()
    wb.save(buf)
    return normalize_workbook(buf.getvalue(), FIXED_DATE)


def _team_features(wb: Any, ws: Any, brand: Brand, budgets: dict[str, dict[str, int]], col: dict[str, int],
                   last_row: int) -> None:
    lists = wb.create_sheet("_lists")
    lists.sheet_state = "hidden"
    layouts = brand.layout_names()
    for i, name in enumerate(layouts, start=1):
        _text(lists, i, 1, name)
    r = 1
    for field, per_layout in budgets.items():
        for lname, limit in per_layout.items():
            _text(lists, r, 3, f"{lname}|{field}")
            lists.cell(row=r, column=4, value=limit)
            r += 1
    dv = DataValidation(type="list", formula1=f"={LISTS}!$A$1:$A${max(1, len(layouts))}", allow_blank=False)
    dv.error = "Pick a layout from the list (deck-builder brand show <slug>)."
    dv.errorTitle = "Unknown layout"
    ws.add_data_validation(dv)
    lcol = get_column_letter(col["layout"])
    dv.add(f"{lcol}2:{lcol}{last_row}")
    red = PatternFill(start_color="FCA5A5", end_color="FCA5A5", fill_type="solid")
    for field in budgets:
        fcol = get_column_letter(col[field])
        hcol = get_column_letter(col[f"#chars:{field}"])
        look = f'VLOOKUP(${lcol}{{r}}&"|{field}",{LISTS}!$C:$D,2,FALSE)'
        for row in range(2, last_row + 1):
            lk = look.format(r=row)
            cell = ws.cell(row=row, column=col[f"#chars:{field}"],
                           value=f'=IF({fcol}{row}="","",IFERROR(LEN({fcol}{row})&"/"&{lk},LEN({fcol}{row})))')
            cell.font = Font(color="6B7280")
        first = look.format(r=2)
        ws.conditional_formatting.add(
            f"{fcol}2:{fcol}{last_row}",
            FormulaRule(formula=[f"AND(ISNUMBER({first}),LEN({fcol}2)>{first})"], fill=red))
        ws.conditional_formatting.add(
            f"{hcol}2:{hcol}{last_row}",
            FormulaRule(formula=[f"AND(ISNUMBER({first}),LEN({fcol}2)>{first})"], fill=red))
