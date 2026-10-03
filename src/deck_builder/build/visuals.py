"""Charts, tables and pictures into placeholders. Colors and sizes come from tokens.yaml."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from PIL import Image as PILImage
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from deck_builder.brand.kit import contrast
from deck_builder.brand.registry import Brand
from deck_builder.build.text import CODE_FONT_DEFAULT, add_runs
from deck_builder.model import Chart, Table

CHART_TYPES = {
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "stacked-column": XL_CHART_TYPE.COLUMN_STACKED,
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "stacked-bar": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
}
# Documented defaults, used only when tokens.yaml doesn't set the value.
CHART_DEFAULTS = {"font_size": 12, "gap_width": 70, "line_width": 2.75, "gridlines": True}
HAIRLINE = Pt(0.5)  # gridlines
BASELINE = Pt(0.75)  # the category axis line, and the line between stacked segments
CLUSTER_OVERLAP = -8  # bars in one group stand a hair apart, as a share of a bar's width
THIN_SEGMENT = 0.06  # a stacked segment under this share of the tallest stack shows no label
TABLE_DEFAULTS = {"font_size": 14, "header_font_size": 14, "row_height_factor": 2.2}


def _rgb(brand: Brand, ref: Any) -> RGBColor | None:
    hexv = brand.color(str(ref)) if ref is not None else None
    return RGBColor.from_string(hexv) if hexv else None


def _rules(cell: Any, rule: RGBColor, row: int) -> None:
    """Thin rules between body rows and no other borders, replacing the table style's white grid lines.

    Body rows get the rule on both edges, since some renderers draw a shared edge from the lower cell's
    top border. The header has none: its fill marks it.
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    tags = ("a:lnL", "a:lnR", "a:lnT", "a:lnB")
    for tag in tags:
        for old in tc_pr.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{tag[2:]}"):
            tc_pr.remove(old)
    ruled = {"a:lnT": row >= 2, "a:lnB": row >= 1}
    for i, tag in enumerate(tags):  # borders come first in tcPr, before the fill
        ln = OxmlElement(tag)
        ln.set("w", "9525")  # 0.75 pt
        if ruled.get(tag):
            fill = OxmlElement("a:solidFill")
            clr = OxmlElement("a:srgbClr")
            clr.set("val", str(rule))
            fill.append(clr)
            ln.append(fill)
        else:
            ln.append(OxmlElement("a:noFill"))
        tc_pr.insert(i, ln)


def _take_geometry(ph: Any) -> tuple[int, int, int, int]:
    geo = (ph.left, ph.top, ph.width, ph.height)
    ph._element.getparent().remove(ph._element)
    return geo


def _alt(frame: Any, text: str) -> None:
    """Alt text on a chart or table, for screen readers; images get theirs from the deck."""
    frame._element.nvGraphicFramePr.cNvPr.set("descr", text[:1].upper() + text[1:])


def _square_corners(chart: Any) -> None:
    """PowerPoint draws a chart with rounded corners when c:roundedCorners is missing, so write it as off.

    It belongs after c:date1904 and c:lang and before everything else in c:chartSpace.
    """
    cs = chart._chartSpace
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
    if cs.find(f"{ns}roundedCorners") is not None:
        return
    rc = OxmlElement("c:roundedCorners")
    rc.set("val", "0")
    lead = [el for el in cs if el.tag in (f"{ns}date1904", f"{ns}lang")]
    cs.insert(cs.index(lead[-1]) + 1 if lead else 0, rc)


def fill_chart(slide: Any, ph: Any, spec: Chart, brand: Brand) -> None:
    """A native, editable chart in the placeholder, styled from tokens.yaml chart, with alt text.

    The styling is quiet: hairline gridlines, no tick marks, axis labels in the muted axis color, a thin
    baseline, lines without markers, and a hairline of the background between stacked segments. A
    clustered bar or column chart with data labels drops its value axis and gridlines, since each bar
    shows its own number. Horizontal bars list their categories top to bottom in the order written.
    """
    ct = CHART_TYPES[spec.type]
    data = CategoryChartData(number_format=spec.number_format)
    data.categories = spec.categories
    for sr in spec.series:
        data.add_series(sr.name, sr.values)
    if hasattr(ph, "insert_chart"):
        frame = ph.insert_chart(ct, data)
    else:
        frame = slide.shapes.add_chart(ct, *_take_geometry(ph), data)
    chart = frame.chart
    _alt(frame, f"{spec.type.replace('-', ' ')} chart{': ' + spec.title if spec.title else ''} of "
                f"{', '.join(s.name for s in spec.series)} by {', '.join(map(str, spec.categories))}")

    tok = {**CHART_DEFAULTS, **(brand.tokens.get("chart") or {})}
    names = spec.colors or tok.get("colors") or []
    colors = [c for c in (_rgb(brand, n) for n in names) if c is not None]
    if tok.get("font"):
        chart.font.name = tok["font"]
    chart.font.size = Pt(tok["font_size"])
    text_color = _rgb(brand, tok.get("text_color"))
    if text_color:
        chart.font.color.rgb = text_color
    pie = spec.type in ("pie", "doughnut")
    stacked = spec.type.startswith("stacked-")
    plot = chart.plots[0]
    if colors:
        if pie:
            for i, pt in enumerate(plot.series[0].points):
                pt.format.fill.solid()
                pt.format.fill.fore_color.rgb = colors[i % len(colors)]
        else:
            for i, s in enumerate(plot.series):
                c = colors[i % len(colors)]
                if spec.type == "line":
                    s.format.line.color.rgb = c
                else:
                    s.format.fill.solid()
                    s.format.fill.fore_color.rgb = c
    if spec.type == "line":
        for s in plot.series:
            s.format.line.width = Pt(tok["line_width"])
            s.marker.style = XL_MARKER_STYLE.NONE
            s.smooth = False
    if stacked:  # a hairline of the background between segments, so each reads as its own block
        paper = _rgb(brand, "background") or RGBColor.from_string("FFFFFF")
        for s in plot.series:
            s.format.line.color.rgb = paper
            s.format.line.width = BASELINE
    if not pie and hasattr(plot, "gap_width"):
        plot.gap_width = tok["gap_width"]
        if not stacked and len(spec.series) > 1:
            plot.overlap = CLUSTER_OVERLAP
        for s in plot.series:  # without this, LibreOffice draws a negative bar as positive
            s.invert_if_negative = False
    if not pie:
        plot.vary_by_categories = False  # one color per series, never a rainbow across one series
    _square_corners(chart)
    values = [v for s in spec.series for v in s.values if v is not None]
    if spec.type in ("column", "stacked-column", "bar", "stacked-bar") and values and min(values) >= 0:
        chart.value_axis.minimum_scale = 0  # bars start at zero; a cut axis exaggerates differences
    if not pie and values and min(values) < 0:  # category names at the edge, not under the negative bars
        chart.category_axis.tick_label_position = XL_TICK_LABEL_POSITION.LOW
    legend = spec.legend if spec.legend is not None else (pie or len(spec.series) > 1)
    chart.has_legend = bool(legend)
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
    if spec.labels:
        plot.has_data_labels = True
        plot.data_labels.show_value = True  # the doughnut template's own labels have it off
        plot.data_labels.number_format = spec.number_format
        plot.data_labels.number_format_is_linked = False
        ink = str(text_color or RGBColor(0, 0, 0))

        def readable(fill: Any) -> RGBColor:  # white or ink, whichever reads on this fill
            return RGBColor.from_string("FFFFFF" if contrast(str(fill), "FFFFFF") >= contrast(str(fill), ink) else ink)

        if stacked and colors:  # labels sit on the bars
            for i, s in enumerate(plot.series):
                dl = s.data_labels
                dl.show_value = True
                dl.number_format = spec.number_format
                dl.number_format_is_linked = False
                dl.font.color.rgb = readable(colors[i % len(colors)])
        if stacked:  # a segment too thin to hold its number keeps none, rather than a number over the edges
            stacks = [sum(abs(v) for v in vals if isinstance(v, int | float))
                      for vals in zip(*(sr.values for sr in spec.series), strict=False)]
            tallest = max(stacks, default=0)
            for s, sr in zip(plot.series, spec.series, strict=False):
                dlbls = s._element.find(qn("c:dLbls"))
                thin = [j for j, v in enumerate(sr.values)
                        if isinstance(v, int | float) and tallest and abs(v) < THIN_SEGMENT * tallest]
                for k, j in enumerate(thin if dlbls is not None else []):
                    lbl = OxmlElement("c:dLbl")  # one point's label, deleted; these lead c:dLbls, in idx order
                    for tag, val in (("c:idx", str(j)), ("c:delete", "1")):
                        el = OxmlElement(tag)
                        el.set("val", val)
                        lbl.append(el)
                    dlbls.insert(k, lbl)
        if pie and colors:  # labels sit on the slices
            for i, pt in enumerate(plot.series[0].points):
                pt.data_label.font.color.rgb = readable(colors[i % len(colors)])
    chart.has_title = bool(spec.title)
    if spec.title:
        chart.chart_title.text_frame.text = spec.title
    axis_text = _rgb(brand, tok.get("axis_text_color")) or text_color
    if legend and axis_text:
        chart.legend.font.color.rgb = axis_text
    if not pie:
        axis, cat = chart.value_axis, chart.category_axis
        # Each bar shows its own number, so the scale and its gridlines would only repeat it.
        labeled = spec.labels and not stacked and spec.type != "line"
        axis.has_major_gridlines = bool(tok["gridlines"]) and not labeled
        grid = _rgb(brand, tok.get("gridline_color"))
        if axis.has_major_gridlines:
            axis.major_gridlines.format.line.width = HAIRLINE
            if grid:
                axis.major_gridlines.format.line.color.rgb = grid
        axis.format.line.fill.background()
        axis.tick_labels.number_format = spec.number_format
        axis.tick_labels.number_format_is_linked = False
        if labeled:
            axis.tick_label_position = XL_TICK_LABEL_POSITION.NONE
        for ax in (axis, cat):
            ax.major_tick_mark = XL_TICK_MARK.NONE
            ax.minor_tick_mark = XL_TICK_MARK.NONE
            if axis_text:
                ax.tick_labels.font.color.rgb = axis_text
        cat.format.line.width = BASELINE
        if axis_text:
            cat.format.line.color.rgb = axis_text
        if spec.type in ("bar", "stacked-bar"):  # top to bottom in the order written, scale still below
            cat.reverse_order = True
            crosses = axis._element.find(qn("c:crosses"))
            if crosses is not None:
                crosses.set("val", "max")


NUMERIC = re.compile(r"^[+\-−]?[$€£]?\d[\d,.]*\s*(%|pts?|x)?$")


def numeric_columns(spec: Table) -> list[bool]:
    """Columns whose every filled body cell is a number, an amount, a percentage or a change in points."""
    from deck_builder.validate import plain  # validate imports this module's table defaults

    out = []
    for c in range(len(spec.header)):
        cells = [plain(row[c]).strip() for row in spec.rows if plain(row[c]).strip()]
        out.append(bool(cells) and all(NUMERIC.match(v) for v in cells))
    return out


CHAR_EM = 0.55  # an average character's width in ems, for sizing table columns before rendering
CELL_PAD_EMU = 2 * 91440 + 45720  # the cell's left and right insets, and a little air


def column_widths(spec: Table, total: int, font_size: float = 14) -> list[int]:
    """Column widths in EMU that follow content, so words never break mid-word and long names wrap less.

    Each column first gets room for its longest single word at the table's size (the bold header counted
    a little wider, a status dot included), so "November" never splits. What's left goes by content length:
    a column's weight is its longest cell in characters, capped at 33 so one long cell can't squeeze the
    others. When even the longest words don't fit, every column shrinks in proportion. The last column
    takes the rounding remainder, so the widths sum to total.
    """
    def word_emu(chars: float) -> int:
        return int(chars * font_size * CHAR_EM * 12700) + CELL_PAD_EMU

    ncols = len(spec.header)
    mins = []
    for c in range(ncols):
        longest = max([len(w) * 1.1 for w in spec.header[c].split()] +
                      [len(w) + 2 for row in spec.rows for w in row[c].split()] + [1.0])
        mins.append(word_emu(longest))
    if sum(mins) >= total:
        widths = [total * m // sum(mins) for m in mins]
    else:
        weights = [min(33, max(6, round(len(spec.header[c]) * 1.15), *(len(row[c]) for row in spec.rows)))
                   for c in range(ncols)]
        spare = total - sum(mins)
        widths = [m + spare * w // sum(weights) for m, w in zip(mins, weights, strict=True)]
    widths[-1] = total - sum(widths[:-1])
    return widths


def fill_table(slide: Any, ph: Any, spec: Table, brand: Brand) -> None:
    """A native table in the placeholder: fills, fonts and status dots from tokens.yaml table, with alt text."""
    nrows, ncols = len(spec.rows) + 1, len(spec.header)
    if hasattr(ph, "insert_table"):
        frame = ph.insert_table(nrows, ncols)
    else:
        frame = slide.shapes.add_table(nrows, ncols, *_take_geometry(ph))
    _alt(frame, f"table with columns {', '.join(spec.header)}, {len(spec.rows)} row{'s' * (len(spec.rows) != 1)}")
    table = frame.table
    table.horz_banding = False  # only the token fills band rows; the default style's bands don't
    tok = {**TABLE_DEFAULTS, **(brand.tokens.get("table") or {})}
    code_font = (brand.tokens.get("text") or {}).get("code_font", CODE_FONT_DEFAULT)
    row_h = Pt(tok.get("row_height", max(tok["font_size"], tok["header_font_size"]) * tok["row_height_factor"]))
    for r in range(nrows):
        table.rows[r].height = row_h
    frame.height = row_h * nrows
    for c, width in enumerate(column_widths(spec, frame.width, max(tok["font_size"], tok["header_font_size"]))):
        table.columns[c].width = width
    numeric = numeric_columns(spec)
    header_text, body_text = _rgb(brand, tok.get("header_text")), _rgb(brand, tok.get("text"))
    fills = {"header": _rgb(brand, tok.get("header_fill")), "row": _rgb(brand, tok.get("row_fill")),
             "band": _rgb(brand, tok.get("band_fill"))}
    status = {str(k).strip().casefold(): _rgb(brand, v) for k, v in (tok.get("status") or {}).items()}
    rule = _rgb(brand, tok.get("rule"))
    for r in range(nrows):
        for c in range(ncols):
            cell = table.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text = ""
            value = spec.header[c] if r == 0 else spec.rows[r - 1][c]
            dot_color = status.get(value.strip().casefold()) if r else None
            if dot_color is not None:
                cell.text_frame.paragraphs[0].add_run().text = "● "
            add_runs(cell.text_frame.paragraphs[0], value, code_font)
            for p in cell.text_frame.paragraphs:
                if numeric[c]:
                    p.alignment = PP_ALIGN.RIGHT  # figures line up by place value
                for run in p.runs:
                    if tok.get("font"):
                        run.font.name = tok["font"]
                    run.font.size = Pt(tok["header_font_size"] if r == 0 else tok["font_size"])
                    if r == 0:
                        run.font.bold = True
                    color = header_text if r == 0 else body_text
                    if color:
                        run.font.color.rgb = color
            if dot_color is not None:
                cell.text_frame.paragraphs[0].runs[0].font.color.rgb = dot_color
            fill = fills["header"] if r == 0 else (fills["band"] if r % 2 == 0 else fills["row"])
            if fill:
                cell.fill.solid()
                cell.fill.fore_color.rgb = fill
            elif r:
                cell.fill.background()  # no fill, so the default style's accent tint doesn't show through
            if rule is not None:
                _rules(cell, rule, r)


def color_bold(ph: Any, hex_color: str) -> None:
    """**bold** runs in a filled text field take this color (a field's tokens.yaml emphasis), so the
    keywords on cards and bands read in the primary color."""
    for p in ph.text_frame.paragraphs:
        for run in p.runs:
            if run.font.bold:
                run.font.color.rgb = RGBColor.from_string(hex_color)


def _visible(path: Path) -> tuple[int, int, int, int, int, int]:
    """(image width, height, then the left, top, right, bottom of its visible pixels): a logo drawn on a
    larger transparent canvas centers by its art, not by its canvas."""
    with PILImage.open(path) as im:
        iw, ih = im.size
        box = im.convert("RGBA").getchannel("A").getbbox() if im.mode in ("RGBA", "LA", "PA", "P") else None
    left, top, right, bottom = box or (0, 0, iw, ih)
    return iw, ih, left, top, right, bottom


def fill_picture(slide: Any, ph: Any, path: Path, alt: str, crop: bool, share: float = 1.0,
                 trim: bool = False) -> tuple[tuple[int, int, int, int], Any]:
    """Crop to fill (photos in picture placeholders) or fit inside (logos, icons, other placeholders).
    A fitted image takes at most `share` of the box in each direction, centered, so a logo stays modest.
    With trim, a fitted image's fully transparent margins are cropped away first, so a logo centers by
    its art; icons don't trim, since an icon set's shared canvas keeps its icons the same size.

    Returns the placeholder's geometry, used for the low-resolution check, and the picture shape.
    """
    geo = (ph.left, ph.top, ph.width, ph.height)
    if crop and hasattr(ph, "insert_picture"):
        pic = ph.insert_picture(str(path))
    else:
        left, top, w, h = _take_geometry(ph)
        if trim:
            iw, ih, vl, vt, vr, vb = _visible(path)
        else:
            with PILImage.open(path) as im:
                iw, ih = im.size
            vl, vt, vr, vb = 0, 0, iw, ih
        scale = min(w / (vr - vl), h / (vb - vt)) * share
        nw, nh = int((vr - vl) * scale), int((vb - vt) * scale)
        pic = slide.shapes.add_picture(str(path), left + (w - nw) // 2, top + (h - nh) // 2, nw, nh)
        if (vl, vt, vr, vb) != (0, 0, iw, ih):  # show only the art; the file itself is untouched
            pic.crop_left, pic.crop_top = vl / iw, vt / ih
            pic.crop_right, pic.crop_bottom = (iw - vr) / iw, (ih - vb) / ih
    if alt:
        pic._element.nvPicPr.cNvPr.set("descr", alt)
    return geo, pic
