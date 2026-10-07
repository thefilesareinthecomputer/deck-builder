"""Build a resolved deck into a PPTX and its manifest."""
from __future__ import annotations

import copy
import datetime as dt
import hashlib
import io
import json
import re
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from PIL import Image as PILImage
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

from deck_builder import __version__, confine, images
from deck_builder import template as tpl
from deck_builder.assets import recolor_icon, sha256_file
from deck_builder.brand.registry import Brand
from deck_builder.build.motion import Step, add_fade
from deck_builder.build.normalize import content_digest, normalize
from deck_builder.build.text import CODE_FONT_DEFAULT, fill_text
from deck_builder.build.visuals import color_bold, fill_chart, fill_code, fill_picture, fill_table, mark_current
from deck_builder.errors import EnvError, Issue
from deck_builder.model import Chart, Code, Deck, Icon, Image, Slide, Table
from deck_builder.validate import asset_source, plain, resolve_asset, text_of

EMU_PER_INCH = 914400
MIN_DPI = 150
DEFAULT_DATE = dt.date(2000, 1, 1)
SLIDENUM_FIELD = "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"  # used when the layout's field has no usable id
FIELD_ID = re.compile(r"\{[0-9A-Fa-f]{8}(-[0-9A-Fa-f]{4}){3}-[0-9A-Fa-f]{12}\}")
OFF = {"false", "no", "off", "0"}


def furniture_settings(meta: dict[str, Any]) -> tuple[bool, str | None]:
    """(slide numbers on, footer text) from front matter: `slide_numbers: false` and `footer: <text>`."""
    numbers = str(meta.get("slide_numbers", True)).strip().lower() not in OFF
    footer = str(meta["footer"]).strip() if meta.get("footer") not in (None, "") else None
    return numbers, footer or None


def first_number(meta: dict[str, Any]) -> int:
    """`first_slide_number:` in front matter, for a deck that is an excerpt of a larger one (default 1).
    check refuses anything but a whole number from 1 up, so a bad value never reaches the build."""
    raw = meta.get("first_slide_number", 1)
    return raw if isinstance(raw, int) and not isinstance(raw, bool) and raw >= 1 else 1


def add_furniture(slide: Any, layout: Any, n: int, numbers: bool, footer: str | None, fixed: bool = False) -> None:
    """Copy the layout's slide-number placeholder (and its footer when there's text for it) onto the slide.

    python-pptx never copies these. The whole placeholder is copied, so the template's position and
    styling stay as they are (LibreOffice doesn't inherit them), and only its text is replaced. Dates
    are never copied. A fixed number (an excerpt's, from first_slide_number) is plain text, since
    LibreOffice recounts a slide-number field from 1 whatever the presentation's first number says.
    """
    tree = slide.shapes._spTree
    ids = [str(e.get("id", "")) for e in tree.iter(qn("p:cNvPr"))]
    next_id = max((int(i) for i in ids if i.isdecimal()), default=1) + 1
    for ph in layout.placeholders:
        kind = ph.placeholder_format.type
        if kind == PP_PLACEHOLDER.SLIDE_NUMBER and numbers and fixed:
            para = f'<a:r><a:rPr lang="en-US"/><a:t>{n}</a:t></a:r>'
        elif kind == PP_PLACEHOLDER.SLIDE_NUMBER and numbers:
            fld = ph._element.find(f".//{qn('a:fld')}[@type='slidenum']")
            fid = str(fld.get("id")) if fld is not None else ""
            fid = fid if FIELD_ID.fullmatch(fid) else SLIDENUM_FIELD  # an adopted template is untrusted XML
            para = f'<a:fld id="{fid}" type="slidenum"><a:rPr lang="en-US"/><a:t>{n}</a:t></a:fld>'
        elif kind == PP_PLACEHOLDER.FOOTER and footer:
            para = f'<a:r><a:rPr lang="en-US"/><a:t>{escape(footer)}</a:t></a:r>'
        else:
            continue
        sp = copy.deepcopy(ph._element)
        sp.find(f"{qn('p:nvSpPr')}/{qn('p:cNvPr')}").set("id", str(next_id))
        next_id += 1
        body = sp.find(qn("p:txBody"))
        for p in body.findall(qn("a:p")):
            body.remove(p)
        body.append(parse_xml(f'<a:p {nsdecls("a")}>{para}</a:p>'))
        tree.append(sp)


def deck_date(meta: dict[str, Any]) -> dt.datetime:
    raw = meta.get("date", DEFAULT_DATE)
    if isinstance(raw, dt.datetime):
        d = raw.date()
    elif isinstance(raw, dt.date):
        d = raw
    else:
        d = dt.date.fromisoformat(str(raw))
    return dt.datetime(d.year, d.month, d.day, tzinfo=dt.UTC)


def _dpi(path: Path, geo: tuple[int, int, int, int], crop: bool) -> float:
    with PILImage.open(path) as im:
        iw, ih = im.size
    _, _, w, h = geo
    scale = max(w / iw, h / ih) if crop else min(w / iw, h / ih)  # EMU per source pixel
    return EMU_PER_INCH / scale if scale else float("inf")


SLOT = re.compile(r"\D+(\d+)$")  # logo3 and caption3 are slot 3 of a row layout
ROW_GROW = 2.0  # with fewer slots filled, each grows up to this much wider, so two logos aren't lost in a row


def lay_out_row(fields: dict[str, Any], filled: set[str], phs: dict[int, Any]) -> None:
    """Give a row layout's filled slots the whole row before anything is placed in them.

    Slots are the numbered field groups (logo1 and caption1 are slot 1). With k slots filled out of n, each
    takes 1/k of the row, centered in it, and its boxes widen by up to ROW_GROW, keeping their tops and
    heights. Every coordinate is set, since a placeholder that inherits its position loses the rest when
    only its left is written.
    """
    groups: dict[str, list[Any]] = {}
    for name, f in fields.items():
        m, ph = SLOT.match(name), phs.get(f["idx"])
        if m and ph is not None:
            groups.setdefault(m.group(1), []).append((name, ph))
    if not groups:
        return
    edges = {k: (min(ph.left for _, ph in g), max(ph.left + ph.width for _, ph in g)) for k, g in groups.items()}
    row_left, row_right = min(a for a, _ in edges.values()), max(b for _, b in edges.values())
    used = sorted((k for k, g in groups.items() if any(name in filled for name, _ in g)), key=int)
    if not used:
        return
    pitch, old_pitch = (row_right - row_left) / len(used), (row_right - row_left) / len(groups)
    grow = min(pitch / old_pitch, ROW_GROW)
    for j, k in enumerate(used):
        old_center, new_center = sum(edges[k]) / 2, row_left + pitch * (j + 0.5)
        for _, ph in groups[k]:
            left, top, width, height = ph.left, ph.top, ph.width, ph.height
            ph.left = int(new_center + (left - old_center) * grow)
            ph.top, ph.width, ph.height = top, int(width * grow), height


def build_steps(build: str, shapes: dict[str, Any], fields: dict[str, Any]) -> list[Step]:
    """What appears on each click: `slots` brings in one numbered group (a card, step or band with all its
    fields) at a time; a list field one item at a time; any other field as a whole."""
    if build == "slots":
        groups: dict[int, Step] = {}
        for name in fields:  # the layout's field order, so a slot's parts keep their reading order
            m = SLOT.match(name)
            if m and name in shapes:
                groups.setdefault(int(m.group(1)), []).append((shapes[name]._element, None))
        return [groups[k] for k in sorted(groups)]
    el = shapes[build]._element
    if el.tag == qn("p:sp") and len(paras := el.findall(f".//{qn('a:p')}")) > 1:
        return [[(el, i)] for i in range(len(paras))]
    return [[(el, None)]]


class Builder:
    """Adds a validated deck's slides to a presentation one at a time, filling each layout's placeholders
    from the slide's fields and collecting build-time issues (such as ASSET_LOW_RES) in `issues`."""

    def __init__(self, brand: Brand, deck_dir: Path, cache_dir: Path, numbers: bool = True,
                 footer: str | None = None, first: int = 1) -> None:
        self.brand = brand
        self.deck_dir = deck_dir
        self.cache_dir = cache_dir
        self.numbers, self.footer, self.first = numbers, footer, first
        self.code_font = (brand.tokens.get("text") or {}).get("code_font", CODE_FONT_DEFAULT)
        self.issues: list[Issue] = []

    def _picture(self, slide: Any, ph: Any, ref: str, alt: str, at: dict[str, Any],
                 fentry: dict[str, Any], color: str | None, fs: dict[str, Any], s: Slide) -> Any:
        path, _ = resolve_asset(ref, self.brand, self.deck_dir)
        assert path is not None  # validation guarantees it resolves
        source = asset_source(path, ref, self.brand, self.deck_dir)  # validation confines every asset
        if color:
            path = recolor_icon(path, color, self.cache_dir)
        # Photos fill their box. Contained images (screenshots, diagrams, `fit: contain`) are scaled to fit
        # inside it whole: ratio and edges exactly as the file has them, transparent margins included, such
        # as a window capture's shadow. Logos, icons and logo slots fit inside it too; a logo is trimmed to
        # its art and takes at most its share of a large box, and icons keep their set's canvas.
        crop = images.crops(s, ref, fs, path)
        contained = not color and not ref.startswith("brand:") and not fs.get("fit") and \
            images.fit_mode(s, ref) == "contain"
        logo = not crop and not contained
        geo, pic = fill_picture(slide, ph, path, alt, crop=crop, share=fs.get("fit_max", 1.0) if logo else 1.0,
                                trim=logo and not color and fs.get("kind") != "icon")
        fentry["asset"] = {"ref": ref, "source": source, "sha256": sha256_file(path)}
        if contained:
            fentry["fit"] = "contain"
            pic.name = images.WHOLE
        if _dpi(path, geo, crop=crop) < MIN_DPI:
            self.issues.append(Issue("ASSET_LOW_RES", f"{ref!r} shows below {MIN_DPI} DPI at this size",
                                     severity="warning", **at))
        return pic

    def slide(self, prs: Any, layouts: dict[Any, Any], s: Slide, n: int) -> dict[str, Any]:
        """Add slide n on its template layout, fill its fields, drop unused placeholders, add the slide number
        and footer, and return its manifest entry."""
        ls = self.brand.tokens["layouts"][s.layout]
        tl = ls["template_layout"]
        layout = tpl.find_layout(layouts, tl, ls.get("master"))
        slide = prs.slides.add_slide(layout)
        phs = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        entry: dict[str, Any] = {"slide": n, "layout": s.layout, "template_layout": tl, "title": s.title,
                                 "fields": {}}
        used: set[int] = set()
        shapes: dict[str, Any] = {}  # field -> the shape it ended up in, for `build:`
        if ls.get("row"):
            lay_out_row(ls["fields"], set(s.fields), phs)
        # Fill in the layout's declared field order, so the output doesn't depend on input order
        # (a workbook's columns can order fields differently than the markdown did).
        for name in [f for f in ls["fields"] if f in s.fields]:
            val = s.fields[name]
            fs = ls["fields"][name]
            idx, kind = fs["idx"], fs.get("kind", "text")
            at: dict[str, Any] = {"slide": n, "field": name, "file": s.where.file if s.where else None}
            ph = phs.get(idx)
            if ph is None:
                self.issues.append(Issue("TEMPLATE_MISMATCH", f"placeholder idx {idx} isn't on layout {tl!r}", **at))
                continue
            if isinstance(val, str) and not val.strip():
                continue  # `kicker: ""` leaves the field out, and its placeholder goes with the unused ones
            used.add(idx)
            fentry: dict[str, Any] = {"idx": idx, "kind": kind}
            shapes[name] = ph
            if isinstance(val, Chart):
                shapes[name] = fill_chart(slide, ph, val, self.brand)
            elif isinstance(val, Table):
                shapes[name] = fill_table(slide, ph, val, self.brand)
            elif isinstance(val, Code):
                fill_code(slide, ph, val, self.brand, self.code_font)
                fentry.update({"language": val.language, "lines": len(val.lines), "shape": ph.name})
            elif isinstance(val, Image):
                shapes[name] = self._picture(slide, ph, val.ref, val.alt, at, fentry, None, fs, s)
            elif isinstance(val, Icon):
                icons = self.brand.meta.get("icons") or {}
                color = self.brand.color(fs.get("color") or icons.get("default_color"))
                shapes[name] = self._picture(slide, ph, val.ref, "", at, fentry, color, fs, s)
            else:
                fill_text(ph, val, self.code_font)
                emphasis = self.brand.color(fs.get("emphasis"))
                if emphasis:
                    color_bold(ph, emphasis)
                fentry["chars"] = len(plain(text_of(val)))
                fentry["shape"] = ph.name  # lets the render step name the field that overflowed
            entry["fields"][name] = fentry
        if s.current is not None:
            mark_current(slide, phs, ls["fields"], s.current, self.brand, cards=s.layout.startswith("cards-"))
            entry["current"] = s.current
        if s.build is not None:
            add_fade(slide, build_steps(s.build, shapes, ls["fields"]))
            entry["build"] = s.build
        for idx, ph in phs.items():
            if idx not in used:  # no empty "Click to add text" boxes left behind
                ph._element.getparent().remove(ph._element)
        add_furniture(slide, layout, n + self.first - 1, self.numbers, self.footer, fixed=self.first != 1)
        if s.notes:
            slide.notes_slide.notes_text_frame.text = s.notes
            entry["notes_chars"] = len(s.notes)
        return entry


def build(deck: Deck, brand: Brand, deck_path: Path, out_path: Path, cache_dir: Path, content_sha: str,
          keep_template_slides: bool = False, force: bool = False) -> tuple[dict[str, Any], list[Issue]]:
    """content_sha identifies the deck's content independent of its file format (see pipeline.content_sha)."""
    assert brand.template is not None
    mpath = manifest_path(out_path)
    # out_path and its manifest sidecar can each already be an existing symlink; confine both before
    # writing either (a no-op outside an MCP call, where there's no configured area to confine to).
    confine.guard(out_path, "the build output")
    confine.guard(mpath, "the build manifest")
    if out_path.exists():
        if not mpath.is_file():
            raise EnvError(f"{out_path} exists and has no {mpath.name} beside it, so deck-builder "
                           "didn't build it; move it or write somewhere else")
        if not force and _recorded_sha256(mpath) != sha256_file(out_path):
            raise EnvError(f"{out_path} was saved after deck-builder built it, so it may hold someone's edits, and "
                           "nothing was built. To build without touching it, write to a new file: front matter "
                           f"`output: {out_path.stem}-v2.pptx`, or `build -o <new path>`. Replacing it with "
                           "--force loses those edits, so only the user decides that.", code="OUTPUT_EDITED")
    prs = tpl.open_template(brand.template)
    if not keep_template_slides:
        tpl.remove_all_slides(prs)
    layouts = tpl.layouts(prs)
    b = Builder(brand, deck_path.parent, cache_dir, *furniture_settings(deck.meta), first_number(deck.meta))
    slides = [b.slide(prs, layouts, s, n) for n, s in enumerate(deck.slides, start=1)]
    if b.first != 1:  # so a slide-number field added in PowerPoint later counts from here too
        prs.part._element.set("firstSlideNum", str(b.first))

    when = deck_date(deck.meta)
    author = str(deck.meta.get("author", ""))
    cp = prs.core_properties
    cp.title = str(deck.meta.get("title", ""))
    cp.author = author
    cp.last_modified_by = author
    cp.comments = ""
    cp.revision = 1
    cp.created = when.replace(tzinfo=None)
    cp.modified = when.replace(tzinfo=None)

    input_sha = sha256_file(deck_path)
    template_sha = sha256_file(brand.template)
    props = {
        "deck-builder:engine": __version__,
        "deck-builder:brand": brand.slug,
        "deck-builder:brand-version": brand.version,
        "deck-builder:template-sha256": template_sha,
        "deck-builder:content-sha256": content_sha,
    }
    buf = io.BytesIO()
    prs.save(buf)
    blob = normalize(buf.getvalue(), when.strftime("%Y-%m-%dT%H:%M:%SZ"), props)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    manifest = {
        "engine": __version__,
        "brand": {"slug": brand.slug, "version": brand.version, "template_sha256": template_sha},
        "input": {"file": deck_path.name, "path": str(deck_path.resolve()), "sha256": input_sha,
                  "content_sha256": content_sha},
        "output": {"file": out_path.name, "sha256": hashlib.sha256(blob).hexdigest(),
                   "content_sha256": content_digest(blob)},
        "slides": slides,
    }
    # Write both to temporary names in the same folder, then rename: a failure partway never leaves a
    # half-written .pptx, or a .pptx and manifest that don't match each other.
    tmp_out = out_path.with_name(out_path.name + ".tmp")
    tmp_manifest = mpath.with_name(mpath.name + ".tmp")
    tmp_out.write_bytes(blob)
    tmp_manifest.write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    tmp_out.replace(out_path)
    tmp_manifest.replace(mpath)
    return manifest, b.issues


def _recorded_sha256(mpath: Path) -> str | None:
    try:
        return str(json.loads(mpath.read_text(encoding="utf-8"))["output"]["sha256"])
    except (OSError, ValueError, KeyError, TypeError):
        return None


def manifest_path(out_path: Path) -> Path:
    return out_path.with_name(out_path.stem + ".manifest.json")
