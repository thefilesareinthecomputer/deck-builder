"""`import`: an existing .pptx back into deck.md, its images, and a report of what needs a decision.

The .pptx is untrusted. It opens through the template size cap, nothing it links to is fetched, and
extracted images get names made here, never names taken from the file. Local formatting (fonts, sizes,
colors, moved boxes) is dropped on purpose, since the rebuild takes the brand's styling; the report
counts what was dropped. Content is never dropped silently: what can't be placed in a field goes to
the slide's speaker notes and the report.
"""
from __future__ import annotations

import hashlib
import re
from collections import Counter
from collections.abc import Iterable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.text.text import _Run

from deck_builder import template as tpl
from deck_builder.assets import recolor_icon, sha256_file
from deck_builder.brand.registry import Brand
from deck_builder.build.deck import DEFAULT_DATE
from deck_builder.build.text import CODE_FONT_DEFAULT
from deck_builder.build.visuals import CHART_TYPES
from deck_builder.model import Bullets, Chart, Deck, Image, Series, Slide, Table, Value
from deck_builder.parse.markdown import number
from deck_builder.validate import IMAGE_EXT, confined
from deck_builder.write.markdown import table_md

LINK_SCHEMES = ("http://", "https://", "mailto:")
FURNITURE = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
TITLES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
FITS = {"text": {"text", "bullets"}, "image": {"image"}, "icon": {"icon"}, "table": {"table"},
        "chart": {"chart"}}
EXACT_ONLY = ("takeaway",)  # filled only from the matching placeholder idx, never by type or position
CHART_NAMES = {v: k for k, v in CHART_TYPES.items()}
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
NOT_FORMATTING = {"lang", "altLang", "dirty", "err", "smtClean", "smtId", "noProof", "bmk", "b", "i"}
FILLS = ("a:solidFill", "a:gradFill", "a:pattFill", "a:blipFill", "a:noFill")


@dataclass
class Found:
    """One piece of slide content before it's placed in a field."""

    kind: str  # text | image | icon | table | chart
    value: Value
    what: str  # how the report and notes name it
    box: tuple[int, int, int, int]  # left, top, width, height in EMU
    idx: int | None = None  # placeholder idx; None for a free shape
    title: bool = False


@dataclass
class SlideReport:
    n: int
    title: str = ""
    layout: str = ""
    match: str = ""
    unplaced: list[str] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)
    dropped: Counter[str] = field(default_factory=Counter)


@dataclass
class Imported:
    deck: Deck
    reports: list[SlideReport]
    assets: dict[str, bytes]  # file name in assets/ -> bytes


# ---------------------------------------------------------------- text


def _style(run: _Run, code_font: str) -> tuple[bool, bool, bool, str]:
    rpr = run._r.rPr
    bold = rpr is not None and rpr.get("b") in ("1", "true")
    italic = rpr is not None and rpr.get("i") in ("1", "true")
    link = str(run.hyperlink.address or "")
    # Only http(s) and mailto survive import as a link; anything else (javascript:, file:, a UNC
    # path, ...) becomes plain text - the rebuild never writes a scheme deck.md authors didn't type.
    if link and not link.lower().startswith(LINK_SCHEMES):
        link = ""
    return bold, italic, run.font.name == code_font, link


def _md(text: str, bold: bool, italic: bool, code: bool, link: str) -> str:
    """One styled stretch as deck.md inline markup, with surrounding spaces kept outside the markers."""
    core = text.strip()
    if not core:
        return text
    lead, trail = text[: len(text) - len(text.lstrip())], text[len(text.rstrip()) :]
    if link:
        core = f"[{core}]({link})"
    elif code:
        core = f"`{core}`"
    elif bold:
        core = f"**{core}**"
    elif italic:
        core = f"*{core}*"
    return lead + core + trail


def paragraph_md(p: Any, code_font: str) -> str:
    stretches: list[tuple[str, tuple[bool, bool, bool, str]]] = []
    for el in p._p:
        if el.tag == qn("a:r"):
            run = _Run(el, p)
            stretches.append((str(run.text), _style(run, code_font)))
        elif el.tag == qn("a:br"):
            stretches.append((" ", (False, False, False, "")))
        elif el.tag == qn("a:fld"):
            t = el.find(qn("a:t"))
            stretches.append((str(t.text or "") if t is not None else "", (False, False, False, "")))
    merged: list[list[Any]] = []
    for text, style in stretches:
        if merged and merged[-1][1] == style:
            merged[-1][0] += text
        else:
            merged.append([text, style])
    return "".join(_md(t, *s) for t, s in merged).strip()


def text_value(tf: Any, code_font: str) -> Value | None:
    """A text frame as a string (one top-level paragraph) or bullets; None when it's empty."""
    items: Bullets = [(p.level, md) for p in tf.paragraphs if (md := paragraph_md(p, code_font))]
    if not items:
        return None
    return items[0][1] if len(items) == 1 and items[0][0] == 0 else items


def _plain_lines(value: Value) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, list):
        return [f"{'  ' * lvl}- {t}" for lvl, t in value]
    return []


def overrides(shape: Any, code_font: str) -> Counter[str]:
    """Local formatting on a shape's text that the rebuild drops, by kind."""
    out: Counter[str] = Counter()
    if not getattr(shape, "has_text_frame", False):
        return out
    body_pr = shape.text_frame._txBody.find(qn("a:bodyPr"))
    if body_pr is not None and (len(body_pr) or body_pr.attrib):
        out["text box settings"] += 1
    for p in shape.text_frame.paragraphs:
        ppr = p._p.pPr
        if ppr is not None and ({k for k in ppr.attrib if k != "lvl"} or len(ppr)):
            out["paragraph"] += 1
        for el in p._p.iter(qn("a:rPr"), qn("a:endParaRPr")):
            if el.tag == qn("a:endParaRPr"):
                continue
            if el.get("sz"):
                out["size"] += 1
            latin = el.find(qn("a:latin"))
            if latin is not None and latin.get("typeface") != code_font or el.find(qn("a:ea")) is not None:
                out["font"] += 1
            if any(el.find(qn(f)) is not None for f in FILLS) or el.find(qn("a:highlight")) is not None:
                out["color"] += 1
            if set(el.attrib) - NOT_FORMATTING - {"sz"}:
                out["other"] += 1
    return out


# ---------------------------------------------------------------- tables and charts


def table_value(tbl: Any, code_font: str, status_keys: Iterable[str] = ()) -> Table:
    status = {k.strip().casefold() for k in status_keys}

    def cell(c: Any, header: bool) -> str:
        text = " ".join(md for p in c.text_frame.paragraphs if (md := paragraph_md(p, code_font)))
        if header and text.startswith("**") and text.endswith("**") and text.count("**") == 2:
            return text[2:-2]  # a bold header row is the table's styling, not markup
        if not header and text.startswith("● "):
            rest = text[2:]
            if rest.strip().casefold() in status:
                return rest  # the status dot is rendering, not the cell's own text
        return text

    rows = [[cell(c, n == 0) for c in row.cells] for n, row in enumerate(tbl.rows)]
    return Table(header=rows[0] if rows else [], rows=rows[1:])


def _chart_type(ct: Any) -> tuple[str | None, bool]:
    """A python-pptx chart type -> (deck.md chart type, exact match)."""
    if ct in CHART_NAMES:
        return CHART_NAMES[ct], True
    name = getattr(ct, "name", "") or ""
    stacked = "STACKED" in name
    for key, kind in (("DOUGHNUT", "doughnut"), ("PIE", "pie"), ("BAR", "bar"), ("COLUMN", "column"),
                      ("CYLINDER", "column"), ("CONE", "column"), ("PYRAMID", "column"), ("LINE", "line"),
                      ("AREA", "line")):
        if key in name:
            return (f"stacked-{kind}" if stacked and kind in ("bar", "column") else kind), False
    return None, False


def _colors(chart: Any, kind: str) -> list[str]:
    plot = chart.plots[0]._element
    if kind in ("pie", "doughnut"):
        sers = plot.findall(qn("c:ser"))
        els = sers[0].findall(f"{qn('c:dPt')}/{qn('c:spPr')}/{qn('a:solidFill')}/{qn('a:srgbClr')}") if sers else []
    else:
        path = (f"{qn('c:spPr')}/{qn('a:ln')}/{qn('a:solidFill')}/{qn('a:srgbClr')}" if kind == "line"
                else f"{qn('c:spPr')}/{qn('a:solidFill')}/{qn('a:srgbClr')}")
        els = [s.find(path) for s in plot.findall(qn("c:ser"))]
    if not els or any(e is None for e in els):
        return []
    return [str(e.get("val")).upper() for e in els]


def _shortest_cycle(xs: list[str]) -> list[str]:
    for n in range(1, len(xs) + 1):
        if all(xs[i] == xs[i % n] for i in range(len(xs))):
            return xs[:n]
    return xs


def chart_value(chart: Any, brand: Brand) -> tuple[Chart | None, str]:
    """A chart as a deck.md chart, or None with the reason it can't be one."""
    kind, exact = _chart_type(chart.chart_type)
    if kind is None:
        return None, f"a {getattr(chart.chart_type, 'name', 'chart').lower()} chart has no deck.md chart type"
    plot = chart.plots[0]
    series = [Series(name=str(s.name), values=[number(v) for v in s.values]) for s in plot.series]
    fmt = plot._element.find(f"{qn('c:ser')}/{qn('c:val')}/{qn('c:numRef')}/{qn('c:numCache')}/{qn('c:formatCode')}")
    pie = kind in ("pie", "doughnut")
    has_legend = bool(chart.has_legend)
    title = None
    if chart.has_title and chart.chart_title.has_text_frame:
        title = chart.chart_title.text_frame.text.strip() or None
    found = _colors(chart, kind)
    colors: list[str] | None = None
    if found:
        tok = brand.tokens.get("chart") or {}
        default = [h for h in (brand.color(str(n)) for n in tok.get("colors") or []) if h]
        n = len(found)
        if not default or [default[i % len(default)] for i in range(n)] != found:
            palette = list((brand.meta.get("palette") or {}).items())
            names = {str(v).lstrip("#").upper(): k for k, v in reversed(palette)}  # the first name wins
            colors = [names.get(h, h) for h in _shortest_cycle(found)]
    value = Chart(
        type=kind,
        categories=[str(c) for c in plot.categories],
        series=series,
        number_format=str(fmt.text or "General") if fmt is not None else "General",
        labels=bool(plot.has_data_labels),
        legend=None if has_legend == (pie or len(series) > 1) else has_legend,
        title=title,
        colors=colors,
    )
    return value, "" if exact else f"chart type {chart.chart_type.name} imported as {kind}"


def _extra_plot_series(chart: Any) -> list[str]:
    """Series from every plot after the first: a combo chart's extra plots (such as a line layered over
    a bar plot) would otherwise vanish, since only plots[0] becomes the field's chart."""
    out = []
    for plot in chart.plots[1:]:
        for s in plot.series:
            out.append(f"{s.name}: {', '.join(str(number(v)) for v in s.values)}")
    return out


# ---------------------------------------------------------------- the importer


class Importer:
    def __init__(self, brand: Brand, cache_dir: Path) -> None:
        self.brand = brand
        self.code_font = (brand.tokens.get("text") or {}).get("code_font", CODE_FONT_DEFAULT)
        self.assets: dict[str, bytes] = {}
        self.footers: dict[str, list[int]] = {}  # footer text -> slides showing it
        self.numbered = False  # any slide shows a slide number
        self.can_number = False  # any slide's layout has a slide-number placeholder
        self.known: dict[str, tuple[str, str]] = {}  # image sha256 -> (kind, brand ref)
        # logos: and icons: dir come from brand.yaml; confine each to the kit before it's hashed or
        # decoded, the same boundary `check` enforces for a deck's own asset references (A4).
        for lid, rel in (brand.meta.get("logos") or {}).items():
            p = brand.path / str(rel)
            if p.is_file() and confined(p, brand.path):
                self.known.setdefault(sha256_file(p), ("image", f"brand:logo/{lid}"))
        icons = brand.meta.get("icons") or {}
        icon_dir = brand.path / str(icons.get("dir", ""))
        if icons.get("dir") and icon_dir.is_dir() and confined(icon_dir, brand.path):
            colors = {icons.get("default_color")}
            for ls in (brand.tokens.get("layouts") or {}).values():
                colors |= {fs.get("color") for fs in (ls.get("fields") or {}).values() if fs.get("kind") == "icon"}
            for ref in sorted(c for c in colors if c):
                hexv = brand.color(str(ref))
                if not hexv:
                    continue
                for png in sorted(icon_dir.glob("*.png")):
                    if not confined(png, brand.path):
                        continue
                    self.known.setdefault(sha256_file(recolor_icon(png, hexv, cache_dir)),
                                          ("icon", f"brand:icon/{png.stem}"))

    # -- shapes

    def picture(self, shape: Any, n: int, rep: SlideReport) -> tuple[str, Value] | None:
        blip = shape._element.find(f"{qn('p:blipFill')}/{qn('a:blip')}")
        if blip is None or not blip.get(qn("r:embed")):
            link = blip.get(qn("r:link")) if blip is not None else None
            target = shape.part.rels[link].target_ref if link and link in shape.part.rels else "an unknown target"
            rep.skipped.append(f"linked image {shape.name!r} points at {target}; not fetched")
            return None
        img = shape.image
        sha = hashlib.sha256(img.blob).hexdigest()
        alt = str(shape._element.find(f"{qn('p:nvPicPr')}/{qn('p:cNvPr')}").get("descr") or "")
        if sha in self.known:
            kind, ref = self.known[sha]
            return (kind, ref) if kind == "icon" else (kind, Image(ref=ref, alt=alt))
        ext = "." + img.ext.lower()
        if ext not in IMAGE_EXT:
            rep.skipped.append(f"picture {shape.name!r} is {img.ext.upper()}, which decks can't use; convert it to PNG")
            return None
        name = f"slide{n:02d}-{sha[:12]}{'.jpg' if ext == '.jpeg' else ext}"
        self.assets[name] = img.blob
        return "image", Image(ref=f"assets/{name}", alt=alt)

    def shapes(self, shapes: Any, n: int, rep: SlideReport) -> list[Found]:
        found: list[Found] = []
        for sh in shapes:
            box = (int(sh.left or 0), int(sh.top or 0), int(sh.width or 0), int(sh.height or 0))
            ph = sh.placeholder_format if sh.is_placeholder else None
            idx = ph.idx if ph is not None else None
            if ph is not None and ph.type in FURNITURE:  # becomes front matter, not slide content
                text = sh.text_frame.text.strip() if sh.has_text_frame else ""
                if ph.type == PP_PLACEHOLDER.FOOTER and text:
                    self.footers.setdefault(text, []).append(n)
                self.numbered |= ph.type == PP_PLACEHOLDER.SLIDE_NUMBER
                continue
            st = sh.shape_type
            if st == MSO_SHAPE_TYPE.GROUP:
                inner = [t for f in self.shapes(sh.shapes, n, SlideReport(n)) for t in _describe(f)]
                rep.skipped.append(f"grouped shapes {sh.name!r}; their content is in the notes")
                rep.unplaced.append(f"Group {sh.name!r}: " + " / ".join(inner) if inner else f"Group {sh.name!r}")
                continue
            if st == MSO_SHAPE_TYPE.MEDIA:
                rep.skipped.append(f"video or audio {sh.name!r} wasn't imported; noted on the slide")
                rep.unplaced.append(f"Video or audio {sh.name!r} wasn't imported; add it in PowerPoint")
                continue
            if st in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
                rep.skipped.append(f"embedded object {sh.name!r} wasn't imported; noted on the slide")
                rep.unplaced.append(f"Embedded object {sh.name!r} wasn't imported; rebuild it as a table or image")
                continue
            if getattr(sh, "has_chart", False) and sh.has_chart:
                chart, why = chart_value(sh.chart, self.brand)
                if why:
                    rep.skipped.append(why)
                if chart is None:
                    rep.unplaced.append(f"Chart {sh.name!r}: {why}")
                    continue
                extra = _extra_plot_series(sh.chart)
                if extra:
                    rep.skipped.append(f"chart {sh.name!r} has {len(sh.chart.plots)} plots; only the first "
                                       "is placed")
                    rep.unplaced.append(f"Chart {sh.name!r} extra plot series (not placed): " + " / ".join(extra))
                found.append(Found("chart", chart, f"chart {sh.name!r}", box, idx))
                continue
            if getattr(sh, "has_table", False) and sh.has_table:
                status_keys = (self.brand.tokens.get("table") or {}).get("status") or {}
                found.append(Found("table", table_value(sh.table, self.code_font, status_keys),
                                   f"table {sh.name!r}", box, idx))
                continue
            if sh._element.tag == qn("p:graphicFrame"):
                gd = sh._element.find(f".//{qn('a:graphicData')}")
                what = "SmartArt diagram" if gd is not None and gd.get("uri") == DIAGRAM_URI else "graphic"
                rep.skipped.append(f"{what} {sh.name!r} wasn't imported; noted on the slide")
                rep.unplaced.append(f"{what[0].upper() + what[1:]} {sh.name!r} wasn't imported; rebuild it as "
                                    "bullets or an image")
                continue
            if st == MSO_SHAPE_TYPE.PICTURE or sh._element.tag == qn("p:pic"):
                got = self.picture(sh, n, rep)
                if got is not None:
                    found.append(Found(got[0], got[1], f"picture {sh.name!r}", box, idx))
                continue
            rep.dropped += overrides(sh, self.code_font)
            if ph is not None and sh._element.find(f"{qn('p:spPr')}/{qn('a:xfrm')}") is not None:
                rep.dropped["position"] += 1  # a placeholder moved or resized on the slide
            if sh.has_text_frame:
                value = text_value(sh.text_frame, self.code_font)
                if value is not None:
                    is_title = ph is not None and ph.type in TITLES
                    what = f"placeholder {sh.name!r}" if ph is not None else f"text box {sh.name!r}"
                    found.append(Found("text", value, what, box, idx, title=is_title))
                    continue
            if ph is None:
                rep.dropped["shape"] += 1  # lines, rectangles and other decoration without text
        return found

    # -- layouts

    def _candidates(self, slide: Any) -> list[str]:
        name = slide.slide_layout.name
        master = slide.slide_layout.slide_master.name
        out = []
        for key, ls in (self.brand.tokens.get("layouts") or {}).items():
            if ls.get("template_layout") == name and ls.get("master") in (None, master):
                out.append(key)
        return out

    def _assign(self, key: str, found: list[Found], by_idx: bool, geometry: dict[str, tuple[int, int, int, int]]
                ) -> tuple[dict[str, Found], Found | None, list[Found], int]:
        """Place content in a layout's fields: by placeholder idx, then by position, then the one free field.

        Also returns how many text boxes were placed the last way: that says nothing about how well the
        layout fits, so choosing a layout doesn't count them.
        """
        ls = self.brand.tokens["layouts"][key]
        fspecs: dict[str, Any] = ls.get("fields") or {}
        hf = ls.get("heading_field", "title")
        placed: dict[str, Found] = {}
        heading: Found | None = None
        rest: list[Found] = []
        by_field_idx = {fs["idx"]: f for f, fs in fspecs.items()}
        for f in sorted(found, key=lambda f: (f.box[1], f.box[0])):
            target = by_field_idx.get(f.idx) if by_idx and f.idx is not None else None
            if target is None and not by_idx and f.title and f.kind == "text":
                target = hf
            if target == hf and f.kind == "text" and heading is None:
                heading = f
            elif target and target not in placed and fspecs[target].get("kind", "text") in FITS[f.kind]:
                placed[target] = f
            else:
                rest.append(f)
        left: list[Found] = []
        for f in rest:
            free = sorted((k for k, fs in fspecs.items() if k not in (hf, *EXACT_ONLY) and k not in placed
                           and fs.get("kind", "text") in FITS[f.kind]),
                          key=lambda k: fspecs[k].get("kind", "text") != _preferred(f))  # stable: field order
            cx, cy = f.box[0] + f.box[2] // 2, f.box[1] + f.box[3] // 2
            inside = [k for k in free if k in geometry and _contains(geometry[k], cx, cy)]
            if by_idx and len(inside) == 1:
                placed[inside[0]] = f
            elif not by_idx and free and f.idx is not None:
                placed[free[0]] = f  # a placeholder, by type: the next free field of its kind, in field order
            else:
                left.append(f)
        if heading is None and hf in fspecs and not by_idx:
            texts = [f for f in left if f.kind == "text" and isinstance(f.value, str)]
            if texts:
                heading = texts[0]
                left.remove(heading)
        fallback = 0
        for f in list(left):
            free = [k for k, fs in fspecs.items() if k not in (hf, *EXACT_ONLY) and k not in placed
                    and fs.get("kind", "text") in FITS[f.kind]]
            if len(free) == 1:
                placed[free[0]] = f
                left.remove(f)
                fallback += f.kind == "text"  # a lone chart or picture fits its field; a stray text box says nothing
        return placed, heading, left, fallback

    def _geometry(self, key: str, prs_layouts: dict[Any, Any]) -> dict[str, tuple[int, int, int, int]]:
        ls = self.brand.tokens["layouts"][key]
        layout = tpl.find_layout(prs_layouts, ls["template_layout"], ls.get("master"))
        if layout is None:
            return {}
        boxes = {ph.placeholder_format.idx: (int(ph.left or 0), int(ph.top or 0), int(ph.width or 0),
                                             int(ph.height or 0)) for ph in layout.placeholders}
        return {f: boxes[fs["idx"]] for f, fs in (ls.get("fields") or {}).items() if fs["idx"] in boxes}

    def slide(self, slide: Any, n: int, prs_layouts: dict[Any, Any], same_template: bool) -> tuple[Slide, SlideReport]:
        rep = SlideReport(n)
        self.can_number |= any(p.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
                               for p in slide.slide_layout.placeholders)
        found = self.shapes(slide.shapes, n, rep)
        names = self._candidates(slide) if same_template else []
        best: tuple[float, float, str, dict[str, Found], Found | None, list[Found]] | None = None
        keys = names or list(self.brand.tokens.get("layouts") or {})
        for key in keys:
            ls = self.brand.tokens["layouts"][key]
            fspecs = ls.get("fields") or {}
            placed, heading, left, fallback = self._assign(key, found, bool(names),
                                                           self._geometry(key, prs_layouts) if names else {})
            fit = (len(placed) - fallback + (heading is not None)) / (len(found) or 1)
            missing = [f for f, fs in fspecs.items()
                       if fs.get("required") and f not in placed and f != ls.get("heading_field", "title")]
            exact = sum(fspecs[k].get("kind", "text") == _preferred(f) for k, f in placed.items())
            named = bool(_words(slide.slide_layout.name) & (_words(key) | _words(str(ls.get("template_layout")))))
            # then the fewest required gaps, the closest kinds, a shared word in the layout name, the simplest
            score = fit - 0.01 * len(missing) + 0.005 * exact + 0.002 * named - 0.001 * len(fspecs)
            if best is None or score > best[0]:
                best = (score, fit, key, placed, heading, left)
        assert best is not None
        _, fit, key, placed, heading, left = best
        if names:
            rep.match = f"matched by layout name ({slide.slide_layout.name!r})"
        else:
            rep.match = (f"best match on placeholder types, confidence {fit:.0%} "
                         f"(the original used {slide.slide_layout.name!r})")
        rep.layout = key
        title = ""
        if heading is not None:
            hv = heading.value
            title = hv if isinstance(hv, str) else " ".join(t for _, t in hv)  # type: ignore[union-attr]
        fspecs = self.brand.tokens["layouts"][key].get("fields") or {}
        fields: dict[str, Value] = {}
        for name in fspecs:
            if name in placed:
                v = placed[name].value
                if isinstance(v, str) and fspecs[name].get("kind") == "bullets" and placed[name].kind == "text":
                    v = [(0, v)]  # one paragraph builds the same either way; the field's kind decides
                fields[name] = v
        for f in left:
            rep.unplaced.append(_unplaced(f))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        if rep.unplaced:
            lines = "\n".join(f"- {u}" for u in rep.unplaced)
            # rstrip: an unplaced table's own dump ends with a newline that a reparse would trim anyway,
            # so trim it here too; otherwise this slide's notes could never round-trip through deck.md.
            notes = ((notes + "\n\n" if notes else "") + f"Unplaced from the original:\n{lines}").rstrip()
        rep.title = title
        return Slide(title=title, layout=key, fields=fields, notes=notes), rep


def _words(name: str) -> set[str]:
    return set(re.findall(r"[a-z]+", name.lower())) - {"and", "with", "the", "a", "of", "slide", "layout"}


def _preferred(f: Found) -> str:
    """The field kind that fits this content best: bullets for a list, text for one paragraph."""
    if f.kind == "text":
        return "bullets" if isinstance(f.value, list) else "text"
    return f.kind


def _contains(box: tuple[int, int, int, int], x: int, y: int) -> bool:
    left, top, w, h = box
    return left <= x <= left + w and top <= y <= top + h


def _describe(f: Found) -> list[str]:
    if f.kind == "text":
        return _plain_lines(f.value)
    if isinstance(f.value, Image):
        return [f"image {f.value.ref}"]
    if f.kind == "icon":
        return [f"icon {f.value}"]
    return [f.what]


def _unplaced(f: Found) -> str:
    v = f.value
    if f.kind == "text":
        return f"{f.what[0].upper() + f.what[1:]}: " + " / ".join(_plain_lines(v))
    if isinstance(v, Image):
        return f"{f.what[0].upper() + f.what[1:]}: ![{v.alt}]({v.ref})"
    if isinstance(v, Table):
        return f"{f.what[0].upper() + f.what[1:]}:\n\n{table_md(v)}\n"
    if isinstance(v, Chart):
        series = "; ".join(f"{s.name}: {', '.join(map(str, s.values))}" for s in v.series)
        return f"{f.what[0].upper() + f.what[1:]} ({v.type}) by {', '.join(v.categories)}: {series}"
    return f"{f.what[0].upper() + f.what[1:]}: {v}"


def import_pptx(pptx: Path, brand: Brand, cache_dir: Path) -> Imported:
    """Read every slide onto the brand's layouts; front matter from core properties and slide furniture."""
    prs = tpl.open_template(pptx)
    assert brand.template is not None
    kit_layouts = tpl.layouts(tpl.open_template(brand.template))
    same = _same_layouts(prs, kit_layouts)
    imp = Importer(brand, cache_dir)
    slides, reports = [], []
    for n, s in enumerate(prs.slides, start=1):
        slide, rep = imp.slide(s, n, kit_layouts, same)
        slides.append(slide)
        reports.append(rep)
    meta: dict[str, Any] = {"brand": brand.slug}
    cp = prs.core_properties
    if (cp.title or "").strip():
        meta["title"] = cp.title.strip()
    if (cp.author or "").strip():
        meta["author"] = cp.author.strip()
    if cp.created and cp.created.date() != DEFAULT_DATE:
        meta["date"] = cp.created.date()
    if len(imp.footers) == 1:
        meta["footer"] = next(iter(imp.footers))
    elif imp.footers:  # several footer texts: the deck gets one, so the user picks it
        for text, ns in imp.footers.items():
            reports[ns[0] - 1].skipped.append(f"footer text {text!r} on slides {', '.join(map(str, ns))}; the "
                                              "footers differ, so set `footer:` in the front matter yourself")
    if imp.can_number and not imp.numbered:
        meta["slide_numbers"] = False
    return Imported(Deck(meta=meta, slides=slides, source=str(pptx)), reports, imp.assets)


def _same_layouts(prs: Any, kit_layouts: dict[Any, Any]) -> bool:
    """The deck was made from this kit's template: every layout its slides use exists there by name."""
    return all((s.slide_layout.slide_master.name, s.slide_layout.name) in kit_layouts for s in prs.slides)


def report_md(pptx: Path, brand: Brand, adopted: bool, imp: Imported, mismatched: list[int]) -> str:
    unplaced = sum(len(r.unplaced) for r in imp.reports)
    dropped = sum(sum(r.dropped.values()) for r in imp.reports)
    how = "a new kit adopted from this file" if adopted else "an existing kit"
    out = [f"# Import report: {pptx.name}", "",
           f"Brand: `{brand.slug}`, {how}. Slides: {len(imp.reports)}. Images extracted: {len(imp.assets)}. "
           f"Unplaced items: {unplaced}. Dropped formatting: {dropped}.", "",
           "Unplaced items are in each slide's speaker notes under \"Unplaced from the original:\". Move them "
           "into fields or cut them, then run `deck-builder check`. Dropped formatting is intended: the rebuild "
           "takes the brand's styling.", ""]
    for r in imp.reports:
        out.append(f"## Slide {r.n}: {r.title or '(no title)'}")
        out.append("")
        out.append(f"- Layout `{r.layout}`, {r.match}.")
        for u in r.unplaced:
            out.append(f"- Unplaced: {u.splitlines()[0]}")
        for s in r.skipped:
            out.append(f"- Not imported: {s}.")
        if r.dropped:
            out.append("- Dropped formatting: " + ", ".join(f"{v} {k}" for k, v in sorted(r.dropped.items())) + ".")
        if r.n in mismatched:
            out.append("- deck.md doesn't hold this slide exactly as imported; compare it with the original.")
        out.append("")
    return "\n".join(out).rstrip() + "\n"
