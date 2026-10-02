"""`brand init`: generate template.potx and tokens.yaml from brand.yaml. Deterministic.

Starts from python-pptx's bundled template (one master, a theme, 11 stock 4:3 layouts), resizes it,
writes the brand's colors and fonts into the theme, replaces the stock layouts with the chosen
layout set, and optionally places a logo on the master.
"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

from lxml import etree
from PIL import Image as PILImage
from pptx import Presentation
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.parts.slide import SlideLayoutPart
from pptx.util import Emu

from deck_builder.brand.inspect import estimate_chars
from deck_builder.brand.kit import contrast
from deck_builder.brand.layouts import PH, LayoutDef, layout_set
from deck_builder.build.normalize import read_parts, rezip
from deck_builder.template import POTX_CT, PPTX_CT

EMU = 914400
SIZES_EMU = {"16:9": (12192000, 6858000), "4:3": (9144000, 6858000)}  # PowerPoint's exact sizes
NSDECL = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
          'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
          'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
SLOTS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink",
         "folHlink")
DEFAULT_SLOT_COLORS = {"dk1": "1B1B1B", "lt1": "FFFFFF", "dk2": "1F3A5F", "lt2": "F2F2F2", "accent1": "1F3A5F",
                       "accent2": "E07A2F", "accent3": "6B7280", "accent4": "9DB4C0", "accent5": "C8553D",
                       "accent6": "5B8C5A", "hlink": "1F5FBF", "folHlink": "6B7280"}
# Slide furniture: a small number at the bottom left, then an optional footer text, centered on the
# logo's line in the footer band. Layout idx values sit above the layouts' own fields (up to 17).
LOGO_H, LOGO_BOTTOM = 0.42, 0.25  # the master logo's height and its gap to the slide's bottom edge
FURNITURE_PT, FURNITURE_H = 10, 0.3
NUMBER_W = 0.4  # a two-digit number at 10 pt is about 0.15 in, so the footer starts about 0.25 in after it
FOOTER_IDX, NUMBER_IDX = 20, 21
FURNITURE_KINDS = {"SLIDE_NUMBER": "sldNum", "FOOTER": "ftr", "DATE": "dt"}  # dt never gets a box
SLIDENUM_FIELD = "{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}"  # any fixed id keeps builds byte-identical
PH_TYPE = {"title": "title", "body": "body", "pic": "pic", "chart": "chart", "tbl": "tbl"}
KIND_FOR = {"pic": "image", "chart": "chart", "tbl": "table"}
PROMPT = {"title": "Title", "body": "Text", "pic": "Picture", "chart": "Chart", "tbl": "Table"}


def emu(inches: float) -> int:
    return int(round(inches * EMU))


# ---------------------------------------------------------------- theme


def slot_colors(meta: dict[str, Any]) -> dict[str, str]:
    palette = {k: str(v).lstrip("#").upper() for k, v in (meta.get("palette") or {}).items()}
    mapping = meta.get("theme_colors") or {}
    out = dict(DEFAULT_SLOT_COLORS)
    for role, slot in (("ink", "dk1"), ("background", "lt1"), ("primary", "dk2"), ("surface", "lt2"),
                       ("primary", "accent1"), ("accent", "accent2"), ("muted", "accent3")):
        if role in palette:
            out[slot] = palette[role]
    for slot, ref in mapping.items():
        ref = str(ref)
        out[slot] = palette.get(ref, ref.lstrip("#").upper())
    return out


def write_theme(theme_part: Any, colors: dict[str, str], fonts: dict[str, Any]) -> None:
    root = etree.fromstring(theme_part.blob)
    scheme = root.find(f".//{A}clrScheme")
    assert scheme is not None
    for slot in SLOTS:
        el = scheme.find(f"{A}{slot}")
        if el is None:
            continue
        for child in list(el):
            el.remove(child)
        etree.SubElement(el, f"{A}srgbClr", val=colors[slot])
    for role, tag in (("heading", "majorFont"), ("body", "minorFont")):
        fam = (fonts.get(role) or {}).get("family")
        latin = root.find(f".//{A}fontScheme/{A}{tag}/{A}latin")
        if fam and latin is not None:
            latin.set("typeface", fam)
    theme_part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


# ---------------------------------------------------------------- layouts


def _rpr(ph: PH) -> str:
    attrs = ""
    if ph.size:
        attrs += f' sz="{int(ph.size * 100)}"'
    if ph.bold:
        attrs += ' b="1"'
    if ph.italic:
        attrs += ' i="1"'
    fill = f'<a:solidFill><a:schemeClr val="{ph.color}"/></a:solidFill>' if ph.color else ""
    return f"<a:defRPr{attrs}>{fill}</a:defRPr>" if attrs or fill else ""


def _sp(shape_id: int, ph: PH) -> str:
    ph_el = (f'<p:ph type="{PH_TYPE[ph.kind]}"/>' if ph.kind == "title"
             else f'<p:ph type="{PH_TYPE[ph.kind]}" idx="{ph.idx}"/>')
    xfrm = (f'<a:xfrm><a:off x="{emu(ph.x)}" y="{emu(ph.y)}"/>'
            f'<a:ext cx="{emu(ph.w)}" cy="{emu(ph.h)}"/></a:xfrm>')
    if ph.kind in ("title", "body"):
        algn = f' algn="{ph.align}"' if ph.align else ""
        bullet = "" if ph.bullets or ph.kind == "title" else '<a:buNone/>'
        indent = "" if ph.bullets or ph.kind == "title" else ' marL="0" indent="0"'
        lvl1 = f"<a:lvl1pPr{indent}{algn}>{bullet}{_rpr(ph)}</a:lvl1pPr>"
        body = (f'<p:txBody><a:bodyPr anchor="{ph.anchor}"><a:noAutofit/></a:bodyPr><a:lstStyle>{lvl1}'
                f'</a:lstStyle><a:p><a:r><a:rPr lang="en-US"/><a:t>{escape(PROMPT[ph.kind])}</a:t></a:r></a:p>'
                "</p:txBody>")
    else:
        body = '<p:txBody><a:bodyPr/><a:lstStyle/><a:p><a:endParaRPr lang="en-US"/></a:p></p:txBody>'
    name = f"{PROMPT[ph.kind]} Placeholder {shape_id - 1}"
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{shape_id}" name="{escape(name)}"/>'
            f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr>{ph_el}</p:nvPr></p:nvSpPr>'
            f"<p:spPr>{xfrm}</p:spPr>{body}</p:sp>")


def furniture_boxes(w_in: float, h_in: float, numbers: bool) -> dict[str, tuple[float, float, float]]:
    """kind -> (x, y, w) in inches: the number at the margin, the footer after it, centered on the logo's line."""
    y = h_in - LOGO_BOTTOM - LOGO_H / 2 - FURNITURE_H / 2
    footer_x = 0.6 + NUMBER_W if numbers else 0.6
    boxes = {"ftr": (footer_x, y, w_in / 2 - footer_x)}
    if numbers:
        boxes["sldNum"] = (0.6, y, NUMBER_W)
    return boxes


def _furniture_text(color: str) -> str:
    """Plain 10 pt text in one color: no box, fill, rule or bullet. The left inset matches the title's and
    body's, so the text lines up with theirs."""
    return ('<a:bodyPr lIns="91440" tIns="0" rIns="0" bIns="0" anchor="ctr" wrap="none"/><a:lstStyle>'
            f'<a:lvl1pPr algn="l"><a:buNone/><a:defRPr sz="{FURNITURE_PT * 100}" b="0"><a:solidFill>'
            f'<a:srgbClr val="{color}"/></a:solidFill></a:defRPr></a:lvl1pPr></a:lstStyle>')


def _furniture_sp(shape_id: int, kind: str, idx: int, box: tuple[float, float, float], color: str) -> str:
    """A layout's slide-number or footer placeholder, fully specified: LibreOffice doesn't inherit these
    from the master."""
    name = {"sldNum": "Slide Number Placeholder", "ftr": "Footer Placeholder"}[kind]
    para = (f'<a:fld id="{SLIDENUM_FIELD}" type="slidenum"><a:rPr lang="en-US"/><a:t>&#8249;#&#8250;</a:t></a:fld>'
            if kind == "sldNum" else '<a:endParaRPr lang="en-US"/>')
    x, y, w = box
    return (f'<p:sp><p:nvSpPr><p:cNvPr id="{shape_id}" name="{name} {shape_id - 1}"/>'
            f'<p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr><p:nvPr><p:ph type="{kind}" sz="quarter" idx="{idx}"/>'
            f'</p:nvPr></p:nvSpPr><p:spPr><a:xfrm><a:off x="{emu(x)}" y="{emu(y)}"/>'
            f'<a:ext cx="{emu(w)}" cy="{emu(FURNITURE_H)}"/></a:xfrm></p:spPr>'
            f"<p:txBody>{_furniture_text(color)}<a:p>{para}</a:p></p:txBody></p:sp>")


def layout_xml(ld: LayoutDef, boxes: dict[str, tuple[float, float, float]] | None = None,
               color: str = "6B7280") -> bytes:
    bg = ""
    if ld.background:
        bg = (f'<p:bg><p:bgPr><a:solidFill><a:schemeClr val="{ld.background}"/></a:solidFill>'
              "<a:effectLst/></p:bgPr></p:bg>")
    sps = "".join(_sp(i + 2, ph) for i, ph in enumerate(ld.phs))
    if boxes and not ld.hide_master:  # title, section and closing slides get no furniture
        n = len(ld.phs) + 2
        sps += "".join(_furniture_sp(n + i, kind, {"ftr": FOOTER_IDX, "sldNum": NUMBER_IDX}[kind], box, color)
                       for i, (kind, box) in enumerate(boxes.items()))
    show = ' showMasterSp="0"' if ld.hide_master else ""
    return (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<p:sldLayout {NSDECL} preserve="1"{show}>'
            f'<p:cSld name="{escape(ld.name)}">{bg}<p:spTree><p:nvGrpSpPr><p:cNvPr id="1" name=""/>'
            "<p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr><p:grpSpPr><a:xfrm><a:off x=\"0\" y=\"0\"/>"
            '<a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/><a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr>'
            f"{sps}</p:spTree></p:cSld><p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>").encode()


def replace_layouts(prs: Any, defs: list[LayoutDef], boxes: dict[str, tuple[float, float, float]] | None = None,
                    color: str = "6B7280") -> None:
    master = prs.slide_masters[0]
    for layout in list(master.slide_layouts):
        master.slide_layouts.remove(layout)
    master_part = master.part
    package = master_part.package
    lst = master._element.get_or_add_sldLayoutIdLst()
    next_id = 2147483649
    for n, ld in enumerate(defs, start=1):
        part = SlideLayoutPart(PackURI(f"/ppt/slideLayouts/slideLayout{n}.xml"), CT.PML_SLIDE_LAYOUT, package,
                               parse_xml(layout_xml(ld, boxes, color)))
        part.relate_to(master_part, RT.SLIDE_MASTER)
        rid = master_part.relate_to(part, RT.SLIDE_LAYOUT)
        el = etree.SubElement(lst, qn("p:sldLayoutId"))
        el.set("id", str(next_id + n))
        el.set(qn("r:id"), rid)


# ---------------------------------------------------------------- master


def furniture_color(meta: dict[str, Any]) -> tuple[str, str]:
    """(color ref, hex) for slide numbers and footers: muted, or ink when muted is under 4.5:1 on the background."""
    palette = {k: str(v).lstrip("#").upper() for k, v in (meta.get("palette") or {}).items()}
    slots = slot_colors(meta)
    muted = palette.get("muted")
    if muted and contrast(muted, slots["lt1"]) >= 4.5:
        return "muted", muted
    return ("ink" if "ink" in palette else slots["dk1"]), slots["dk1"]


def _style_furniture(ph: Any, box: tuple[float, float, float], color: str) -> None:
    x, y, w = box
    ph.left, ph.top, ph.width, ph.height = emu(x), emu(y), emu(w), emu(FURNITURE_H)
    body = ph.text_frame._txBody
    styled = parse_xml(f"<p:txBody {NSDECL}>{_furniture_text(color)}</p:txBody>")
    for tag in ("a:bodyPr", "a:lstStyle"):
        body.replace(body.find(qn(tag)), styled.find(qn(tag)))


def style_master(prs: Any, w_in: float, h_in: float, boxes: dict[str, tuple[float, float, float]] | None = None,
                 color: str = "6B7280") -> None:
    """Fit the master's placeholders to the new size and set the type scale the layouts inherit."""
    master = prs.slide_masters[0]
    boxes = boxes or {}
    for ph in list(master.placeholders):
        t = ph.placeholder_format.type
        name = t.name if t is not None else ""
        if name == "TITLE":
            ph.left, ph.top, ph.width, ph.height = emu(0.6), emu(0.45), emu(w_in - 1.2), emu(1.05)
        elif name == "BODY":
            ph.left, ph.top, ph.width, ph.height = emu(0.6), emu(1.7), emu(w_in - 1.2), emu(h_in - 2.5)
        elif name in FURNITURE_KINDS:
            box = boxes.get(FURNITURE_KINDS[name])
            if box:
                _style_furniture(ph, box, color)
            else:
                ph._element.getparent().remove(ph._element)  # no date, ever; no number when they're off
    tx = master._element.find(qn("p:txStyles"))
    if tx is None:
        return
    sizes = {"titleStyle": [3000], "bodyStyle": [2000, 1800, 1600, 1400, 1400]}
    for style, szs in sizes.items():
        el = tx.find(qn(f"p:{style}"))
        if el is None:
            continue
        for lvl, sz in enumerate(szs, start=1):
            rpr = el.find(f"{qn(f'a:lvl{lvl}pPr')}/{qn('a:defRPr')}")
            if rpr is not None:
                rpr.set("sz", str(sz))
    title_ppr = tx.find(f"{qn('p:titleStyle')}/{qn('a:lvl1pPr')}")
    if title_ppr is not None:
        title_ppr.set("algn", "l")  # the stock master centers titles; body text is left-aligned
    title_rpr = tx.find(f"{qn('p:titleStyle')}/{qn('a:lvl1pPr')}/{qn('a:defRPr')}")
    if title_rpr is not None:
        for fill in title_rpr.findall(qn("a:solidFill")):
            title_rpr.remove(fill)
        fill = etree.SubElement(title_rpr, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:schemeClr"), val="tx2")
        title_rpr.insert(0, fill)


def add_master_logo(prs: Any, logo: Path, w_in: float, h_in: float) -> None:
    master = prs.slide_masters[0]
    image_part, rid = master.part.get_or_add_image_part(str(logo))
    with PILImage.open(logo) as im:
        iw, ih = im.size
    h = LOGO_H
    w = h * iw / ih
    x, y = w_in - 0.6 - w, h_in - LOGO_BOTTOM - h
    pic = parse_xml(
        f'<p:pic {NSDECL}><p:nvPicPr><p:cNvPr id="900" name="Logo" descr="Logo"/>'
        '<p:cNvPicPr><a:picLocks noChangeAspect="1"/></p:cNvPicPr><p:nvPr userDrawn="1"/></p:nvPicPr>'
        f'<p:blipFill><a:blip r:embed="{rid}"/><a:stretch><a:fillRect/></a:stretch></p:blipFill>'
        f'<p:spPr><a:xfrm><a:off x="{emu(x)}" y="{emu(y)}"/><a:ext cx="{emu(w)}" cy="{emu(h)}"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></p:spPr></p:pic>'
    )
    master.shapes._spTree.append(pic)


# ---------------------------------------------------------------- tokens


def tokens_for(defs: list[LayoutDef], meta: dict[str, Any]) -> dict[str, Any]:
    mapping = meta.get("theme_colors") or {}
    palette = meta.get("palette") or {}

    def name_for(slot: str, fallback: str) -> str:
        ref = mapping.get(slot)
        if ref:
            return str(ref)
        return fallback if fallback in palette else DEFAULT_SLOT_COLORS[slot]

    layouts: dict[str, Any] = {}
    for ld in defs:
        fields: dict[str, Any] = {}
        for ph in ld.phs:
            kind = KIND_FOR.get(ph.kind) or ("bullets" if ph.bullets else "text")
            spec: dict[str, Any] = {"idx": ph.idx, "kind": kind}
            if kind in ("text", "bullets"):
                spec["max_chars"] = estimate_chars(ph.w, ph.h, ph.size or 18)
            if kind == "bullets":
                spec["max_bullets"] = max(2, min(7, int(ph.h * 72 / ((ph.size or 18) * 1.2 * 1.5))))
                spec["max_bullet_chars"] = estimate_chars(ph.w, (ph.size or 18) * 2.4 / 72, ph.size or 18)
                spec["max_level"] = 1
            if kind == "table":
                spec["max_rows"] = 8
                spec["max_cols"] = 6
            if ld.key == "icon-row" and kind == "image":
                spec["kind"] = "icon"
            if ph.required:
                spec["required"] = True
            fields[ph.field] = spec
        entry: dict[str, Any] = {"template_layout": ld.name, "description": ld.description}
        if ld.heading != "title":
            entry["heading_field"] = ld.heading
        entry["fields"] = fields
        layouts[ld.key] = entry
    return {
        "spec_version": 1,
        "chart": {"font_size": 12, "text_color": name_for("dk1", "ink"),
                  "colors": [name_for("accent1", "primary"), name_for("accent2", "accent"),
                             name_for("accent3", "muted"), name_for("accent4", "accent4")]},
        "table": {"font_size": 14, "header_font_size": 14, "header_fill": name_for("dk2", "primary"),
                  "header_text": name_for("lt1", "background"), "band_fill": name_for("lt2", "surface"),
                  "status": {"Green": "15803D", "Amber": "B45309", "Red": "B91C1C"}},
        "layouts": layouts,
        "furniture": {"slide_numbers": bool((meta.get("generate") or {}).get("slide_numbers", True)),
                      "color": furniture_color(meta)[0]},  # decided here: muted, or ink when muted is too light
    }


# ---------------------------------------------------------------- package


def _finish(blob: bytes) -> bytes:
    """Mark the package as a template, drop the stock thumbnail, fix every date and timestamp."""
    when = "2000-01-01T00:00:00Z"
    parts = []
    for name, data in read_parts(blob):
        if name == "docProps/thumbnail.jpeg":
            continue
        if name == "[Content_Types].xml":
            data = data.replace(PPTX_CT, POTX_CT)
        elif name == "_rels/.rels":
            data = re.sub(rb'<Relationship [^>]*thumbnail[^>]*/>', b"", data)
        elif name == "docProps/core.xml":
            data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                          lambda m: m.group(1) + when.encode() + m.group(2), data)
        parts.append((name, data))
    return rezip(parts)  # rewrites every entry with the fixed timestamp


def generate(meta: dict[str, Any], source_dir: Path) -> tuple[bytes, dict[str, Any]]:
    gen = meta.get("generate") or {}
    w_emu, h_emu = SIZES_EMU[gen.get("slide_size", "16:9")]
    w_in, h_in = w_emu / EMU, h_emu / EMU
    defs = layout_set(gen.get("layout_set", "standard"), w_in, h_in)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(w_emu), Emu(h_emu)
    sld_sz = prs.part._element.find(qn("p:sldSz"))
    if sld_sz is not None and "type" in sld_sz.attrib:
        del sld_sz.attrib["type"]
    master = prs.slide_masters[0]
    theme_part = master.part.part_related_by(RT.THEME)
    write_theme(theme_part, slot_colors(meta), meta.get("fonts") or {})
    boxes = furniture_boxes(w_in, h_in, bool(gen.get("slide_numbers", True)))
    color = furniture_color(meta)[1]
    style_master(prs, w_in, h_in, boxes, color)
    replace_layouts(prs, defs, boxes, color)
    logo_id = gen.get("logo_on_master")
    if logo_id:
        rel = (meta.get("logos") or {}).get(logo_id)
        if rel:
            add_master_logo(prs, source_dir / rel, w_in, h_in)

    cp = prs.core_properties
    cp.title = str(meta.get("name", ""))
    cp.author = ""
    cp.last_modified_by = ""
    cp.comments = ""
    cp.revision = 1
    buf = io.BytesIO()
    prs.save(buf)
    return _finish(buf.getvalue()), tokens_for(defs, meta)
