"""Read a template: layouts, placeholders, theme colors and fonts. Feeds `inspect` and `brand adopt`."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any

from lxml import etree
from pptx.enum.shapes import PP_PLACEHOLDER

from deck_builder import template as tpl

EMU_PER_INCH = 914400
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
KIND = {
    PP_PLACEHOLDER.TITLE: "text",
    PP_PLACEHOLDER.CENTER_TITLE: "text",
    PP_PLACEHOLDER.SUBTITLE: "text",
    PP_PLACEHOLDER.BODY: "bullets",
    PP_PLACEHOLDER.OBJECT: "bullets",
    PP_PLACEHOLDER.PICTURE: "image",
    PP_PLACEHOLDER.CHART: "chart",
    PP_PLACEHOLDER.TABLE: "table",
}
SKIP = {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}
THEME_SLOTS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
               "hlink", "folHlink")
SLOT_NAMES = {"dk1": "ink", "lt1": "background", "dk2": "primary", "lt2": "surface", "hlink": "link",
              "folHlink": "link-visited"}
DEFAULT_SIZE = {"title": 32.0, "body": 18.0}


def slug(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-") or "layout"


def estimate_chars(w_in: float, h_in: float, font_pt: float) -> int:
    """Rough capacity: lines that fit times characters per line, less 20% for wrapping. Tune after a render."""
    lines = max(1, int(h_in * 72 / (font_pt * 1.2)))
    per_line = max(1, int(w_in * 72 / (font_pt * 0.5)))
    return max(4, int(lines * per_line * 0.8))


def font_size(ph: Any, is_title: bool) -> float:
    """The placeholder's own first-level size if it sets one, else a typical default."""
    sz = ph._element.xpath(".//a:lstStyle/a:lvl1pPr/a:defRPr/@sz")
    return int(sz[0]) / 100 if sz else DEFAULT_SIZE["title" if is_title else "body"]


def field_spec(ph: Any) -> tuple[str, dict[str, Any]]:
    pf = ph.placeholder_format
    kind = KIND.get(pf.type, "bullets")
    is_title = pf.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
    spec: dict[str, Any] = {"idx": pf.idx, "kind": kind}
    if kind in ("text", "bullets"):
        w, h = (ph.width or 0) / EMU_PER_INCH, (ph.height or 0) / EMU_PER_INCH
        spec["max_chars"] = estimate_chars(w, h, font_size(ph, is_title))
    if kind == "bullets":
        spec["max_bullets"] = 6
    role = "title" if is_title else "subtitle" if pf.type == PP_PLACEHOLDER.SUBTITLE else kind
    return role, spec


def layouts_report(path: Path) -> list[dict[str, Any]]:
    prs = tpl.open_template(path)
    out = []
    for mi, master in enumerate(prs.slide_masters):
        for li, layout in enumerate(master.slide_layouts):
            phs = []
            for ph in layout.placeholders:
                pf = ph.placeholder_format
                if pf.type in SKIP:
                    continue
                phs.append({
                    "idx": pf.idx,
                    "type": pf.type.name if pf.type else "?",
                    "name": ph.name,
                    "at_in": [round((ph.left or 0) / EMU_PER_INCH, 2), round((ph.top or 0) / EMU_PER_INCH, 2)],
                    "size_in": [round((ph.width or 0) / EMU_PER_INCH, 2), round((ph.height or 0) / EMU_PER_INCH, 2)],
                })
            out.append({"master": mi, "index": li, "name": layout.name, "placeholders": phs})
    return out


def starter_tokens(path: Path) -> dict[str, Any]:
    prs = tpl.open_template(path)
    layouts: dict[str, Any] = {}
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            fields: dict[str, Any] = {}
            counts: dict[str, int] = {}
            for ph in layout.placeholders:
                if ph.placeholder_format.type in SKIP:
                    continue
                role, spec = field_spec(ph)
                counts[role] = counts.get(role, 0) + 1
                name = "body" if role == "bullets" and counts[role] == 1 else role
                if counts[role] > 1:
                    name = f"{'body' if role == 'bullets' else role}{counts[role]}"
                if role == "title":
                    spec["required"] = True
                fields[name] = spec
            if not fields:
                continue
            key = slug(layout.name)
            while key in layouts:
                key += "-2"
            layouts[key] = {"template_layout": layout.name, "fields": fields}
    return {"spec_version": 1, "layouts": layouts}


def theme(path: Path) -> dict[str, Any]:
    """Theme colors (slot -> hex) and heading/body fonts of the first master's theme."""
    with zipfile.ZipFile(path) as z:
        names = sorted(n for n in z.namelist() if re.fullmatch(r"ppt/theme/theme\d+\.xml", n))
        if not names:
            return {"colors": {}, "fonts": {}}
        root = etree.fromstring(z.read(names[0]))
    colors: dict[str, str] = {}
    for slot in THEME_SLOTS:
        el = root.find(f".//a:clrScheme/a:{slot}", NS)
        if el is None:
            continue
        srgb = el.find("a:srgbClr", NS)
        sys = el.find("a:sysClr", NS)
        val = srgb.get("val") if srgb is not None else (sys.get("lastClr") if sys is not None else None)
        if val:
            colors[slot] = val.upper()
    fonts = {}
    for role, tag in (("heading", "majorFont"), ("body", "minorFont")):
        latin = root.find(f".//a:fontScheme/a:{tag}/a:latin", NS)
        if latin is not None and latin.get("typeface"):
            fonts[role] = latin.get("typeface")
    return {"colors": colors, "fonts": fonts}


def brand_skeleton(slug_: str, path: Path) -> dict[str, Any]:
    th = theme(path)
    palette: dict[str, str] = {}
    theme_colors: dict[str, str] = {}
    for slot, hexv in th["colors"].items():
        name = SLOT_NAMES.get(slot, slot)
        palette[name] = hexv
        theme_colors[slot] = name
    meta: dict[str, Any] = {
        "spec_version": 1,
        "name": slug_.replace("-", " ").title(),
        "slug": slug_,
        "version": "0.1.0",
        "description": f"Adopted from {path.name}",
        "palette": palette or {"primary": "1F3A5F"},
        "theme_colors": theme_colors,
        "voice": ["One idea per slide. The title states the takeaway."],
        "lint": {"max_slides": 40, "banned_patterns": []},
    }
    if th["fonts"]:
        meta["fonts"] = {role: {"family": fam} for role, fam in th["fonts"].items()}
    return meta
