"""`inspect --text` and `inspect --index`: the words on every slide of any .pptx, including one a person made.

Text keeps the inline markup `import` writes (`code`, **bold**, *italic*, links), so a dump of a hand-edited
deck can be compared with deck.md. Nothing is written; the file is read as untrusted.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.oxml.ns import qn

from deck_builder import template as tpl
from deck_builder.importer import TITLES, paragraph_md

P14 = "{http://schemas.microsoft.com/office/powerpoint/2010/main}"
TYPED_NUMBER = re.compile(r"^\d{1,3}$")  # a text box holding only a number, typed where the field isn't used


def _sections(prs: Any) -> dict[int, str]:
    """Slide id -> the name of the PowerPoint section it's in."""
    out: dict[int, str] = {}
    for sec in prs.part._element.iter(f"{P14}section"):
        for sid in sec.iter(f"{P14}sldId"):
            out[int(sid.get("id"))] = str(sec.get("name") or "")
    return out


def _lines(tf: Any) -> list[str]:
    return [f"{'  ' * p.level}{md}" for p in tf.paragraphs if (md := paragraph_md(p, ""))]


def _shapes(shapes: Any) -> list[dict[str, Any]]:
    """Each shape with words, in the slide's reading order; a group's shapes come in its place."""
    out: list[dict[str, Any]] = []
    for sh in shapes:
        ph = sh.placeholder_format if sh.is_placeholder else None
        kind = "title" if ph is not None and ph.type in TITLES else \
            "slide number" if ph is not None and ph.type == PP_PLACEHOLDER.SLIDE_NUMBER else "text"
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            out += _shapes(sh.shapes)
        elif getattr(sh, "has_table", False) and sh.has_table:
            rows = [" | ".join(" ".join(_lines(c.text_frame)) for c in r.cells) for r in sh.table.rows]
            out.append({"name": sh.name, "kind": "table", "lines": rows})
        elif getattr(sh, "has_chart", False) and sh.has_chart:
            ct = sh.chart.chart_title if sh.chart.has_title else None
            title = ct.text_frame.text.strip() if ct is not None and ct.has_text_frame else ""
            out.append({"name": sh.name, "kind": "chart", "lines": [title] if title else []})
        elif sh._element.tag == qn("p:pic"):
            alt = sh._element.find(f"{qn('p:nvPicPr')}/{qn('p:cNvPr')}").get("descr") or ""
            out.append({"name": sh.name, "kind": "picture", "lines": [alt] if alt else []})
        elif getattr(sh, "has_text_frame", False) and sh.has_text_frame and (lines := _lines(sh.text_frame)):
            out.append({"name": sh.name, "kind": kind, "lines": lines})
    return out


def _number(shapes: list[dict[str, Any]]) -> str | None:
    """The number shown on the slide: its slide-number box, or else a text box holding only a number."""
    for kinds in (("slide number",), ("text",)):
        for sh in shapes:
            if sh["kind"] in kinds and len(sh["lines"]) == 1 and TYPED_NUMBER.match(sh["lines"][0].strip()):
                return sh["lines"][0].strip()
    return None


def slides(path: Path) -> list[dict[str, Any]]:
    """Every slide: position, layout, section, shown number, title, hidden, each shape's words, and notes."""
    prs = tpl.open_template(path)
    sections = _sections(prs)
    out = []
    for n, s in enumerate(prs.slides, start=1):
        shapes = _shapes(s.shapes)
        title = next((" ".join(sh["lines"]) for sh in shapes if sh["kind"] == "title"), None)
        notes = s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide and \
            s.notes_slide.notes_text_frame is not None else ""
        out.append({"slide": n, "layout": s.slide_layout.name, "section": sections.get(s.slide_id),
                    "number": _number(shapes), "title": title, "hidden": s._element.get("show") == "0",
                    "shapes": shapes, "notes": notes})
    return out


def text_report(found: list[dict[str, Any]]) -> str:
    lines: list[str] = []
    for s in found:
        extra = [f"number {s['number']}"] * bool(s["number"]) + [f"section {s['section']!r}"] * bool(s["section"])
        lines.append(f"--- slide {s['slide']}{' (hidden)' if s['hidden'] else ''}: layout {s['layout']!r}"
                     + "".join(f", {e}" for e in extra))
        for sh in s["shapes"]:
            if sh["kind"] == "slide number":
                continue
            label = sh["kind"] if sh["kind"] in ("title", "table", "chart", "picture") else f"text {sh['name']!r}"
            if len(sh["lines"]) == 1:
                lines.append(f"{label}: {sh['lines'][0]}")
            else:
                lines += [f"{label}:", *(f"  {t}" for t in sh["lines"])]
        if s["notes"]:
            lines.append("notes:")
            lines += [f"  {t}" for t in s["notes"].splitlines() if t.strip()]
    return "\n".join(lines)


def index_report(found: list[dict[str, Any]]) -> str:
    return "\n".join(f"{s['slide']:>3}  number {s['number'] or '-':<4} {s['section'] or '-'} | "
                     f"{s['title'] or '(no title)'}{' (hidden)' if s['hidden'] else ''}" for s in found)
