"""Render a built deck: PDF, slide PNGs, contact sheets, measured overflow, font check, flagged slides."""
from __future__ import annotations

import json
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER

from deck_builder import confine
from deck_builder.brand.registry import Brand
from deck_builder.build.deck import manifest_path
from deck_builder.errors import EnvError, Issue
from deck_builder.qa import backends, fonts, images, measure

FLAG_CODES = {"OVERFLOW_MEASURED", "EMPTY_PLACEHOLDER", "ASSET_LOW_RES"}
VISUAL = {PP_PLACEHOLDER.PICTURE, PP_PLACEHOLDER.CHART, PP_PLACEHOLDER.TABLE}
RENDER_FILES = re.compile(r"deck\.pdf|slide-\d+\.png|contact-\d+\.png")


@dataclass
class Rendered:
    backend: str
    out_dir: Path
    pdf: Path
    slides: list[Path]
    contact_sheets: list[Path]
    issues: list[Issue] = field(default_factory=list)
    hidden: list[int] = field(default_factory=list)

    @property
    def flagged(self) -> dict[int, list[str]]:
        out: dict[int, list[str]] = {}
        for i in self.issues:
            if i.code in FLAG_CODES and i.slide is not None:
                out.setdefault(i.slide, []).append(i.code)
        return dict(sorted(out.items()))


def hidden_slide_numbers(prs: Any) -> list[int]:
    """Slide numbers (1-indexed) marked hidden (`<p:sld show="0">`): LibreOffice leaves these out of the
    PDF entirely, so PDF page order skips them."""
    return [n for n, slide in enumerate(prs.slides, start=1) if slide._element.get("show") == "0"]


def stale_build_check(pptx: Path, manifest: dict[str, Any] | None) -> list[Issue]:
    """Warn when the deck source the manifest points at was edited after this .pptx was built."""
    if not manifest:
        return []
    src = manifest.get("input", {}).get("path")
    if not src:
        return []
    source = Path(src)
    if not source.is_file():
        return []
    if source.stat().st_mtime > pptx.stat().st_mtime:
        return [Issue("STALE_BUILD", f"{source.name} was edited after {pptx.name} was built",
                     severity="warning")]
    return []


def empty_placeholders(pptx: Path) -> list[Issue]:
    issues = []
    for n, slide in enumerate(Presentation(str(pptx)).slides, start=1):
        for ph in slide.placeholders:
            unfilled_visual = ph._element.tag.endswith("}sp")  # filled ones become p:pic or p:graphicFrame
            if ph.placeholder_format.type in VISUAL:
                if unfilled_visual:
                    issues.append(Issue("EMPTY_PLACEHOLDER", f"an unfilled {ph.name!r} is on the slide", slide=n))
            elif ph.has_text_frame and not ph.text_frame.text.strip():
                issues.append(Issue("EMPTY_PLACEHOLDER", f"{ph.name!r} has no text", slide=n))
    return issues


def render(pptx: Path, backend: str, dpi: int, batch: int, pages: list[int] | None,
           brand: Brand | None) -> Rendered:
    """PDF, slide PNGs and contact sheets in <deck>.render/, plus measured overflow, empty placeholders and fonts."""
    prs = Presentation(str(pptx))
    total = len(prs.slides)
    past = [str(n) for n in pages or [] if n > total]
    if past:
        raise EnvError(f"--slides {','.join(past)}: the deck has {total} slides")
    if dpi < 1:
        raise EnvError(f"dpi must be 1 or more, not {dpi}")
    if batch < 1:
        raise EnvError(f"contact_batch in deck-builder.toml must be 1 or more, not {batch}")
    out_dir = pptx.with_name(pptx.stem + ".render")
    # out_dir sits beside pptx and can already be an existing symlink; confine it (and so what render
    # deletes inside it) before touching it. A no-op outside an MCP call.
    confine.guard(out_dir, "the render folder")
    if out_dir.exists():
        # only ever touch a folder this command made (it holds deck.pdf, or nothing), and in it only the files
        # render writes, so anything else kept there survives
        if not (out_dir / "deck.pdf").is_file() and any(out_dir.iterdir()):
            raise EnvError(f"{out_dir} exists and wasn't made by deck-builder render; move it and retry")
        for f in out_dir.iterdir():
            if f.is_file() and RENDER_FILES.fullmatch(f.name):
                f.unlink()
    out_dir.mkdir(parents=True, exist_ok=True)
    pdf = out_dir / "deck.pdf"
    backends.to_pdf(backend, pptx, pdf)
    hidden = hidden_slide_numbers(prs)
    visible = [n for n in range(1, total + 1) if n not in hidden]
    w_emu, h_emu = int(prs.slide_width or 0), int(prs.slide_height or 0)
    size = (round(w_emu / 914400 * dpi), round(h_emu / 914400 * dpi))  # 13.33 in x 96 dpi = 1280 px
    pngs = images.rasterize(pdf, out_dir, size, pages, total, visible)
    issues: list[Issue] = []
    if backend == "powerpoint":
        issues.append(Issue("RENDER_UNVERIFIED", "the PowerPoint backend hasn't been verified on a real Mac yet",
                            severity="warning"))
    mpath = manifest_path(pptx)
    manifest: dict[str, Any] | None = json.loads(mpath.read_text()) if mpath.is_file() else None
    issues += stale_build_check(pptx, manifest)
    issues += measure.overflow(pptx, pdf, manifest, visible)
    issues += empty_placeholders(pptx)
    if brand is not None:
        issues += fonts.check(pdf, brand.meta)
    if pages:
        issues = [i for i in issues if i.slide is None or i.slide in pages]
    result = Rendered(backend, out_dir, pdf, pngs, [], issues, hidden)
    result.contact_sheets = images.contact_sheets(pngs, out_dir, batch, set(result.flagged))
    return result
