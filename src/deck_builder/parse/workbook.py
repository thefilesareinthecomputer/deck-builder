"""deck.xlsx -> Deck. Reads cell values only, so re-saving in Excel or LibreOffice changes nothing."""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook

from deck_builder.errors import EnvError, Issue
from deck_builder.model import Chart, Deck, Series, Table, Value
from deck_builder.parse.csvfile import slides_from_rows
from deck_builder.parse.markdown import number, substitute


def cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "true" if v else "false"
    if isinstance(v, float | int):
        return str(number(v))
    if isinstance(v, datetime | date):
        return v.isoformat()
    return str(v)


def _rows(ws: Any) -> list[list[Any]]:
    return [list(r) for r in ws.iter_rows(values_only=True)]


def _flag(v: Any) -> bool | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        return v
    loaded = yaml.safe_load(str(v))
    return bool(loaded) if isinstance(loaded, bool) else None


def chart_from_rows(rows: list[list[Any]]) -> Chart:
    opts: dict[str, Any] = {}
    i = 0
    while i < len(rows) and any(c not in (None, "") for c in rows[i]):
        key = cell_str(rows[i][0]).strip()
        opts[key] = rows[i][1] if len(rows[i]) > 1 else None
        i += 1
    i += 1  # the blank row
    c = Chart(type=cell_str(opts.get("type")) or "column",
              number_format=cell_str(opts.get("number_format")) or "General",
              labels=bool(_flag(opts.get("labels"))),
              legend=_flag(opts.get("legend")),
              title=cell_str(opts.get("title")) or None,
              colors=[x.strip() for x in cell_str(opts.get("colors")).split(",") if x.strip()] or None)
    if i >= len(rows):
        return c
    names = [cell_str(x) for x in rows[i][1:] if x not in (None, "")]
    c.series = [Series(name=n, values=[]) for n in names]
    for row in rows[i + 1 :]:
        if all(x in (None, "") for x in row):
            continue
        c.categories.append(cell_str(row[0]))
        for j, s in enumerate(c.series, start=1):
            v = row[j] if j < len(row) else None
            s.values.append(number(v) if isinstance(v, int | float) and not isinstance(v, bool) else cell_str(v))
    return c


def table_from_rows(rows: list[list[Any]]) -> Table:
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        return Table(header=[], rows=[])
    width = max(i + 1 for r in rows for i, c in enumerate(r) if c not in (None, ""))
    cells = [[cell_str(c) for c in (r + [None] * width)[:width]] for r in rows]
    return Table(header=cells[0], rows=cells[1:])


def parse(path: Path, row: dict[str, str] | None = None) -> tuple[Deck, list[Issue]]:
    """deck.xlsx -> Deck: the `deck` sheet is front matter, `slides` the slides, chart and table sheets by name."""
    issues: list[Issue] = []
    name = path.name
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises several types for a bad file
        raise EnvError(f"{path}: not a readable .xlsx workbook ({e})") from e
    if "slides" not in wb.sheetnames:
        issues.append(Issue("PARSE", "no `slides` sheet", file=name))
        return Deck(meta={}, slides=[], source=str(path)), issues

    def sub(text: str) -> str:
        return substitute(text, row, name, issues)

    meta: dict[str, Any] = {}
    if "deck" in wb.sheetnames:
        for r in _rows(wb["deck"])[1:]:
            if not r or r[0] in (None, ""):
                continue
            raw = sub(cell_str(r[1] if len(r) > 1 else None))
            try:
                meta[cell_str(r[0]).strip()] = yaml.safe_load(raw) if raw else None
            except yaml.YAMLError as e:
                issues.append(Issue("PARSE", f"deck sheet, key {r[0]!r}: {e}", file=name))

    grid = _rows(wb["slides"])
    if not grid:
        issues.append(Issue("PARSE", "the slides sheet is empty", file=name))
        return Deck(meta=meta, slides=[], source=str(path)), issues
    header = [cell_str(h).strip() for h in grid[0]]
    records = []
    for rownum, r in enumerate(grid[1:], start=2):
        if not any(c not in (None, "") for c in r):
            continue
        records.append((rownum, {h: sub(cell_str(v)) for h, v in zip(header, r, strict=False) if h}))

    def resolve_sheet(ref: str) -> Value | None:
        sheet = ref[len("sheet:"):].strip()
        if sheet not in wb.sheetnames:
            return None
        rows = _rows(wb[sheet])
        return chart_from_rows(rows) if sheet.lower().startswith("chart") else table_from_rows(rows)

    if any(i.code in ("UNKNOWN_TOKEN", "BAD_DATA_VALUE") for i in issues):
        return Deck(meta={}, slides=[], source=str(path)), issues
    slides = slides_from_rows(records, name, issues, sheet_resolver=resolve_sheet)
    wb.close()
    if not slides:
        issues.append(Issue("PARSE", "no slide rows", file=name))
    return Deck(meta=meta, slides=slides, source=str(path)), issues


def read_table_rows(path: Path) -> list[dict[str, str]]:
    """The first sheet as rows of strings, for bulk --data."""
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises several types for a bad file
        raise EnvError(f"{path}: not a readable .xlsx workbook ({e})") from e
    rows = _rows(wb.worksheets[0])
    wb.close()
    if not rows:
        return []
    header = [cell_str(h).strip() for h in rows[0]]
    return [{h: cell_str(v) for h, v in zip(header, r, strict=False) if h}
            for r in rows[1:] if any(c not in (None, "") for c in r)]
